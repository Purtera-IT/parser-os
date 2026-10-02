"""One line, many documents: the earliest owns it, the rest keep a copy.

A sentence that appears in two documents is one fact for counting, scope and
pricing -- but it is a line in BOTH documents, and a person labelling the
later one has to see it there. Live 000132:

* a HubSpot note written in May was pasted into a June email; the dedup kept
  the email's copy (the later, larger document) and the note showed 5 of its
  13 lines;
* the v2 SOW repeats v1's clauses; 30 of v2's lines were folded onto v1 and
  vanished from v2, and 19 of the Deal Kit's went the other way.

Two rules, applied by every dedup stage that can fold across documents:

1. **The earliest document owns the canonical atom.** Earliest by the
   document's own date (the note's date, the email's sent date, the file's
   stated or authored date), falling back to the order the documents were
   read in. :func:`document_order` builds that order once per compile, and
   the stages pick their survivor from the earliest document first.

2. **A cross-document fold leaves a copy, not a hole.** The folded atom is
   kept -- with its own document and locator -- flagged
   :data:`COPY_FLAG` and pointing at the canonical atom through
   ``value["duplicate_of"]``. :func:`split_copies` finds them after a stage.
   The compiler holds copies out of every later stage (exactly as it holds
   admission chatter) and puts them back only for the result, so no count,
   scope roll-up or price ever reads one twice.

Within one document nothing changes: two copies of a line in the same file
are still folded into one atom and recorded in the suppression ledger.
"""
from __future__ import annotations

import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable, Iterable, Mapping

#: Marks an atom kept only so its document shows its own copy of a line that
#: another (earlier) document owns. Consumers that count, price or roll up
#: scope must skip it; the labelling walk lists it under its own document.
COPY_FLAG = "cross_doc_copy"

#: An order key for a document nothing could date or place: after everything.
_LAST = (2, 0.0, 10**9)


def is_cross_doc_copy(atom: Any) -> bool:
    return COPY_FLAG in (getattr(atom, "review_flags", None) or [])


def _value(atom: Any) -> dict:
    v = getattr(atom, "value", None)
    return v if isinstance(v, dict) else {}


def _parse_date(raw: Any) -> datetime | None:
    if not raw:
        return None
    text = str(raw).strip()
    dt = None
    try:
        dt = datetime.fromisoformat(text.replace("Z", "+00:00"))
    except Exception:
        try:
            from email.utils import parsedate_to_datetime

            dt = parsedate_to_datetime(text)
        except Exception:
            try:
                from dateutil import parser as _dp

                dt = _dp.parse(text, fuzzy=True)
            except Exception:
                return None
    if dt is None:
        return None
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt


def _own_date(items: list[Any]) -> datetime | None:
    """The date a document carries about ITSELF, read off its atoms.

    A note's CRM header; an email's own (outermost, unquoted) message; a
    transcript's header. A quoted message's date is the date of the message
    it quotes, not of this document, so it is never read here.
    """
    for a in items:
        v = _value(a)
        if str(v.get("kind") or "") == "hubspot_note_meta":
            d = _parse_date(v.get("date"))
            if d is not None:
                return d
    for a in items:
        v = _value(a)
        if v.get("quoted"):
            continue
        if str(v.get("kind") or "") == "email_header":
            d = _parse_date(v.get("date"))
            if d is not None:
                return d
    dates = []
    for a in items:
        v = _value(a)
        if v.get("quoted"):
            continue
        et = v.get("email_thread") if isinstance(v.get("email_thread"), dict) else {}
        msg = et.get("message") if isinstance(et.get("message"), dict) else {}
        if msg.get("quoted"):
            continue
        for raw in (msg.get("sent_at"), v.get("authored_at"), et.get("date"), v.get("note_date")):
            d = _parse_date(raw)
            if d is not None:
                dates.append(d)
                break
    if dates:
        # The outermost message is the newest one the file holds.
        return max(dates)
    for a in items:
        v = _value(a)
        if str(v.get("kind") or "").endswith("_header") and v.get("document_date"):
            d = _parse_date(v.get("document_date"))
            if d is not None:
                return d
    return None


def _filename(items: list[Any]) -> str:
    for a in items:
        for r in getattr(a, "source_refs", None) or []:
            name = str(getattr(r, "filename", "") or "")
            if name:
                return name
    return ""


def document_order(
    atoms: Iterable[Any],
    *,
    provenance: Mapping[str, Mapping[str, Any]] | None = None,
) -> dict[str, tuple]:
    """``artifact_id -> sort key``; a smaller key is an earlier document.

    Dated documents sort by their own date; undated ones follow, in the order
    the compile read them. ``provenance`` is the manifest's per-filename
    record (``authored_at``), used when the atoms state no date of their own.
    """
    by_doc: dict[str, list[Any]] = {}
    first_seen: dict[str, int] = {}
    for i, a in enumerate(atoms):
        aid = str(getattr(a, "artifact_id", "") or "")
        if not aid:
            continue
        by_doc.setdefault(aid, []).append(a)
        first_seen.setdefault(aid, i)
    out: dict[str, tuple] = {}
    for aid, items in by_doc.items():
        when = _own_date(items)
        if when is None and provenance:
            rec = provenance.get(_filename(items)) or {}
            when = _parse_date(rec.get("authored_at"))
        if when is None:
            out[aid] = (1, 0.0, first_seen[aid])
        else:
            try:
                ts = when.timestamp()
            except Exception:
                ts = 0.0
            out[aid] = (0, ts, first_seen[aid])
    return out


def doc_key(atom: Any, order: Mapping[str, tuple] | None) -> tuple:
    """The atom's document's position in ``order`` (all equal when no order)."""
    if not order:
        return (0, 0.0, 0)
    return order.get(str(getattr(atom, "artifact_id", "") or ""), _LAST)


def _ref_key(ref: Any) -> tuple:
    loc = getattr(ref, "locator", None) or {}
    try:
        loc_key = tuple(sorted((str(k), repr(v)) for k, v in dict(loc).items()))
    except Exception:
        loc_key = ()
    return (str(getattr(ref, "id", "") or ""), str(getattr(ref, "filename", "") or ""), loc_key)


#: A leading list marker: a bullet glyph ("- ", "* ", "• ") or an
#: ordinal ("1. ", "2) ", "(3) "). Needs whitespace after it, so "1.5 hours",
#: "-48V" and "*required" are left alone.
_LIST_MARKER_RE = re.compile(
    r"^\s*(?:[-*•·▪◦‣‧∙⁃–—+>]+|\(?\d{1,3}[.)])\s+"
)


def strip_list_marker(text: Any) -> str:
    """The line without its leading list marker(s), for COMPARISON keys only.

    A note bullet "- Install 4 APs" and the email line "Install 4 APs" are
    one line; so are "1. Install 4 APs" and "2) Install 4 APs" in two pastes
    of a list renumbered differently. Folding to ``[a-z0-9]`` drops a glyph
    but keeps an ordinal, and a key truncated BEFORE folding shifts by the
    marker's width -- either way the copies used to key apart. Displayed atom
    text is never touched.
    """
    s = str(text or "")
    for _ in range(3):  # "- 1. item" nests
        t = _LIST_MARKER_RE.sub("", s, count=1)
        if t == s or not t.strip():
            break
        s = t
    return s


def _text_key(atom: Any) -> str:
    t = getattr(atom, "raw_text", "") or getattr(atom, "normalized_text", "") or ""
    return re.sub(r"[^a-z0-9]+", " ", strip_list_marker(t).lower()).strip()


def _same_words(a: str, b: str) -> bool:
    if not a or not b:
        return False
    if a == b or a in b or b in a:
        return True
    wa, wb = set(a.split()), set(b.split())
    return bool(wa and wb) and len(wa & wb) / len(wa | wb) >= 0.6


def _survivor(dropped: Any, kept_by_ref: dict[tuple, list[Any]], kept_by_text: dict[str, list[Any]]) -> Any | None:
    """The kept atom ``dropped`` was folded into, when that is another document's.

    A fold carries the loser's source refs onto the winner, so the winner is
    the kept atom holding the loser's primary ref. Every atom of a document can
    share that document's ref (a note's atoms all cite the note), so a
    candidate in the loser's own document is not evidence of a cross-document
    fold, and the words must agree. With no ref trail, the same words in a
    kept atom of another document stand in.
    """
    own = str(getattr(dropped, "artifact_id", "") or "")
    words = _text_key(dropped)
    if not words:
        return None
    # The same line survives in its own document: that was an in-document
    # fold, which stays a suppression.
    for k in kept_by_text.get(words, ()):
        if str(getattr(k, "artifact_id", "") or "") == own:
            return None
    refs = list(getattr(dropped, "source_refs", None) or [])
    if refs:
        for cand in kept_by_ref.get(_ref_key(refs[0]), ()):
            if str(getattr(cand, "artifact_id", "") or "") == own:
                continue
            if _same_words(words, _text_key(cand)):
                return cand
    for cand in kept_by_text.get(words, ()):
        if str(getattr(cand, "artifact_id", "") or "") != own:
            return cand
    return None


def split_copies(before: list[Any], after: list[Any], *, stage: str) -> list[Any]:
    """The atoms a stage folded onto ANOTHER document's atom, marked as copies.

    Each returned atom keeps its own ``artifact_id`` and locator, gains
    :data:`COPY_FLAG`, and records the canonical atom in
    ``value["duplicate_of"]``; the canonical atom lists the documents that
    also carry the line in ``value["also_in_documents"]``.

    Not copies, and left to the suppression ledger: an in-document fold,
    held chatter, and anything whose survivor cannot be found.
    """
    after_ids = {id(a) for a in after}
    dropped = [a for a in before if id(a) not in after_ids]
    if not dropped:
        return []
    kept_by_ref: dict[tuple, list[Any]] = {}
    kept_by_text: dict[str, list[Any]] = {}
    for k in after:
        for r in getattr(k, "source_refs", None) or []:
            kept_by_ref.setdefault(_ref_key(r), []).append(k)
        kept_by_text.setdefault(_text_key(k), []).append(k)

    copies: list[Any] = []
    for d in dropped:
        v = _value(d)
        if is_cross_doc_copy(d):
            continue
        if "admission_regex" in (getattr(d, "review_flags", None) or []):
            continue
        w = _survivor(d, kept_by_ref, kept_by_text)
        if w is None:
            continue
        if _is_quoted_mail_echo(d):
            # An email quoting another message: the line is that message's,
            # and belongs in ITS section, not shown again as the reply's
            # (010003). The survivor remembers the reply so neither the sweep
            # below nor the envelope credits the line to it. A quoted line in
            # a NOTE is not an echo: the note shows its own copy (000132).
            _note_quoted_in(w, d)
            continue
        _mark_copy(d, w, stage)
        copies.append(d)
    return copies


def _mark_copy(d: Any, w: Any, stage: str) -> None:
    """Make ``d`` its document's copy of canonical atom ``w``."""
    v = dict(_value(d))
    v["duplicate_of"] = {
        "atom_id": str(getattr(w, "id", "") or ""),
        "artifact_id": str(getattr(w, "artifact_id", "") or ""),
        "stage": stage,
    }
    d.value = v
    flags = list(getattr(d, "review_flags", None) or [])
    if COPY_FLAG not in flags:
        flags.append(COPY_FLAG)
    d.review_flags = flags
    wv = _value(w)
    if isinstance(getattr(w, "value", None), dict):
        docs = list(wv.get("also_in_documents") or [])
        aid = str(getattr(d, "artifact_id", "") or "")
        if aid and aid not in docs:
            docs.append(aid)
            wv["also_in_documents"] = docs


from app.core.suppression_ledger import SURVIVOR_KEY  # noqa: E402  (the kept atom a fold went into)


def settle_folds(
    before: list[Any],
    after: list[Any],
    copies: list[Any],
    folds: Mapping[int, tuple[Any, Any]],
    *,
    stage: str,
    standing: Iterable[Any] = (),
    make_copies: bool = True,
) -> tuple[list[Any], list[Any], list[Any]]:
    """Every atom a dedup stage dropped either names a survivor or comes back.

    ``folds`` is ``id(loser) -> (loser, winner)`` as the stage's folds recorded
    it (``semantic_dedup.take_folds``). For each atom in ``before`` that is in
    neither ``after`` nor ``copies``:

    * its survivor is the kept atom its fold chain ends at (a winner that was
      itself folded hands on to ITS winner; a winner kept as a copy hands on
      to that copy's canonical atom), or else a kept atom with the same words;
    * survivor in the SAME document (or the atom is a quoted echo of another
      message): it stays suppressed and records the survivor under
      :data:`SURVIVOR_KEY`;
    * survivor in ANOTHER document: it is kept as that document's copy
      (:data:`COPY_FLAG`), never suppressed -- a document keeps its own line;
    * no survivor: it was folded into nothing that stands, so it is put back
      where it was.

    An atom the stage dropped on purpose rather than folded (a hallucinated
    site; see ``semantic_dedup.mark_dropped_not_folded``) is left alone.

    ``standing`` are atoms outside ``after`` that a fold may also name (a
    stage run over a side list -- held chatter, vision transcriptions --
    folds onto the main list). ``make_copies=False`` keeps a fold across
    documents suppressed, naming its survivor, instead of making a copy (a
    quoted echo of another message is that message's line).
    Returns ``(after_with_restored, new_copies, restored)``.
    """
    from app.core.suppression_ledger import DROPPED_NOT_FOLDED_KEY

    after = list(after)
    extra = [a for a in standing if a is not None]
    kept_ids = {id(a) for a in after} | {id(a) for a in extra}
    copy_ids = {id(c) for c in copies}
    kept_by_atom_id = {str(getattr(a, "id", "") or ""): a for a in after + extra}
    by_text: dict[str, list[Any]] = {}
    for a in after + extra:
        by_text.setdefault(_text_key(a), []).append(a)

    def _standing(atom: Any) -> Any | None:
        seen: set[int] = set()
        cur = atom
        while cur is not None and id(cur) not in seen and len(seen) < 32:
            seen.add(id(cur))
            if id(cur) in kept_ids:
                return cur
            if id(cur) in copy_ids:
                dup = _value(cur).get("duplicate_of") or {}
                return kept_by_atom_id.get(str(dup.get("atom_id") or ""))
            nxt = folds.get(id(cur))
            cur = nxt[1] if nxt and nxt[0] is cur else None
        return None

    new_copies: list[Any] = []
    restored: list[Any] = []
    for d in before:
        if id(d) in kept_ids or id(d) in copy_ids:
            continue
        if _value(d).get(DROPPED_NOT_FOLDED_KEY):
            continue
        own = str(getattr(d, "artifact_id", "") or "")
        nxt = folds.get(id(d))
        w = _standing(nxt[1]) if nxt and nxt[0] is d else None
        if w is None:
            same = by_text.get(_text_key(d), ()) if _text_key(d) else ()
            w = next((k for k in same if str(getattr(k, "artifact_id", "") or "") == own),
                     next(iter(same), None))
        if w is None:
            restored.append(d)
            continue
        if make_copies and str(getattr(w, "artifact_id", "") or "") != own and not _is_quoted_mail_echo(d) \
                and not is_cross_doc_copy(d):
            _mark_copy(d, w, stage)
            new_copies.append(d)
            continue
        if isinstance(getattr(d, "value", None), dict):
            d.value[SURVIVOR_KEY] = {
                "atom_id": str(getattr(w, "id", "") or ""),
                "artifact_id": str(getattr(w, "artifact_id", "") or ""),
                "stage": stage,
            }
    if not restored:
        return list(after), new_copies, []
    # Put each restored atom back after the nearest kept atom preceding it.
    kept_ids = {id(a) for a in after}
    restored_ids = {id(a) for a in restored}
    follow: dict[int, list[Any]] = {}
    lead: list[Any] = []
    anchor = None
    for atom in before:
        if id(atom) in kept_ids:
            anchor = atom
        elif id(atom) in restored_ids:
            (lead if anchor is None else follow.setdefault(id(anchor), [])).append(atom)
    out: list[Any] = list(lead)
    for atom in after:
        out.append(atom)
        out.extend(follow.get(id(atom), ()))
    return out, new_copies, restored


#: ``value`` key on a survivor: documents whose copy of the line was only a
#: QUOTE of the survivor's message (a reply quoting it). Not their line.
QUOTED_IN_KEY = "quoted_in_documents"


def _is_quoted_mail_echo(atom: Any) -> bool:
    """A quoted line of an email message (not of a HubSpot note)."""
    v = _value(atom)
    if not v.get("quoted"):
        return False
    if "hubspot_note_parser" in (getattr(atom, "review_flags", None) or []):
        return False
    if v.get("email_thread") or v.get("message_index") is not None or str(v.get("kind") or "").startswith("email_"):
        return True
    refs = getattr(atom, "source_refs", None) or []
    t = getattr(refs[0], "artifact_type", None) if refs else None
    return str(getattr(t, "value", t) or "") in {"email", "msg", "mbox"}


def _note_quoted_in(survivor: Any, echo: Any) -> None:
    aid = str(getattr(echo, "artifact_id", "") or "")
    if not aid or aid == str(getattr(survivor, "artifact_id", "") or ""):
        return
    if not isinstance(getattr(survivor, "value", None), dict):
        return
    docs = list(survivor.value.get(QUOTED_IN_KEY) or [])
    if aid not in docs:
        docs.append(aid)
        survivor.value[QUOTED_IN_KEY] = docs


def quoted_in(atom: Any) -> set[str]:
    """Documents that only QUOTE this atom's line (see :data:`QUOTED_IN_KEY`)."""
    return {str(x) for x in (_value(atom).get(QUOTED_IN_KEY) or [])}


def resolve_canonical(copies: list[Any], final_atoms: list[Any]) -> None:
    """Point every copy at a canonical atom that is still in the result.

    A canonical atom can itself be folded by a later stage (a note line that
    owned a pasted copy, then retyped away by semantic dedup); its own copy
    entry, or a kept atom of another document with the same words, takes its
    place. A copy whose canonical cannot be found says so rather than pointing
    at nothing.
    """
    final_ids = {str(getattr(a, "id", "") or ""): a for a in final_atoms}
    copy_by_id = {str(getattr(c, "id", "") or ""): c for c in copies}
    by_text: dict[str, list[Any]] = {}
    for a in final_atoms:
        by_text.setdefault(_text_key(a), []).append(a)
    for c in copies:
        v = _value(c)
        ref = v.get("duplicate_of") if isinstance(v.get("duplicate_of"), dict) else None
        if not ref:
            continue
        target = str(ref.get("atom_id") or "")
        hops = 0
        while target and target not in final_ids and target in copy_by_id and hops < 16:
            nxt = _value(copy_by_id[target]).get("duplicate_of") or {}
            target = str(nxt.get("atom_id") or "")
            hops += 1
        if target in final_ids:
            if target != ref.get("atom_id"):
                w = final_ids[target]
                v["duplicate_of"] = {**ref, "atom_id": target,
                                     "artifact_id": str(getattr(w, "artifact_id", "") or "")}
            continue
        own = str(getattr(c, "artifact_id", "") or "")
        alt = next((a for a in by_text.get(_text_key(c), ())
                    if str(getattr(a, "artifact_id", "") or "") != own), None)
        if alt is not None:
            v["duplicate_of"] = {**ref, "atom_id": str(getattr(alt, "id", "") or ""),
                                 "artifact_id": str(getattr(alt, "artifact_id", "") or "")}
        else:
            v["duplicate_of"] = {**ref, "canonical_missing": True}


#: Atoms a stage BUILT from several documents (a conflict between two of them,
#: a signal derived from another atom, the declared-scope question): their refs
#: name the documents they compare, not lines those documents hold.
_SYNTHESIZED_FLAGS = frozenset({"cross_document_conflict", "derived_signal", "declared_scope"})


def _synthesized(atom: Any) -> bool:
    if _SYNTHESIZED_FLAGS & set(getattr(atom, "review_flags", None) or []):
        return True
    v = _value(atom)
    return bool(v.get("derived_from") or v.get("backfill")
                or str(v.get("kind") or "") == "cross_document_conflict")


def _type_of(atom: Any) -> str:
    t = getattr(atom, "atom_type", None)
    return str(getattr(t, "value", t) or "")


def line_key(atom: Any) -> tuple[str, str]:
    """``(atom type, folded words)``: what :func:`holds_own_line` compares."""
    return (_type_of(atom), _text_key(atom))


def holds_own_line(atom: Any, own_keys: Iterable[tuple[str, str]]) -> bool:
    """Whether a document whose atoms have ``own_keys`` (see :func:`line_key`)
    holds ``atom``'s line as an atom of the same type."""
    t, words = line_key(atom)
    return any(ot == t and _jaccard_same(words, ow) for ot, ow in own_keys)


def _jaccard_same(a: str, b: str) -> bool:
    """Same line, worded within a token or two. No containment: a list line
    ("Supported Locations: Delphos, OH, Plymouth, MI") CONTAINS its pieces and
    is not a copy of any one of them."""
    if not a or not b:
        return False
    if a == b:
        return True
    wa, wb = set(a.split()), set(b.split())
    return bool(wa and wb) and len(wa & wb) / len(wa | wb) >= 0.6


def ensure_own_copies(
    kept: list[Any],
    copies: list[Any],
    *,
    stage: str = "own_copy_sweep",
    doc_lines: Callable[[str], list[str] | None] | None = None,
    dropped: Iterable[Any] = (),
) -> list[Any]:
    """A copy for every document a kept atom cites but does not belong to.

    Every fold that merges a loser's source refs onto a survivor in another
    document credits that document's line to the survivor. :func:`split_copies`
    turns the folds it can see into copies; this sweep, run once on the final
    atoms, catches every other path -- a table row or cell folded by a
    type-specific pass, a list-split city merged into another document's site,
    a quoted note bullet folded onto the email it quotes -- so each document
    shows its own atom for each line it holds (live 000132: 35 v2-SOW lines
    and 18 v2-Deal-Kit rows were v1's atoms; six note bullets were an email's).

    A document that already holds the atom (its own atom of the same type with
    the same words, or a copy of this survivor) is left alone. The type
    matters: a Deal Kit row is a ``raw_table_row`` AND a ``bom_line``, and
    when only the ``bom_line`` folded onto v1's, v2 still needs its own. Returns the new copies; the
    survivor lists the documents in ``value["also_in_documents"]``.

    A ref is not the line. With ``doc_lines`` (``artifact_id -> the
    document's own source lines``) a copy is made only where that document's
    own text -- its source, or one of its own atoms (``kept`` / ``dropped``)
    -- holds the survivor's line, outside its header and author metadata.
    A document that merely names the same person is cited, not copied
    (010353: the SOW's contacts row cloned into two intake files that list
    the same people their own way; 010003: a vendor quote's "<name> |
    <phone> | <email>" line cloned into a SOW and four emails). The survivor
    records such documents in ``value["cited_not_held_in"]`` so no
    document listing credits it to them.
    """
    held_by_doc: dict[tuple[str, str], list[str]] = {}
    copy_of: set[tuple[str, str]] = set()
    for a in list(kept) + list(copies):
        aid = str(getattr(a, "artifact_id", "") or "")
        held_by_doc.setdefault((aid, _type_of(a)), []).append(_text_key(a))
        dup = _value(a).get("duplicate_of")
        if isinstance(dup, dict) and dup.get("atom_id"):
            copy_of.add((aid, str(dup["atom_id"])))
    taken = {str(getattr(a, "id", "") or "") for a in list(kept) + list(copies)}
    own_texts: dict[str, list[str]] = {}
    if doc_lines is not None:
        for a in list(kept) + list(dropped):
            if is_cross_doc_copy(a) or _synthesized(a) or _is_metadata_atom(a):
                continue
            k = _text_key(a)
            if k:
                own_texts.setdefault(str(getattr(a, "artifact_id", "") or ""), []).append(k)
    dropped_by_doc: dict[str, list[Any]] = {}
    for d in dropped:
        if not is_cross_doc_copy(d) and not _synthesized(d):
            dropped_by_doc.setdefault(str(getattr(d, "artifact_id", "") or ""), []).append(d)
    restored_ids: set[int] = set()
    out: list[Any] = []
    for w in kept:
        if is_cross_doc_copy(w) or _synthesized(w):
            continue
        if "admission_regex" in (getattr(w, "review_flags", None) or []):
            continue
        own = str(getattr(w, "artifact_id", "") or "")
        wid = str(getattr(w, "id", "") or "")
        words = _text_key(w)
        wtype = _type_of(w)
        if not own or not wid or not words:
            continue
        by_doc: dict[str, list[Any]] = {}
        for r in getattr(w, "source_refs", None) or []:
            aid = str(getattr(r, "artifact_id", "") or "")
            if aid and aid != own:
                by_doc.setdefault(aid, []).append(r)
        quoting = quoted_in(w)
        for aid, refs in by_doc.items():
            if aid in quoting or (aid, wid) in copy_of:
                continue
            if any(_jaccard_same(words, k) for k in held_by_doc.get((aid, wtype), ())):
                continue
            if doc_lines is not None and not document_holds_line(
                words, doc_lines(aid), own_texts.get(aid, ())
            ):  # unreadable (None) is not evidence either
                _note_cited_not_held(w, aid)
                # The document's OWN atom that was folded into ``w`` (its ref
                # is what ``w`` cites) comes back as its copy, in its own
                # words: the intake's own "csm: <name>" line, not the
                # SOW's contacts row.
                ref_keys = {_ref_key(r) for r in refs}
                for d in dropped_by_doc.get(aid, ()):
                    if id(d) in restored_ids or not ({_ref_key(r) for r in d.source_refs or []} & ref_keys):
                        continue
                    if _is_metadata_atom(d) or is_metadata_line(_value(d).get("context")):
                        continue
                    dv = _value(d)
                    dv = dict(dv) if isinstance(getattr(d, "value", None), dict) else {}
                    dv.pop("_suppression", None)
                    dv["duplicate_of"] = {"atom_id": wid, "artifact_id": own, "stage": stage}
                    d.value = dv
                    d.review_flags = [
                        f for f in (getattr(d, "review_flags", None) or [])
                        if not str(f).startswith("suppressed:")
                    ] + [COPY_FLAG]
                    restored_ids.add(id(d))
                    out.append(d)
                    copy_of.add((aid, wid))
                    if isinstance(getattr(w, "value", None), dict):
                        docs = list(w.value.get("also_in_documents") or [])
                        if aid not in docs:
                            w.value["also_in_documents"] = docs + [aid]
                        cn = [x for x in w.value.get(CITED_NOT_HELD_KEY) or [] if x != aid]
                        if cn:
                            w.value[CITED_NOT_HELD_KEY] = cn
                        else:
                            w.value.pop(CITED_NOT_HELD_KEY, None)
                    break
                continue
            from app.core.ids import stable_id

            cid = stable_id("atm_copy", wid, aid)
            if cid in taken:
                continue
            v = {k: val for k, val in _value(w).items() if k not in DOCUMENT_SPECIFIC_KEYS}
            v["duplicate_of"] = {"atom_id": wid, "artifact_id": own, "stage": stage}
            flags = [f for f in (getattr(w, "review_flags", None) or []) if f != COPY_FLAG] + [COPY_FLAG]
            try:
                c = w.model_copy(deep=True)
            except AttributeError:  # not a pydantic model (tests' stand-ins)
                import copy as _copy

                c = _copy.deepcopy(w)
            c.id = cid
            if getattr(c, "atom_id", None):
                c.atom_id = cid
            c.artifact_id = aid
            c.source_refs = [r.model_copy(deep=True) if hasattr(r, "model_copy") else r for r in refs]
            c.receipts = []
            c.value = v
            c.review_flags = flags
            out.append(c)
            taken.add(cid)
            copy_of.add((aid, wid))
            held_by_doc.setdefault((aid, wtype), []).append(words)
            if isinstance(getattr(w, "value", None), dict):
                docs = list(w.value.get("also_in_documents") or [])
                if aid not in docs:
                    docs.append(aid)
                    w.value["also_in_documents"] = docs
    return out


#: ``value`` keys that describe the ORIGINAL's document (its kind, origin,
#: page, section, message, note, speaker) or its cross-document bookkeeping.
#: A copy in another document never inherits them (010003: six copies of a
#: vendor quote's contact line each said ``document_kind: vendor_quote_bom``).
DOCUMENT_SPECIFIC_KEYS = frozenset({
    "also_in_documents", "cited_not_held_in", QUOTED_IN_KEY,
    "document_kind", "doc_origin", "document_date", "page", "page_number",
    "locator", "section_path", "section", "lead_in", "block_id", "block_index",
    "block_kind", "row_index", "line", "line_start", "line_end",
    "email_thread", "message_index", "message_id", "quoted", "hubspot_note_id",
    "note_id", "note_date", "said_by", "authored_at", "context",
})

#: ``value`` key on a survivor: documents its refs cite whose own text does
#: not hold its line (they name the same person, nothing more).
CITED_NOT_HELD_KEY = "cited_not_held_in"

#: A header / author-metadata line: who wrote or sent a document is not a line
#: of its content (010353: a note's "Author: <name>" header line, and a body that never
#: names her).
_METADATA_LINE_RE = re.compile(
    r"^\s*(?:"
    r"(?:from|to|cc|bcc|sent|date|subject|reply-to|message-id|importance|"
    r"author|author[- ]email|author[- ]affiliation|created by|owner|"
    r"hubspot note(?: id)?|note id)\s*:"
    r"|note_id\s*="
    r")"
    r"|(?:^|\|)\s*(?:author|author_email|note_id|date)\s*=",
    re.I,
)


def is_metadata_line(line: Any) -> bool:
    return bool(_METADATA_LINE_RE.search(str(line or "")))


def _is_metadata_atom(atom: Any) -> bool:
    v = _value(atom)
    kind = str(v.get("kind") or "")
    if kind.endswith("_meta") or kind.endswith("_header"):
        return True
    return is_metadata_line(getattr(atom, "raw_text", "") or "")


def _fold_line(text: Any) -> str:
    s = str(text or "")
    if "&" in s:
        try:
            from app.core.textio import decode_html_entities

            s = decode_html_entities(s)
        except Exception:
            pass
    return re.sub(r"[^a-z0-9]+", " ", strip_list_marker(s).lower()).strip()


def document_holds_line(words: str, lines: list[str] | None, own_texts: Iterable[str] = ()) -> bool | None:
    """Does a document's own text hold the line whose folded words are ``words``?

    The line itself, normalised, must appear: in the document's source
    ``lines`` outside header / author metadata (read end to end, so a wrapped
    line still matches), or as one of its own atoms (``own_texts``, folded).
    The same name, or the same phone digits in a signature, is not the line
    (010003: a quote's "<name> | <phone> | <email>" line against an email
    signature showing the same name and the phone written with dots). None when the source
    cannot be read and no own atom holds the line.
    """
    if not words:
        return False
    padded = f" {words} "
    if any(padded in f" {t} " for t in own_texts):
        return True
    if lines is None:
        return None
    content = " ".join(c for c in (_fold_line(ln) for ln in lines if not is_metadata_line(ln)) if c)
    return padded in f" {content} "


def _only_in_metadata(words: str, lines: list[str] | None) -> bool:
    """The line appears in the document only on its header / metadata lines."""
    if not words or not lines:
        return False
    padded = f" {words} "
    every = " ".join(c for c in (_fold_line(ln) for ln in lines) if c)
    if padded not in f" {every} ":
        return False
    return not document_holds_line(words, lines)


def _note_cited_not_held(survivor: Any, aid: str) -> None:
    if not isinstance(getattr(survivor, "value", None), dict):
        return
    docs = list(survivor.value.get(CITED_NOT_HELD_KEY) or [])
    if aid not in docs:
        docs.append(aid)
        survivor.value[CITED_NOT_HELD_KEY] = docs


def cited_not_held_in(atom: Any) -> set[str]:
    """Documents an atom cites that do not hold its line (see :data:`CITED_NOT_HELD_KEY`)."""
    return {str(x) for x in (_value(atom).get(CITED_NOT_HELD_KEY) or [])}


_TEXT_SUFFIXES = {".txt", ".md", ".eml", ".html", ".htm", ".json", ".vtt", ".srt", ".csv", ".tsv", ".text"}


def _read_source_lines(path: Path) -> list[str] | None:
    """A document's own text as lines, or None when it cannot be read back."""
    try:
        from app.core.filetype import content_suffix

        suffix = content_suffix(path)
    except Exception:
        suffix = path.suffix.lower()
    try:
        if suffix == ".pdf":
            from app.core.text_coverage import _read_pdf_lines

            got = [ln for _, ln in _read_pdf_lines(path)]
            return got or None
        if suffix == ".docx":
            from app.core.source_replay import _docx_full_text

            body = _docx_full_text(path)
            return body.splitlines() if body else None
        if suffix in {".xlsx", ".xlsm"}:
            from openpyxl import load_workbook

            wb = load_workbook(str(path), read_only=True, data_only=True)
            out: list[str] = []
            try:
                for ws in wb.worksheets:
                    for row in ws.iter_rows(values_only=True):
                        cells = [str(c) for c in row if c not in (None, "")]
                        if cells:
                            out.append(" | ".join(cells))
            finally:
                wb.close()
            return out or None
        if suffix in _TEXT_SUFFIXES:
            from app.core.text_coverage import _read_text

            body = _read_text(path)
            return body.splitlines() if body else None
    except Exception:
        return None
    return None


def source_lines_reader(artifact_paths: Mapping[str, Any]) -> Callable[[str], list[str] | None]:
    """``artifact_id -> its own source lines`` (cached; None when unreadable)."""
    cache: dict[str, list[str] | None] = {}

    def read(aid: str) -> list[str] | None:
        if aid not in cache:
            p = artifact_paths.get(aid)
            cache[aid] = _read_source_lines(Path(p)) if p else None
        return cache[aid]

    return read


def drop_unheld_copies(
    copies: list[Any],
    final_atoms: list[Any],
    doc_lines: Callable[[str], list[str] | None],
) -> tuple[list[Any], list[Any]]:
    """Split the held copies into ``(kept, refused)``: a copy read off its
    document's header or author metadata is refused.

    A stage's fold can drop an atom a document's EXTRACTOR minted from its
    header -- a note's "Author: <name>" read as a person (010353) --
    and :func:`split_copies` then kept it as that document's copy of the
    SOW's record of that person. Its own context is author metadata; the note's body
    never names her. A refused copy is just a folded atom: it goes to the
    suppression ledger, and the canonical atom stops listing the document.
    """
    by_id = {str(getattr(a, "id", "") or ""): a for a in final_atoms}
    by_id.update({str(getattr(c, "id", "") or ""): c for c in copies})
    kept: list[Any] = []
    refused: list[Any] = []
    for c in copies:
        aid = str(getattr(c, "artifact_id", "") or "")
        v = _value(c)
        context = v.get("context")
        meta = bool(context) and is_metadata_line(context)
        # A held copy IS its document's own atom, minted by its parser (a JSON
        # value reads "contacts[1].name: ...", not the file's bytes). It is
        # refused only when what it was read from is the document's header or
        # author metadata.
        if not meta and not _only_in_metadata(_text_key(c), doc_lines(aid)):
            kept.append(c)
            continue
        refused.append(c)
        c.review_flags = [f for f in (getattr(c, "review_flags", None) or []) if f != COPY_FLAG]
        dup = v.get("duplicate_of") if isinstance(v.get("duplicate_of"), dict) else {}
        w = by_id.get(str(dup.get("atom_id") or ""))
        if isinstance(v, dict):
            v.pop("duplicate_of", None)
        if w is not None and isinstance(getattr(w, "value", None), dict):
            docs = [d for d in (w.value.get("also_in_documents") or []) if d != aid]
            if docs:
                w.value["also_in_documents"] = docs
            else:
                w.value.pop("also_in_documents", None)
            if any(str(getattr(r, "artifact_id", "") or "") == aid for r in getattr(w, "source_refs", None) or []):
                _note_cited_not_held(w, aid)
    return kept, refused


__all__ = [
    "CITED_NOT_HELD_KEY",
    "COPY_FLAG",
    "DOCUMENT_SPECIFIC_KEYS",
    "cited_not_held_in",
    "document_holds_line",
    "drop_unheld_copies",
    "ensure_own_copies",
    "is_metadata_line",
    "source_lines_reader",
    "holds_own_line",
    "line_key",
    "quoted_in",
    "QUOTED_IN_KEY",
    "resolve_canonical",
    "doc_key",
    "document_order",
    "is_cross_doc_copy",
    "settle_folds",
    "SURVIVOR_KEY",
    "split_copies",
]
