"""SOW text inside Word content controls (``w:sdt``) is read wherever the
control sits: around a block, around runs inside a paragraph, inside a table
cell, around a table row, or around a cell within a row.

python-docx's views read none of the last four, so a templated SOW lost its
controlled fields ("deliver  to the Phoenix site") and controlled table rows.
"""

from __future__ import annotations

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

from app.parsers.docx_parser import DocxParser

W = nsdecls("w")


def _r(t: str) -> str:
    return f'<w:r><w:t xml:space="preserve">{t}</w:t></w:r>'


def _p(inner: str) -> str:
    return f"<w:p {W}>{inner}</w:p>"


def _sdt(content: str, *, placeholder: bool = False) -> str:
    pr = "<w:sdtPr><w:showingPlcHdr/></w:sdtPr>" if placeholder else "<w:sdtPr/>"
    return f"<w:sdt {W}>{pr}<w:sdtContent>{content}</w:sdtContent></w:sdt>"


def _build(path) -> None:
    doc = Document()
    doc.add_heading("Scope of Work", level=1)
    doc.add_paragraph("The contractor shall install the core switch stack in the MDF.")
    sect = doc.element.body[-1]
    sect.addprevious(parse_xml(_sdt(_p(_r("The contractor shall provide 40 hours of onsite cabling labor.")))))
    sect.addprevious(parse_xml(
        f"<w:p {W}>{_r('The contractor shall deliver ')}"
        f"<w:sdt><w:sdtPr/><w:sdtContent>{_r('twelve Cisco access points')}</w:sdtContent></w:sdt>"
        f"{_r(' to the Phoenix site.')}</w:p>"
    ))
    sect.addprevious(parse_xml(
        f"<w:p {W}>{_r('Customer contact: ')}"
        f"<w:sdt><w:sdtPr><w:showingPlcHdr/></w:sdtPr><w:sdtContent>{_r('Click or tap here to enter text.')}"
        f"</w:sdtContent></w:sdt></w:p>"
    ))
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Item"
    t.cell(0, 1).text = "Description"
    t.cell(1, 0).text = "1"
    tc = t.cell(1, 1)._tc
    tc.remove(tc.findall(qn("w:p"))[0])
    tc.append(parse_xml(_sdt(_p(_r("Install fiber patch panel in rack 3")))))
    tbl = t._tbl
    tbl.append(parse_xml(_sdt(
        f"<w:tr><w:tc>{_p(_r('2'))}</w:tc><w:tc>{_p(_r('Label every cable at both ends'))}</w:tc></w:tr>"
    )))
    tbl.append(parse_xml(
        f"<w:tr {W}><w:tc>{_p(_r('3'))}</w:tc>"
        f"<w:sdt><w:sdtPr/><w:sdtContent><w:tc>{_p(_r('Certify each drop with a Fluke tester'))}</w:tc>"
        f"</w:sdtContent></w:sdt></w:tr>"
    ))
    tbl.append(parse_xml(
        f"<w:tr {W}><w:tc>{_p(_r('4'))}</w:tc><w:tc><w:p>{_r('Remove ')}"
        f"<w:sdt><w:sdtPr/><w:sdtContent>{_r('the abandoned coax runs')}</w:sdtContent></w:sdt>"
        f"</w:p></w:tc></w:tr>"
    ))
    doc.save(path)


def _texts(path) -> list[str]:
    out = DocxParser().parse_artifact("p", "a", path)
    atoms = out if isinstance(out, list) else out.atoms
    return [a.raw_text for a in atoms]


def test_content_controls_are_read_everywhere(tmp_path) -> None:
    path = tmp_path / "SOW.docx"
    _build(path)
    blob = "\n".join(_texts(path))
    # block-level (already worked)
    assert "40 hours of onsite cabling labor" in blob
    # inline, inside a paragraph
    assert "deliver twelve Cisco access points to the Phoenix site" in blob
    # block control inside a table cell
    assert "Install fiber patch panel in rack 3" in blob
    # control around a whole row
    assert "Label every cable at both ends" in blob
    # control around a cell within a row
    assert "Certify each drop with a Fluke tester" in blob
    # inline control inside a table cell
    assert "Remove the abandoned coax runs" in blob


def test_placeholder_text_is_not_content(tmp_path) -> None:
    # A control still showing its placeholder is an unfilled field: never
    # content (scope, chatter), but kept as the open question it is.
    path = tmp_path / "SOW.docx"
    _build(path)
    out = DocxParser().parse_artifact("p", "a", path)
    atoms = out if isinstance(out, list) else out.atoms
    hits = [a for a in atoms if "Click or tap here" in a.raw_text]
    assert hits, [a.raw_text for a in atoms]
    for a in hits:
        assert a.atom_type.value == "open_question", (a.atom_type, a.raw_text)
        assert "unfilled_placeholder" in a.review_flags
        assert a.value["placeholder"] is True
        assert a.value["field_label"] == "Customer contact"
