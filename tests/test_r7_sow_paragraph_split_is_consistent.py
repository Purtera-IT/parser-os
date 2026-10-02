"""One rule for a multi-sentence SOW paragraph or list item, on every path.

Live 000132: the same Project Scope paragraph (a list of site cities, then a
"Customer requires ..." sentence) was two atoms in SOW v1 and one atom in v2.
Not a length threshold: v2's list ended "..., and <City>, MS." and the clause
splitter read "MS." as the honorific "Ms.", so no boundary. Live 010003: a
two-sentence numbered assumption was two atoms in the .docx and one in the
signed PDF, because PDF list items were never sentence-split. Synthetic text,
same shapes.
"""

from __future__ import annotations

from pathlib import Path

import docx
import fitz

from app.parsers.clause_split import sentences, split_clauses
from app.parsers.docx_parser import DocxParser
from app.parsers.orbitbrief_pdf import OrbitBriefPdfParser

SERIF = "/usr/share/fonts/truetype/liberation/LiberationSerif-Regular.ttf"
SERIF_BOLD = "/usr/share/fonts/truetype/liberation/LiberationSerif-Bold.ttf"

LEAD = "Provider will deliver monthly on site network support services for Customer’s offices in "
TAIL = (
    "Customer requires on site technician support for building IT systems, including cabling, "
    "wireless, printers, desk moves, troubleshooting, ticket response, and coordination with site managers."
)
SHORT = LEAD + "Akron, OH, Madison, WI, Lansing, MI, and Biloxi, MS. " + TAIL           # ends on MS.
LONG = LEAD + "Akron, OH, Madison, WI, Lansing, MI, Flint, MI, Biloxi, MS, and Dover, DE. " + TAIL


def test_state_code_ms_ends_a_sentence() -> None:
    assert split_clauses(SHORT) == [SHORT[: SHORT.index(" Customer requires")], TAIL]
    assert split_clauses(LONG) == [LONG[: LONG.index(" Customer requires")], TAIL]
    # The honorific still holds its sentence together.
    assert sentences("Please call Ms. Avery for access to the riser room on Monday.") == [
        "Please call Ms. Avery for access to the riser room on Monday."
    ]


def test_docx_splits_both_versions_the_same(tmp_path: Path) -> None:
    for name, para in (("short", SHORT), ("long", LONG)):
        d = docx.Document()
        d.add_heading("Project Scope", 2)
        d.add_paragraph(para)
        p = tmp_path / f"{name}.docx"
        d.save(p)
        texts = [a.raw_text for a in DocxParser().parse_artifact("p", "a", p)]
        assert TAIL in texts, (name, texts)
        assert para not in texts, name


ITEMS = {
    6: ["Provider will complete the Services during the hours stated in this SOW."],
    7: ["Once a Change Order that needs a revised PO is signed, Provider will not start the added work until",
        "the revised PO arrives. If work was finished before the revised PO arrives Provider may bill for",
        "the Change Order amount."],
}
SENT_1 = ("Once a Change Order that needs a revised PO is signed, Provider will not start the added work "
          "until the revised PO arrives.")
SENT_2 = "If work was finished before the revised PO arrives Provider may bill for the Change Order amount."


def _pdf(tmp: Path) -> Path:
    doc = fitz.open()
    page = doc.new_page(width=612.3, height=790.9)
    page.insert_text((54, 90), "PROJECT ASSUMPTIONS", fontfile=SERIF_BOLD, fontname="sb", fontsize=12)
    top = 113.0
    for n, lines in ITEMS.items():
        page.insert_text((72, top + 8), f"{n}.", fontfile=SERIF, fontname="s", fontsize=10)
        for ln in lines:
            page.insert_text((90, top + 8), ln, fontfile=SERIF, fontname="s", fontsize=10)
            top += 11.5
    out = tmp / "sow.pdf"
    doc.save(str(out))
    return out


def test_pdf_list_item_splits_like_docx(tmp_path: Path) -> None:
    atoms = OrbitBriefPdfParser().parse_artifact("p", "a", _pdf(tmp_path)).atoms
    by_text = {a.raw_text: a for a in atoms}
    assert SENT_1 in by_text and SENT_2 in by_text, sorted(by_text)
    l1, l2 = (by_text[t].source_refs[0].locator for t in (SENT_1, SENT_2))
    assert l1["block_kind"] == l2["block_kind"] == "bullet_list"
    assert l1["bullet_path"] == l2["bullet_path"]
    assert (l1["sentence_index"], l2["sentence_index"]) == (0, 1)
    assert l1["section_path"] == l2["section_path"]
    # A one-sentence item is untouched.
    one = "Provider will complete the Services during the hours stated in this SOW."
    assert one in by_text and "sentence_index" not in by_text[one].source_refs[0].locator

    d = docx.Document()
    d.add_heading("PROJECT ASSUMPTIONS", 1)
    for n, lines in ITEMS.items():
        d.add_paragraph(f"{n}. " + " ".join(lines))
    p = tmp_path / "sow.docx"
    d.save(p)
    texts = [a.raw_text for a in DocxParser().parse_artifact("p", "a", p)]
    assert SENT_2 in texts
