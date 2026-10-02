"""A SOW's "Supported Locations: City, ST, City, ST, ..." line (live 000132 v1).

The line was one prose atom that the substance gate dropped as unreadable
(place names are not dictionary words); the cities came back only as sites
minted from mentions and no atom carried "Supported Locations". The lead is
now its own list_lead_in context atom and each city its own physical_site
atom pointing back at it.
"""
from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from app.core.city_site_list import split_city_list_paragraph
from app.parsers.docx_parser import DocxParser

CITIES = ["Delphos, OH", "Hudson, WI", "Plymouth, MI", "Lima, OH", "Troy, OH", "Wilmington, DE"]


def _sow(path: Path, *, inline: bool) -> Path:
    d = Document()
    d.add_paragraph("The vendor will install network equipment at each customer site listed below.")
    if inline:
        d.add_paragraph("Supported Locations: " + ", ".join(CITIES))
    else:
        para = d.add_paragraph("Supported Locations:")
        for c in CITIES:
            para.add_run().add_break()
            para.add_run(c)
    d.add_paragraph("All work will be performed during normal business hours.")
    d.save(path)
    return path


@pytest.mark.parametrize("inline", [True, False])
def test_locations_lead_is_its_own_atom_and_each_city_a_site(tmp_path: Path, inline: bool) -> None:
    atoms = DocxParser().parse_artifact("p", "a", _sow(tmp_path / "sow_v1.docx", inline=inline))
    leads = [a for a in atoms if a.raw_text == "Supported Locations:"]
    assert len(leads) == 1
    lead = leads[0]
    assert lead.value["kind"] == "list_lead_in"
    sites = [a for a in atoms if getattr(a.atom_type, "value", a.atom_type) == "physical_site"]
    assert [a.raw_text for a in sites] == CITIES
    for a in sites:
        assert a.value["list_label"] == "Supported Locations"
        assert a.value["list_lead_atom_id"] == lead.id
        assert a.source_refs[0].locator["lead_in"] == ["Supported Locations:"]
    assert sites[0].entity_keys == ["site:delphos_oh"]
    # Nothing still carries the whole line as one prose atom.
    assert not [a for a in atoms if "Delphos" in a.raw_text and "Hudson" in a.raw_text]


def test_split_city_list_paragraph_abstains_on_prose() -> None:
    assert split_city_list_paragraph("Note: we met in Dallas, TX, and Austin, TX, last week") is None
    assert split_city_list_paragraph("Location: Dallas, TX") is None
    assert split_city_list_paragraph("Delphos, OH, Hudson, WI") is None
    lead, sites = split_city_list_paragraph("Sites: Lima, OH; Troy, OH")
    assert lead == "Sites:" and [s.slug for s in sites] == ["lima_oh", "troy_oh"]
