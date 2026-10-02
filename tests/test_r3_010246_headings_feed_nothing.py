"""A docx section heading is a reject-able line that feeds nothing.

Deal 010246's SOW (re-run on #268) came back with every section heading as
its own ``deal_metadata`` atom -- 17 on one SOW. That is by design: every
source line must be an atom a labeler can reject. But a heading is structure,
not a statement, so it must stay behind "Show small talk" (chatter,
``rejected_by`` section_heading) and must not reach anything built on
content: "2.1 Site Survey" read as ``quantity:1``, and a heading that names a
survey must not stage the deal "awaiting site survey".

Since r4 (010087) a bare heading is held in the suppressed sidecar rather
than the atom list: it is the section_path of the lines under it. Since r7
(000132) a heading that leads child lines is an atom again
(``block_kind: heading``), so a label can govern its group; the document
title stays in the sidecar. Either way it feeds nothing.
"""
from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from app.core.compiler import compile_project


@pytest.fixture(autouse=True)
def _no_llm(monkeypatch):
    monkeypatch.setenv("SOWSMITH_DISABLE_LLM", "1")


HEADINGS = (
    "Statement of Work",
    "1. Project Overview",
    "2. Scope of Work",
    "2.1 Site Survey Required Before Install",
    "3. Out of Scope",
    "4. Pricing",
)


def _build(path: Path) -> None:
    d = Document()
    d.add_heading(HEADINGS[0], 0)
    d.add_heading(HEADINGS[1], 1)
    d.add_paragraph("PurTera will install 24 cameras across the warehouse for Acme Corp.")
    d.add_heading(HEADINGS[2], 1)
    d.add_heading(HEADINGS[3], 2)
    d.add_paragraph("Technician will mount each camera on the existing poles.")
    d.add_heading(HEADINGS[4], 1)
    d.add_paragraph("Electrical work and conduit installation.")
    d.add_heading(HEADINGS[5], 1)
    d.add_paragraph("Fixed fee of $12,500 for the installation.")
    d.save(path)


def _type(a) -> str:
    return str(getattr(a.atom_type, "value", a.atom_type))


def test_headings_are_chatter_and_feed_nothing(tmp_path: Path) -> None:
    deal = tmp_path / "deal"
    deal.mkdir()
    _build(deal / "SOW.docx")
    r = compile_project(deal, project_id="p", allow_errors=True, use_cache=False)
    by_text = {a.raw_text: a for a in r.atoms}
    # The document title is no group's lead line: it waits in the suppressed
    # sidecar (r4), still labeled chatter. Every heading here leads child
    # lines, so each is an atom of its own (r7), section_path ending at itself.
    held = {a.raw_text: a for a in r.suppressed_atoms}

    heading_ids = set()
    for h in HEADINGS:
        if h == HEADINGS[0]:
            assert h not in by_text, h
            a = held.get(h)
            assert a is not None, (h, sorted(held))
            assert "suppressed:section_heading" in a.review_flags
        else:
            a = by_text.get(h)
            assert a is not None, (h, sorted(by_text))
            loc = a.source_refs[0].locator
            assert loc.get("block_kind") == "heading"
            assert loc.get("section_path", [])[-1] == h
            assert not any(str(f).startswith("suppressed:") for f in a.review_flags)
        assert _type(a) == "deal_metadata"
        assert "chatter" in a.review_flags
        assert a.value.get("chatter") is True and a.value.get("rejected_by") == "section_heading"
        assert a.entity_keys == [], (h, a.entity_keys)
        heading_ids.add(a.id)

    # Content under the headings keeps a content type.
    assert _type(by_text["Electrical work and conduit installation."]) == "exclusion"
    assert _type(by_text["Technician will mount each camera on the existing poles."]) != "deal_metadata"

    # No packet is built on a heading, and no derived state reads one.
    for p in r.packets:
        used = set(p.governing_atom_ids or []) | set(p.supporting_atom_ids or [])
        assert not (used & heading_ids), (p.family, used & heading_ids)
    derived = [a for a in r.atoms if _type(a) == "deal_state"]
    assert not any("survey" in a.raw_text for a in derived), [a.raw_text for a in derived]
