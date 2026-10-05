"""Deal 000132: a site list's checkbox column, and the site-count gap.

1. "Delphos, OH | ☐ Assessment ☐ Configuration…" -- a service-type checkbox
   cell was glued onto the site name. In an xlsx site table checkbox cells
   are their own atoms beside the site's name/address atom; in a docx table
   the row is ONE atom, the site first and then only the ticked options
   ("Delphos, OH: Installation"), the options as structure on it.
2. "Customer documents declare 6 locations; 3 identified" beside four live
   sites: the gap counted site: keys on physical_site atoms, the site list
   counts what build_site_readiness lists. They now agree.
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
from docx import Document

from app.core.declared_scope import declared_scope_questions
from app.core.ids import stable_id
from app.core.schemas import (
    ArtifactType, AtomType, AuthorityClass, EvidenceAtom, ReviewStatus, SourceRef,
)
from app.parsers.checkbox_cells import checkbox_options
from app.parsers.docx_parser import DocxParser
from app.parsers.xlsx_parser import XlsxParser

CB = {
    "Delphos, OH": "☐ Assessment ☐ Configuration ☒ Installation ☐ Decommission",
    "Lima, OH": "☒ Assessment ☐ Configuration ☒ Installation ☐ Decommission",
    "Troy, OH": "☐ Assessment ☒ Configuration ☐ Installation ☐ Decommission",
    "Wilmington, OH": "☒ Assessment ☒ Configuration ☒ Installation ☐ Decommission",
}
ADDR = {
    "Delphos, OH": "123 Main St, Delphos, OH 45833",
    "Lima, OH": "45 Elm St, Lima, OH 45801",
    "Troy, OH": "9 Oak Ave, Troy, OH 45373",
    "Wilmington, OH": "77 Pine Rd, Wilmington, OH 45177",
}


def _atoms(out):
    return out if isinstance(out, list) else out.atoms


def _assert_split(atoms):
    for name, cb in CB.items():
        glued = [a.raw_text for a in atoms if name in a.raw_text and "☐" in a.raw_text]
        assert not glued, glued
        cells = [a for a in atoms if a.raw_text == cb]
        assert cells, (name, [a.raw_text for a in atoms])
        v = cells[0].value
        assert v["kind"] == "checkbox_selection" and v["subject"] == name
        assert v["selected"] == [l for c, l in checkbox_options(cb) if c]
        assert "Installation" in v["selected"] or "Configuration" in v["selected"]


def _row_atoms(atoms):
    return [a for a in atoms if a.atom_type != AtomType.raw_table_row
            and isinstance(a.value, dict) and a.value.get("checkbox_row")]


def _assert_docx_rows(atoms):
    """One atom per docx row: site first, then the ticked options."""
    assert not [a for a in atoms if "☐" in a.raw_text or "☒" in a.raw_text]
    rows = _row_atoms(atoms)
    for name, cb in CB.items():
        mine = [a for a in rows if a.raw_text.startswith(name)]
        assert len(mine) == 1, (name, [a.raw_text for a in rows])
        v = mine[0].value
        ticked = [l for c, l in checkbox_options(cb) if c]
        assert v["selected"] == ticked
        assert v["not_selected"] == [l for c, l in checkbox_options(cb) if not c]
        assert mine[0].raw_text.endswith(": " + ", ".join(ticked))
    return rows


def test_checkbox_cell_reader():
    assert checkbox_options(CB["Delphos, OH"]) == [
        (False, "Assessment"), (False, "Configuration"), (True, "Installation"), (False, "Decommission")]
    assert checkbox_options("Delphos, OH") is None
    assert checkbox_options("■ Install cameras") is None  # a bullet, not a box


def test_docx_site_name_table(tmp_path: Path):
    doc = Document()
    doc.add_heading("Site List", 1)
    t = doc.add_table(rows=1, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Site Location", "Service Type"
    for name, cb in CB.items():
        r = t.add_row().cells
        r[0].text, r[1].text = name, cb
    doc.save(tmp_path / "SOW v2.docx")
    atoms = _atoms(DocxParser().parse_artifact("p", "a", tmp_path / "SOW v2.docx"))
    rows = _assert_docx_rows(atoms)
    assert all(a.atom_type == AtomType.scope_item and a.value["site"] in CB for a in rows)


def test_docx_roster_table(tmp_path: Path):
    doc = Document()
    t = doc.add_table(rows=1, cols=3)
    for i, h in enumerate(["Site Location", "Address", "Service Type"]):
        t.cell(0, i).text = h
    for name, cb in CB.items():
        r = t.add_row().cells
        r[0].text, r[1].text, r[2].text = name, ADDR[name], cb
    doc.save(tmp_path / "SOW v2.docx")
    atoms = _atoms(DocxParser().parse_artifact("p", "a", tmp_path / "SOW v2.docx"))
    _assert_docx_rows(atoms)  # one checkbox row atom per site row, never per cell
    sites = [a for a in atoms if a.atom_type == AtomType.physical_site
             and a.source_refs[0].extraction_method == "docx_site_roster_v1"]
    assert len(sites) == 4 and not [a for a in sites if "Service Type" in a.raw_text]


def test_xlsx_roster_sheet(tmp_path: Path):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Site List"
    ws.append(["Site Location", "Address", "Service Type"])
    for name, cb in CB.items():
        ws.append([name, ADDR[name], cb])
    wb.save(tmp_path / "Deal Kit v2.xlsx")
    atoms = _atoms(XlsxParser().parse(tmp_path / "Deal Kit v2.xlsx"))
    _assert_split(atoms)
    assert len([a for a in atoms if a.atom_type == AtomType.physical_site]) == 4


# ── the site-count gap agrees with the site list ────────────────────────────


def _atom(aid, atype, text, *, keys=(), value=None, authority=AuthorityClass.contractual_scope):
    return EvidenceAtom(
        id=aid, project_id="p", artifact_id="art", atom_type=atype, raw_text=text,
        normalized_text=text.lower(), value=value or {"text": text}, entity_keys=list(keys),
        source_refs=[SourceRef(id=stable_id("src", aid), artifact_id="art",
                               artifact_type=ArtifactType.docx, filename="SOW v2.docx",
                               locator={}, extraction_method="t", parser_version="t")],
        authority_class=authority, confidence=0.9, review_status=ReviewStatus.auto_accepted,
        review_flags=[], parser_version="t",
    )


def test_gap_counts_every_listed_site():
    from app.core.orbitbrief_core import build_site_readiness

    atoms = [_atom("d", AtomType.scope_item, "The project covers 6 locations across Ohio.")]
    for i, name in enumerate(["Lima", "Troy", "Wilmington"]):
        sid = name.upper()
        atoms.append(_atom(f"s{i}", AtomType.physical_site, f"facility: {name}, OH",
                           keys=[f"site:{name.lower()}"],
                           value={"kind": "physical_site", "id": sid, "site_id": sid, "name": name}))
    # A listed site with its id but no site: key (a name-only roster row).
    atoms.append(_atom("s3", AtomType.physical_site, "facility: Delphos, OH",
                       value={"kind": "physical_site", "id": "DELPHOS-OH", "site_id": "DELPHOS-OH",
                              "name": "Delphos, OH"}))
    listed = build_site_readiness(atoms=atoms, edges=[])["site_count"]
    assert listed == 4
    (q,) = [a for a in declared_scope_questions(project_id="p", atoms=atoms)
            if a.value["declared_scope"]["kind"] == "site_count_gap"]
    assert q.value["declared_scope"]["found_count"] == listed
    assert "declare 6 locations; 4 identified" in q.raw_text


def test_no_gap_when_the_list_is_complete():
    atoms = [_atom("d", AtomType.scope_item, "The project covers 2 locations.")]
    for i, name in enumerate(["Lima", "Troy"]):
        atoms.append(_atom(f"s{i}", AtomType.physical_site, f"facility: {name}, OH",
                           value={"kind": "physical_site", "id": name.upper(), "name": name}))
    assert not [a for a in declared_scope_questions(project_id="p", atoms=atoms)
                if a.value["declared_scope"]["kind"] == "site_count_gap"]


# ── every row of a name-only site table is a site ───────────────────────────


def test_name_only_site_table_makes_every_row_a_site(tmp_path: Path):
    # "Troy, OH" and "Lima, OH" became sites; "Delphos, OH" and "Wilmington,
    # OH" did not. The hygiene gate kept a two-token phrase only when both
    # tokens were <=6 characters (the site-code shape), so a longer city name
    # had no "positive site signal" and its site: key was dropped.
    from app.core.compiler import compile_project
    from app.core.orbitbrief_core import build_site_readiness

    doc = Document()
    doc.add_heading("Site List", 1)
    t = doc.add_table(rows=1, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Site Location", "Service Type"
    for name, cb in CB.items():
        r = t.add_row().cells
        r[0].text, r[1].text = name, cb
    doc.save(tmp_path / "SOW v2.docx")
    r = compile_project(tmp_path, project_id="p", allow_errors=True, use_cache=False)
    keys = {k for a in r.atoms for k in a.entity_keys if k.startswith("site:")}
    for name in CB:
        slug = "site:" + name.lower().replace(", ", "_")
        assert slug in keys, (slug, keys)
    assert build_site_readiness(atoms=r.atoms, edges=r.edges)["site_count"] == 4


def test_city_state_hygiene_is_narrow():
    from app.core.site_llm_verify import _is_obvious_non_site

    # Outside the document's own site catalog a bare city/state is still not
    # a site by itself; the gate is unchanged.
    assert _is_obvious_non_site("delphos oh")
