"""No artifact may route to nobody.

When the router returned None the compile recorded `skipped_no_parser`, and the
file then existed in the manifest and nowhere else -- not the envelope, not
PM_HANDOFF, not the Deal Kit. The only way to learn it had arrived was to go
and look at the manifest.

Measured across 461 dev envelopes:

    .dwg    2   AutoCAD drawings           -> dwg_parser
    .json   6   historical                 -> JsonParser claims these at 0.55
    .xls    2   legacy Excel               -> XlsxParser via calamine
                (a .xls that cannot be OPENED still reaches the marker --
                 claimed at routing time only when a reader can read it,
                 so an unreadable workbook stays visible instead of
                 producing zero atoms and no error)
    .rpmsg  2   RMS-encrypted Outlook mail -> unread marker
    .doc    1   legacy Word                -> unread marker
    .gif    1   an image                   -> ImageParser, extension was missing
    .tsv    1   tab-separated values       -> unread marker
    .pptx   6   failed_parse, "python-pptx is required" -- the parser shipped
                and its dependency was never declared

Some of these deserve a real parser and some cannot be read at all: an `.rpmsg`
is encrypted by design. The point is that from outside those two look
identical, and both look like the file was never sent.
"""
from __future__ import annotations

from pathlib import Path

import pytest

from app.parsers.registry import choose_parser
from app.parsers.unread_parser import UnreadParser, describe

#: (filename, first bytes) -> the parser that must take it.
BINARY = bytes([0x8B, 0xD3])
#: The OLE2 signature a pre-2007 Word or Excel file starts with.
OLE = bytes([0xD0, 0xCF, 0x11, 0xE0])

ROUTES = [
    ("plan.dwg", b"AC1032" + b"\x00" * 600, "dwg"),
    ("plan.dxf", b"  0\nSECTION\n" + b"x" * 600, "dwg"),
    ("anim.gif", b"GIF89a" + b"x" * 600, "image"),
    ("photo.png", b"\x89PNG\r\n\x1a\n" + b"x" * 600, "image"),
    # Realistic: a pre-2007 Word or Excel file is OLE2 and full of NULs. The
    # first version of these fixtures was the signature plus 600 `x`, which
    # decodes cleanly as text -- so the parser read it, correctly, and the test
    # was asserting against a file that does not exist in the world.
    ("legacy.doc", OLE + bytes([0]) * 4000, "unread"),
    ("book.xls", OLE + bytes([0]) * 4000, "unread"),
    ("protected.rpmsg", bytes([0]) + BINARY * 300, "unread"),
    ("mystery.xyz", bytes([0]) + BINARY * 900, "unread"),
]


@pytest.mark.parametrize("name,body,expected", ROUTES)
def test_every_artifact_reaches_a_parser(tmp_path: Path, name, body, expected):
    path = tmp_path / name
    path.write_bytes(body)
    parser, match, _ = choose_parser(path)
    assert parser is not None, f"{name} routed to nobody"
    assert parser.capability.parser_name == expected, match.reasons


def test_nothing_routes_to_none(tmp_path: Path):
    """The invariant, stated once: whatever arrives, somebody holds it."""
    for name in ("a.rpmsg", "b.7z", "c", "d.dmg", "e.wpd", "f.pages"):
        path = tmp_path / name
        path.write_bytes(bytes([0]) + BINARY * 100)
        parser, _, _ = choose_parser(path)
        assert parser is not None, f"{name} routed to nobody"


def test_the_marker_says_what_arrived_and_why_it_is_missing(tmp_path: Path):
    """Genuinely binary. A file of 4,211 `x` bytes IS text, and the parser is
    right to read it rather than call it unreadable -- which is what the first
    version of this fixture accidentally proved."""
    path = tmp_path / "Signed SOW.rpmsg"
    path.write_bytes(bytes([0, 1, 2]) + BINARY * 2104)
    atom = UnreadParser().parse(path)[0]
    assert "Signed SOW.rpmsg" in atom.raw_text
    assert "encrypted" in atom.raw_text
    assert "4,211 bytes" in atom.raw_text
    assert atom.atom_type.value == "open_question"
    assert atom.value["suffix"] == ".rpmsg"


def test_it_names_a_fix_where_there_is_one():
    """A PM can act on "ask for a .docx". They cannot act on "unsupported"."""
    assert "docx" in describe(Path("x.doc"))
    assert "xlsx" in describe(Path("x.xls"))
    assert "unprotected copy" in describe(Path("x.rpmsg"))


def test_an_unknown_extension_still_gets_an_honest_sentence():
    assert ".wpd" in describe(Path("x.wpd")) or "wpd" in describe(Path("x.wpd"))
    assert "no extension" in describe(Path("noext"))


def test_the_marker_parser_never_competes(tmp_path: Path):
    """It is reached only where the router had already decided on None, so a
    real parser at any confidence beats it."""
    path = tmp_path / "anything.pdf"
    path.write_bytes(b"%PDF-1.4\n" + b"x" * 600)
    assert UnreadParser().match(path, None, None).confidence == 0.0


def test_a_cad_drawing_is_typed_as_a_picture():
    """It fell through to ArtifactType.txt while dwg_parser declared `image` on
    its source refs, and the marker atom never reached the envelope. The image
    parsers' markers do reach it -- 133 atoms across 31 deals -- and this is the
    difference between them."""
    from app.core.manifest import _artifact_type_for_path
    from app.parsers.registry import _artifact_type_for_suffix

    assert _artifact_type_for_path(Path("plan.dwg")).value == "image"
    assert _artifact_type_for_suffix(".dwg").value == "image"
    assert _artifact_type_for_suffix(".dxf").value == "image"


def test_pptx_has_its_dependency_declared():
    """Six decks across the dev corpus come back failed_parse with
    "python-pptx is required for the PPTX parser". The parser shipped; the
    dependency was never declared."""
    import tomllib

    text = Path("pyproject.toml").read_bytes()
    deps = tomllib.loads(text.decode())["project"]["dependencies"]
    assert any(d.startswith("python-pptx") for d in deps), deps


# --------------------------------------------------------------------------
# ...but "nothing was read" is itself a signal, and must survive
# --------------------------------------------------------------------------

def test_filler_too_small_to_be_a_document_still_routes_to_nobody(tmp_path: Path):
    """The floor, and the reason for it.

    `test_contentless_text_still_reports_no_parser` records a measured
    position: every real deal document is at least 4 non-empty lines and 258
    characters, and saying "I did not read this" is more honest than
    manufacturing atoms out of filler. The first version of UnreadParser read
    anything that decoded as text and swallowed that signal.
    """
    path = tmp_path / "random.txt"
    path.write_text("just filler words with no structured signals", encoding="utf-8")
    parser, _, _ = choose_parser(path)
    assert parser is None


def test_a_substantial_unstructured_file_is_read(tmp_path: Path):
    """The other half. NO parser claims a plain .txt with no structure, so
    above the floor this is the difference between reading the file and losing
    its contents to a warning line."""
    path = tmp_path / "site_notes.txt"
    path.write_text(
        "Walkthrough notes from the Penn Plaza visit on the twelfth floor.\n"
        "The team counted 212 Cat6A drops across one hundred and six workstations.\n"
        "Six forty-eight port patch panels are to be installed in the IT closet.\n"
        "Electrical connections will be provided by the landlord, not by us.\n"
        "The freight elevator is the only route for cable reels and panels.\n",
        encoding="utf-8",
    )
    parser, _, _ = choose_parser(path)
    assert parser is not None
    texts = " ".join(a.raw_text for a in parser.parse(path))
    assert "212 Cat6A drops" in texts


def test_the_legacy_bridge_only_claims_what_it_can_read(tmp_path: Path):
    """`can_read` is asked at ROUTING time, and it is honest about failure.

    calamine routes on the EXTENSION, not the content, so a workbook whose name
    lies about its format fails to open and correctly falls through to the
    marker rather than being mis-parsed in silence.

    There is no binary .xls fixture in this repo on purpose -- the real one is
    1.9MB. The live file is covered by the impact run; what belongs here is the
    boundary: which suffixes the bridge claims, and that an unopenable file is
    not claimed.
    """
    pytest.importorskip("python_calamine")
    from openpyxl import Workbook

    from app.parsers.legacy_spreadsheet import can_read, needs_conversion

    src = tmp_path / "pricing.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["PART #", "MATERIAL", "QTY"])
    ws.append(["NX-100", "STATION CABLE", 250])
    wb.save(src)

    # .xlsx already has a parser; the bridge must not intercept it.
    assert needs_conversion(src) is False
    assert can_read(src) is False

    # A convertible suffix over content that does not match it: not claimed.
    mislabelled = tmp_path / "pricing.ods"
    mislabelled.write_bytes(src.read_bytes())
    assert needs_conversion(mislabelled) is True
    assert can_read(mislabelled) is False, "a lying extension must not be claimed"


def test_an_unreadable_legacy_workbook_stays_visible(tmp_path: Path):
    """The miss must not become silent.

    Claiming every .xls and then failing to open one produces zero atoms and no
    error: the file lands in the manifest and nowhere else. Routing asks whether
    a reader can open it BEFORE claiming, so a corrupt workbook keeps its marker.
    """
    path = tmp_path / "corrupt.xls"
    path.write_bytes(OLE + bytes([0]) * 4000)
    parser, match, _ = choose_parser(path)
    assert parser.capability.parser_name == "unread", match.reasons
    atom = parser.parse(path)[0]
    assert "xlsx" in atom.raw_text, "a PM must be told what to ask for"


def test_a_word_template_is_a_word_document(tmp_path: Path):
    """A change order does not stop being one because it was saved as a template.

    `.dotx` is OOXML identical to `.docx` but for one string in
    [Content_Types].xml -- `wordprocessingml.template.main+xml` where a document
    says `document.main+xml`. python-docx refuses on that string alone, so both
    of the corpus's .dotx files produced nothing. Both are CHANGE ORDERS:

        010195- TV Install Change Order 8.20 v1.dotx
        00051- Merrill Gardens (Change Order).dotx
    """
    import docx

    from app.parsers.word_template import is_word_template, to_docx

    src = tmp_path / "change_order.docx"
    d = docx.Document()
    d.add_paragraph("Effective Date: 8.20")
    d.add_paragraph("Requesting Party: CDW")
    d.save(src)

    # Re-declare it a template, exactly as Word does when you Save As .dotx.
    import zipfile
    tpl = tmp_path / "change_order.dotx"
    with zipfile.ZipFile(src) as s, zipfile.ZipFile(tpl, "w") as out:
        for item in s.infolist():
            data = s.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"document.main+xml", b"template.main+xml")
            out.writestr(item, data)

    assert is_word_template(tpl)
    with pytest.raises(ValueError):
        docx.Document(str(tpl))          # the defect, pinned

    rewritten = to_docx(tpl)
    assert rewritten is not None
    body = "\n".join(p.text for p in docx.Document(str(rewritten)).paragraphs)
    assert "Requesting Party: CDW" in body

    parser, match, _ = choose_parser(tpl)
    assert parser.capability.parser_name == "docx", match.reasons


def test_a_broken_template_is_not_claimed(tmp_path: Path):
    """Same rule as the legacy workbook: an unreadable file keeps its marker."""
    from app.parsers.word_template import can_read

    bad = tmp_path / "corrupt.dotx"
    bad.write_bytes(b"PK\x03\x04" + bytes([0]) * 2000)
    assert can_read(bad) is False
    parser, match, _ = choose_parser(bad)
    assert parser.capability.parser_name == "unread", match.reasons
