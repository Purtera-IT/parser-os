"""A schema-typed table row keeps its source row's place in the document.

    Fees
    The work is billed as below.
    | Item          | Unit Rate | Quantity | Subtotal |
    | Widget setup  | $50.00    | 4        | $200.00  |
    | Gadget repair | $75.00    | 2        | $150.00  |
    Totals are estimates.
    [footer] Internal use only.

The raw table rows are re-read as typed rows (bom_line here) by
entity_extraction._enrich_table_atoms and appended after every parser atom.
Their locator carried the table and row numbers but not the row's place
(block_index, line_start, page). A reading-order sort over the locator then
fell back to the bare row number and put the whole table after the page
header and footer, at the end of the document. The typed row now copies the
place of the raw row it was read out of.
"""
from __future__ import annotations

import math
from pathlib import Path

from docx import Document

from app.core.entity_extraction import _enrich_table_atoms
from app.parsers.docx_parser import DocxParser

ROWS = [
    ["Item", "Unit Rate", "Quantity", "Subtotal"],
    ["Widget setup", "$50.00", "4", "$200.00"],
    ["Gadget repair", "$75.00", "2", "$150.00"],
]


def _docx(tmp_path: Path) -> Path:
    doc = Document()
    doc.sections[0].footer.paragraphs[0].text = "Internal use only."
    doc.add_heading("Fees", level=1)
    doc.add_paragraph("The work is billed as below.")
    t = doc.add_table(rows=len(ROWS), cols=len(ROWS[0]))
    for r, row in enumerate(ROWS):
        for c, text in enumerate(row):
            t.cell(r, c).text = text
    doc.add_paragraph("Totals are estimates.")
    p = tmp_path / "fees.docx"
    doc.save(p)
    return p


def _loc(a):
    return a.source_refs[0].locator


def _type(a):
    return getattr(a.atom_type, "value", a.atom_type)


def _num(loc, key, fallback):
    # Mirrors the labeling walk's num(): null reads 0, a missing key falls back.
    if key not in loc:
        return fallback
    v = loc[key]
    if v is None:
        return 0.0
    try:
        n = float(v)
    except (TypeError, ValueError):
        return fallback
    return n if math.isfinite(n) else fallback


def _reading_key(a, i):
    loc = _loc(a)
    return (
        _num(loc, "page", _num(loc, "row", 0)),
        _num(loc, "block_index", _num(loc, "row_index", _num(loc, "row", 0))),
        _num(loc, "line_start", 0),
        i,
    )


def _atoms(tmp_path):
    atoms = DocxParser().parse_artifact_full(project_id="p", artifact_id="a", path=_docx(tmp_path)).atoms
    typed = _enrich_table_atoms(atoms, project_id="p")
    return atoms, typed


def test_typed_row_copies_the_raw_rows_place(tmp_path):
    atoms, typed = _atoms(tmp_path)
    raw = {(_loc(a).get("table_index"), _loc(a).get("row")): a for a in atoms if _type(a) == "raw_table_row"}
    boms = [a for a in typed if _type(a) == "bom_line"]
    assert len(boms) == 2
    for b in boms:
        src = _loc(raw[(_loc(b).get("table_index"), _loc(b).get("row"))])
        assert "block_index" in src
        for key in ("page", "block_index", "line_start", "line_end"):
            if key in src:
                assert _loc(b)[key] == src[key], key


def test_typed_rows_read_inside_the_table_not_after_the_footer(tmp_path):
    atoms, typed = _atoms(tmp_path)
    allatoms = [a for a in atoms if _type(a) != "raw_table_row"] + typed
    ordered = [a for _, a in sorted(((_reading_key(a, i), a) for i, a in enumerate(allatoms)), key=lambda x: x[0])]
    texts = [a.raw_text for a in ordered]

    bom_pos = [i for i, a in enumerate(ordered) if _type(a) == "bom_line"]
    assert bom_pos, texts
    after = next(i for i, t in enumerate(texts) if t.startswith("Totals are estimates"))
    footer = next(i for i, t in enumerate(texts) if t.startswith("Internal use only"))
    lead = next(i for i, t in enumerate(texts) if t.startswith("The work is billed"))
    assert lead < min(bom_pos)
    assert max(bom_pos) < after
    assert max(bom_pos) < footer
    # Row order inside the table is kept.
    assert [ordered[i].raw_text.split(" |")[0] for i in bom_pos] == ["Widget setup", "Gadget repair"]


def test_a_paged_row_hands_its_page_on(tmp_path):
    atoms, _ = _atoms(tmp_path)
    raw = next(a for a in atoms if _type(a) == "raw_table_row")
    _loc(raw)["page"] = 7
    typed = _enrich_table_atoms([raw], project_id="p")
    assert typed and all(_loc(b)["page"] == 7 for b in typed)
