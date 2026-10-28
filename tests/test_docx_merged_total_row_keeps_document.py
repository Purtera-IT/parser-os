"""A total row merged across three of four columns keeps the document (010245, 010355).

"TOTAL | (merged) | (merged) | $191,000.83" reads by grid column as
["TOTAL", "", "", "$191,000.83"]: four cells, so the label/value form-row
rewrite paired it as (TOTAL, '') and ('', $191,000.83), dropped both, and the
empty row atom failed validation -- the whole document was lost with it
(010245 SOW v3/v4, 010355's three revised ROMs).
"""

from __future__ import annotations

from docx import Document

from app.parsers.docx_parser import DocxParser

HEADER = ["Item / Description", "Qty", "Unit Price", "Line Total (USD)"]
BODY = ["Cat6 drop, terminate and test", "12", "$150.00", "$1,800.00"]


def _doc(path, *, form_row=False):
    d = Document()
    d.add_paragraph("Rough order of magnitude")
    t = d.add_table(rows=3, cols=4)
    for i, v in enumerate(HEADER):
        t.rows[0].cells[i].text = v
    for i, v in enumerate(BODY):
        t.rows[1].cells[i].text = v
    t.cell(2, 0).merge(t.cell(2, 2)).text = "ESTIMATED TOTAL:"
    t.cell(2, 3).text = "$1,800.00"
    d.add_paragraph("Customer provides lift access to the ceiling space.")
    if form_row:
        f = d.add_table(rows=2, cols=4)
        for i, v in enumerate(["Contact", "Contact", "Contact", "Contact"]):
            f.rows[0].cells[i].text = v
        for i, v in enumerate(["Tech Contact", "Pat 555-0100", "Escalation Contact", ""]):
            f.rows[1].cells[i].text = v
    d.save(path)
    return DocxParser().parse_artifact("p", "a", path)


def test_the_document_parses_and_the_total_is_one_line(tmp_path):
    atoms = _doc(tmp_path / "rom.docx")
    texts = [a.raw_text for a in atoms]
    assert "ESTIMATED TOTAL: | $1,800.00" in texts
    assert any("lift access" in t for t in texts)  # the rest of the document survives
    assert all(a.raw_text.strip() and a.normalized_text.strip() for a in atoms)


def test_a_real_form_row_still_drops_its_orphan_label(tmp_path):
    atoms = _doc(tmp_path / "rom.docx", form_row=True)
    texts = [a.raw_text for a in atoms]
    assert "Tech Contact | Pat 555-0100" in texts
    assert not any("Escalation Contact" in t for t in texts if "Pat" in t)
