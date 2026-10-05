"""Fields typed as soft-broken lines of one paragraph are one atom per field.

Live 000132 (structure only; the text here is synthetic): a "Normal (Web)"
paragraph held "Label: value" <w:br/> "Label: value" (and, in a second
paragraph, three such lines). paragraph.text reads each w:br as "\\n". The
clause split kept a paragraph under 160 characters whole, and merged a field
shorter than 25 characters into the next one, so both paragraphs came out as
one atom per paragraph ("Term: ... Support Window: ...").
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from app.parsers.clause_split import split_clauses
from app.parsers.docx_parser import DocxParser

W = nsdecls("w")
BR = "<w:r><w:br/></w:r>"


def _r(text: str) -> str:
    return f'<w:r><w:t xml:space="preserve">{text}</w:t></w:r>'


def _para(*lines: tuple[str, ...]) -> str:
    """One NormalWeb paragraph; each line is a label run then value runs, and
    lines are separated by a run holding only <w:br/>."""
    body = BR.join("".join(_r(t) for t in line) for line in lines)
    return (f'<w:p {W}><w:pPr><w:pStyle w:val="NormalWeb"/><w:spacing w:line="240" '
            f'w:lineRule="auto"/><w:contextualSpacing/></w:pPr>{body}</w:p>')


def _texts(tmp_path: Path, *paras: str) -> list[str]:
    d = Document()
    d.add_heading("Services Proposal", 1)
    d.add_heading("Project Scope", 2)
    d.add_paragraph("Provider will deliver monthly support visits to every listed office.")
    body = d.element.body
    for p in paras:
        body.insert(len(body) - 1, parse_xml(p))  # before w:sectPr
    d.add_paragraph("Services are billed monthly in arrears.")
    path = tmp_path / "SOW.docx"
    d.save(path)
    out = DocxParser().parse_artifact("p", "a", path)
    atoms = out if isinstance(out, list) else out.atoms
    return [a.raw_text for a in atoms]


def test_two_short_fields_are_two_atoms(tmp_path):
    texts = _texts(tmp_path, _para(
        ("Rate Basis:", " Hourly"),
        ("Minimum Charge:", " 2 hours", " per call, 6 for scheduled ", "full", " days"),
    ))
    assert "Rate Basis: Hourly" in texts
    assert "Minimum Charge: 2 hours per call, 6 for scheduled full days" in texts
    assert not any("Rate Basis" in t and "Minimum Charge" in t for t in texts)


def test_a_field_shorter_than_a_sentence_is_not_merged_into_the_next(tmp_path):
    texts = _texts(tmp_path, _para(
        ("Duration:", " ", "6", "-months "),
        ("Service Hours:", " Normal office hours unless agreed otherwise in writing"),
        ("Escalation:", " The engineer escalates to the named site lead during the visit if needed"),
    ))
    assert texts[-4:-1] == [
        "Duration: 6-months",
        "Service Hours: Normal office hours unless agreed otherwise in writing",
        "Escalation: The engineer escalates to the named site lead during the visit if needed",
    ]


def test_split_clauses_field_lines():
    assert split_clauses("Rate Basis: Hourly\nMinimum Charge: 2 hours", hard_breaks=True) == [
        "Rate Basis: Hourly", "Minimum Charge: 2 hours"]
    # Without a w:br the same words stay one unit (no break to split on).
    assert split_clauses("Rate Basis: Hourly Minimum Charge: 2 hours", hard_breaks=True) == []
    # A line that does not open with a label keeps the old behaviour.
    assert split_clauses("Rate Basis: Hourly\nand 2 hours minimum", hard_breaks=True) == []
    # Not a docx soft break: unchanged.
    assert split_clauses("Rate Basis: Hourly\nMinimum Charge: 2 hours") == []
