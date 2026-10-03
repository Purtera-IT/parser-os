"""Every row of a headed xlsx table reads in one keyed shape.

Found on deal 000132's Deal Kit Summary tab: a small two-column table
(bold header row, then rows with both cells filled and rows with only the
first) came out in two shapes. The generic row emitter bound a lone cell to
its column ("Header: value") but left a fully-filled row bare
("value | value"). Both now read "Header: value | Header: value".
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.styles import Font

from app.parsers.xlsx_parser import XlsxParser


def _workbook(path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    grid = [
        ("Site", "Drive"),
        ("Springfield, IL", "15 miles"),
        ("Columbus, GA", "40 miles"),
        ("Dayton, OH", None),
        ("Akron, OH", None),
    ]
    for r, (a, b) in enumerate(grid, 1):
        ws.cell(r, 1, a)
        if b is not None:
            ws.cell(r, 2, b)
    ws["A1"].font = Font(bold=True)
    ws["B1"].font = Font(bold=True)
    ws["E5"].font = Font(italic=True)  # styled extent past the data, all empty
    wb.save(path)
    return path


def _rows(path: Path) -> dict[int, set[str]]:
    out: dict[int, set[str]] = {}
    for a in XlsxParser().parse(path):
        row = (a.source_refs[0].locator or {}).get("row") if a.source_refs else None
        out.setdefault(row, set()).add(a.raw_text)
    return out


def test_full_and_lone_cell_rows_share_the_keyed_shape(tmp_path: Path):
    rows = _rows(_workbook(tmp_path / "Deal Kit.xlsx"))
    assert rows[2] == {"Site: Springfield, IL | Drive: 15 miles"}
    assert rows[3] == {"Site: Columbus, GA | Drive: 40 miles"}
    assert rows[4] == {"Site: Dayton, OH"}
    assert rows[5] == {"Site: Akron, OH"}
    assert 1 not in rows  # the header row is no atom

