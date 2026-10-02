"""A rate sheet is rows a person labels, not a count line.

Deal Kit rate sheets ("COST RATES", "SELLL RATES", a per-country matrix,
"Materials") used to collapse into one ``"<sheet>: N pricing lines, $x-$y"``
atom; on one deal ~850 rate rows left the labeller's view that way. Every
priced row is now its own atom (role, rate, unit, country when present), and
the rollup stays beside them only as the summary the pricing packet reads.
"""

from __future__ import annotations

from openpyxl import Workbook

from app.core.schemas import (
    ArtifactType,
    AtomType,
    AuthorityClass,
    EvidenceAtom,
    ReviewStatus,
    SourceRef,
)
from app.core.table_rollup import roll_up_table_rows
from app.parsers.xlsx_parser import XlsxParser

_ROLES = ["Networking L1 Technician", "Networking L2 Technician", "EUC L1 Technician", "Project Manager"]
_MINS = ["2 hr. min", "4hr. Min.", "8 hr. min"]
_COUNTRIES = ["United States", "Indonesia", "United Arab Emirates", "Hong Kong", "Germany", "Brazil"]
_RATE_ROWS = [
    ("PS-L1-ENG-LABOR-ONSITE", "L1 Engineer", "Per Hour", 85),
    ("PS-L2-ENG-LABOR-ONSITE", "L2 Engineer", "Per Hour", 105),
    ("PS-PROJMGMT-REMOTE", "Project Manager", "Per Hour", 125),
    ("PS-TRAVEL-EXPENSE", "Travel", "Per Day", 165),
]


def _deal_kit(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Country Rates"
    ws.append(["Country", "Request"] + [f"{r} {m}" for r in _ROLES for m in _MINS])
    for ci, c in enumerate(_COUNTRIES):
        ws.append([c, 1.0 + ci / 10] + [60 + ci * 3 + j for j in range(len(_ROLES) * len(_MINS))])
    for name in ("COST RATES", "SELLL RATES"):
        w = wb.create_sheet(name)
        w.append(["Code", "Role", "Unit", "Rate"])
        for row in _RATE_ROWS:
            w.append(list(row))
    m = wb.create_sheet("Materials")
    m.append(["ID #", "Material Description", "OEM", "Order QTY", "USA Cost $"])
    for i in range(45):
        m.append([i + 1, f"CAT6 patch cord type {i}", "CommScope", None, 20 + i])
    h = wb.create_sheet("Helper - Do not Edit")
    for x in ("L0", "L1", "L2"):
        h.append([x])
    wb.save(path)


def _atoms(path):
    out = XlsxParser().parse_artifact("proj", "art", path)
    return out if isinstance(out, list) else out.atoms


def _sheet(atoms, name):
    return [a for a in atoms if (a.value or {}).get("sheet_name") == name]


def test_every_rate_row_is_its_own_atom(tmp_path) -> None:
    path = tmp_path / "Deal_Kit.xlsx"
    _deal_kit(path)
    atoms = _atoms(path)

    for name in ("COST RATES", "SELLL RATES"):
        rows = [a for a in _sheet(atoms, name) if not a.value.get("is_summary")]
        assert len(rows) == len(_RATE_ROWS), name
        assert all(a.atom_type == AtomType.service_line for a in rows)
        by_role = {a.value["role"]: a.value for a in rows}
        assert by_role["L1 Engineer"]["rate"] == 85
        assert by_role["L1 Engineer"]["unit"] == "Per Hour"
        assert by_role["Travel"]["unit"] == "Per Day"
        # COST and SELL rows are identical cells, but different facts.
        assert all(a.raw_text.startswith(name) for a in rows)
    cost_ids = {a.id for a in _sheet(atoms, "COST RATES")}
    assert not cost_ids & {a.id for a in _sheet(atoms, "SELLL RATES")}

    matrix = [a for a in _sheet(atoms, "Country Rates") if not a.value.get("is_summary")]
    assert {a.value["country"] for a in matrix} == set(_COUNTRIES)
    us = next(a for a in matrix if a.value["country"] == "United States")
    rates = us.value["rates"]
    assert len(rates) == len(_ROLES) * len(_MINS)
    first = rates[0]
    assert first["role"] == "Networking L1 Technician"
    assert first["minimum_hours"] == 2.0
    assert first["rate"] == 60

    materials = [a for a in _sheet(atoms, "Materials") if not a.value.get("is_summary")]
    assert len(materials) == 45
    assert all(a.atom_type == AtomType.pricing_assumption for a in materials)

    # The helper sheet is still skipped, as one marker.
    assert [a.atom_type for a in atoms if a.atom_type == AtomType.dropped_sheet] == [AtomType.dropped_sheet]
    # Never scope.
    assert not [a for a in atoms if a.atom_type == AtomType.scope_item]


def test_rollup_stays_only_as_a_summary(tmp_path) -> None:
    path = tmp_path / "Deal_Kit.xlsx"
    _deal_kit(path)
    atoms = _atoms(path)
    summaries = [a for a in _sheet(atoms, "COST RATES") if a.value.get("is_summary")]
    assert len(summaries) == 1
    assert summaries[0].value["line_count"] == len(_RATE_ROWS)
    assert len(summaries[0].value["rows"]) == len(_RATE_ROWS)


def _form_row(idx: int) -> EvidenceAtom:
    cells = {"Question": f"Site question {idx}", "Unit Price": ""}
    text = f"Question: Site question {idx}"
    src = SourceRef(
        id=f"src_{idx}", artifact_id="artF", artifact_type=ArtifactType.docx,
        filename="form.docx", locator={"row": idx + 1},
        extraction_method="t", parser_version="t",
    )
    return EvidenceAtom(
        id=f"atm_form_{idx}", project_id="p", artifact_id="artF",
        atom_type=AtomType.scope_item, raw_text=text, normalized_text=text.lower(),
        value={"kind": "table_row", "cells": cells, "sheet": ""},
        source_refs=[src], authority_class=AuthorityClass.contractual_scope,
        confidence=0.5, review_status=ReviewStatus.needs_review, parser_version="t",
    )


def test_money_header_without_money_is_not_rolled_up() -> None:
    # ": 53 table rows (rolled up)" -- a form under a price header with no
    # price in it folded at the commercial threshold and vanished.
    atoms = [_form_row(i) for i in range(53)]
    out, stats = roll_up_table_rows(atoms)
    assert stats["groups_folded"] == 0
    assert len(out) == 53
