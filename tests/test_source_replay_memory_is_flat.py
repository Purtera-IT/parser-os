"""source_replay's memory does not grow with the number of atoms it checks.

Deal 010264 (87 documents, 16,001 source_replay items) was killed for memory
three times on the 8 GB parser worker, each time 4,000-5,700 items into this
stage. Memory climbed ~1.75 GB per 1,000 items; a run on a 40 GB machine
peaked near 30 GB, more than even the 28 GB worker it was moved to.

The cause was the workbook cache. It kept full openpyxl Workbooks, and reading
a row of one creates -- and keeps -- a Cell for every column up to the sheet's
max_column. A single formatted cell out at column XFD makes that 16,384 cells,
about 2 MB, for every row an atom points at. On a synthetic copy of that shape
the old code grew 1,956 MB over 1,000 rows; it now holds the sheet's values
once and grows by single-digit megabytes.

These tests replay a deal of 010264's size and shape in a fresh interpreter and
hold the peak, and the growth over the second half, to fixed ceilings. Some
growth is the output itself: every atom keeps its receipts (a pydantic model
with its locator and snippet), 46 MB per 8,000 items on Python 3.13 and 121 MB
on CI's 3.12. The ceilings sit above that and fifty times below the old leak,
so anything leaking more than ~40 KB per item fails them.
"""
from __future__ import annotations

import json
import subprocess
import sys
import textwrap
from pathlib import Path

import pytest

pytest.importorskip("openpyxl")
pytest.importorskip("docx")

if not Path("/proc/self/status").exists():  # pragma: no cover - non-Linux dev box
    pytest.skip("reads RSS from /proc", allow_module_level=True)

REPO = Path(__file__).resolve().parents[1]

#: 010264's source_replay item count.
ITEMS = 16_001
#: The prod parser worker had 4 GB and dev 8 GB when 010264 was killed. The
#: whole replay of a deal this size must fit with room to spare for every other
#: stage, so it gets at most a fifth of the smallest of them.
PEAK_GROWTH_CEILING_MB = 800
#: The second 8,000 items may not cost more memory than this. The old code
#: spent ~14 GB on them.
SECOND_HALF_CEILING_MB = 300

_SCRIPT = textwrap.dedent(
    r'''
    import json, sys
    from pathlib import Path

    import openpyxl
    from docx import Document
    from openpyxl.styles import PatternFill

    from app.core.schemas import (
        ArtifactType, AtomType, AuthorityClass, EvidenceAtom, ReviewStatus, SourceRef,
    )
    from app.core.source_replay import attach_receipts_to_atoms, summarize_receipts

    out = Path(sys.argv[1])
    items = int(sys.argv[2])

    def rss_mb():
        for line in Path("/proc/self/status").read_text().splitlines():
            if line.startswith("VmRSS:"):
                return int(line.split()[1]) / 1024
        return 0.0

    def peak_mb():
        for line in Path("/proc/self/status").read_text().splitlines():
            if line.startswith("VmHWM:"):
                return int(line.split()[1]) / 1024
        return 0.0

    # The deal: two big workbooks, each with one fill reaching column XFD (the
    # shape that leaked), a long text export cited by line range, and a SOW.
    sheet_rows = 6_000
    paths = {}
    for n in (1, 2):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sites"
        for r in range(1, sheet_rows + 1):
            ws.cell(r, 1, f"Site {n}-{r} install camera")
            ws.cell(r, 2, r)
            ws.cell(r, 3, f"{r} Main St")
        ws.cell(2, 16384).fill = PatternFill("solid", fgColor="FFFF00")
        p = out / f"sites{n}.xlsx"
        wb.save(p)
        paths[f"xlsx{n}"] = p
    text_lines = 4_000
    t = out / "thread.txt"
    t.write_text("\n".join(f"Line {i}: please confirm crew for site {i}" for i in range(1, text_lines + 1)))
    paths["txt"] = t
    doc = Document()
    for i in range(200):
        doc.add_paragraph(f"Section {i}: the contractor installs {i} cameras")
    d = out / "sow.docx"
    doc.save(d)
    paths["docx"] = d

    def atom(i, artifact, loc, text):
        ref = SourceRef(
            id=f"src_{i}", artifact_id=artifact, artifact_type=ArtifactType.xlsx,
            filename=paths[artifact].name, locator=loc, extraction_method="t", parser_version="t",
        )
        return EvidenceAtom(
            id=f"atm_{i}", project_id="p", artifact_id=artifact, atom_type=AtomType.scope_item,
            raw_text=text, normalized_text=text.lower(), value={}, entity_keys=[], source_refs=[ref],
            authority_class=AuthorityClass.customer_current_authored, confidence=0.9,
            review_status=ReviewStatus.auto_accepted, parser_version="t",
        )

    atoms = []
    i = 0
    while len(atoms) < items:
        i += 1
        kind = i % 8
        if kind < 6:  # three in four point into a spreadsheet
            n = 1 + (i % 2)
            r = 1 + (i % sheet_rows)
            loc = {"sheet": "Sites", "row": r}
            if i % 3 == 0:
                loc["columns"] = {"site": "A"}
            atoms.append(atom(i, f"xlsx{n}", loc, f"Site {n}-{r} install camera"))
        elif kind == 6:
            ln = 1 + (i % text_lines)
            atoms.append(atom(i, "txt", {"line_start": ln, "line_end": ln}, f"please confirm crew for site {ln}"))
        else:
            k = i % 200
            atoms.append(atom(i, "docx", {"paragraph_index": k}, f"the contractor installs {k} cameras"))

    artifact_paths = {k: v for k, v in paths.items()}
    half = len(atoms) // 2
    base = rss_mb()
    attach_receipts_to_atoms(atoms[:half], artifact_paths)
    mid = rss_mb()
    attach_receipts_to_atoms(atoms[half:], artifact_paths)
    end = rss_mb()
    print(json.dumps({
        "items": len(atoms),
        "base_mb": base,
        "mid_mb": mid,
        "end_mb": end,
        "peak_mb": peak_mb(),
        "summary": summarize_receipts(atoms),
    }))
    '''
)


def _replay(tmp_path: Path, items: int) -> dict:
    script = tmp_path / "replay_size.py"
    script.write_text(_SCRIPT)
    proc = subprocess.run(
        [sys.executable, str(script), str(tmp_path), str(items)],
        cwd=REPO, capture_output=True, text=True, timeout=900,
        env={"PYTHONPATH": str(REPO), "PATH": "/usr/bin:/bin"},
    )
    assert proc.returncode == 0, proc.stderr[-4000:]
    return json.loads(proc.stdout.strip().splitlines()[-1])


@pytest.fixture(scope="module")
def deal_010264_size(tmp_path_factory) -> dict:
    return _replay(tmp_path_factory.mktemp("replay_size"), ITEMS)


def test_every_item_is_replayed(deal_010264_size):
    s = deal_010264_size["summary"]
    assert deal_010264_size["items"] == ITEMS
    assert s["verified"] + s["failed"] + s["unsupported"] == ITEMS
    # The same rows verify as before the change: the deal is built so every
    # atom's text is really in its source.
    assert s["verified"] == ITEMS, s


def test_peak_memory_of_a_16k_item_deal_stays_small(deal_010264_size):
    growth = deal_010264_size["peak_mb"] - deal_010264_size["base_mb"]
    assert growth < PEAK_GROWTH_CEILING_MB, deal_010264_size


def test_memory_is_flat_over_the_second_half(deal_010264_size):
    second_half = deal_010264_size["end_mb"] - deal_010264_size["mid_mb"]
    assert second_half < SECOND_HALF_CEILING_MB, deal_010264_size
