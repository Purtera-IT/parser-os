"""Every line of a SOW docx is an atom: headings, lead-ins, placeholders.

Deal 000132's SOW v2 lost its section headings ("A. IT Infrastructure
Support"), list lead-ins ("Provider is responsible for the following:") and
Word placeholder text ("Click here to enter phone number") -- none became an
atom, so none could be labeled. Headings / lead-ins are kept as chatter
rejects with a reason; a placeholder is an open question (an unfilled SOW
field), whether it sits in a content control (w:sdt + w:showingPlcHdr) or is
plain text.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from app.parsers.docx_parser import DocxParser

W = nsdecls("w")
PLC_RUN = '<w:r><w:rPr><w:rStyle w:val="PlaceholderText"/></w:rPr><w:t>{}</w:t></w:r>'


def _build(path: Path) -> None:
    d = Document()
    d.add_heading("Scope of Services", 1)
    p = d.add_paragraph(); p.add_run("A. IT Infrastructure Support").bold = True
    d.add_paragraph("Provider is responsible for the following:")
    d.add_paragraph("Install and configure network switches at each site.", style="List Bullet")
    d.add_paragraph("Deploy wireless access points per the approved floor plan.", style="List Bullet")
    d.add_heading("Customer Contact", 1)
    p = d.add_paragraph(); p.add_run("Phone: ")
    p._p.append(parse_xml(f'<w:sdt {W}><w:sdtPr><w:showingPlcHdr/></w:sdtPr><w:sdtContent>'
                          + PLC_RUN.format("Click here to enter phone number") + '</w:sdtContent></w:sdt>'))
    d.element.body.append(parse_xml(
        f'<w:sdt {W}><w:sdtPr><w:showingPlcHdr/></w:sdtPr><w:sdtContent><w:p>'
        + PLC_RUN.format("Click here to enter text.") + '</w:p></w:sdtContent></w:sdt>'))
    d.add_paragraph("Site contact email: Click here to enter text.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Contact"; t.cell(0, 1).text = "Phone"
    t.cell(1, 0).text = "Jane Roe"; t.cell(1, 1).text = "Click or tap here to enter text."
    d.save(path)


def _atoms(tmp_path):
    path = tmp_path / "SOW v2.docx"
    _build(path)
    out = DocxParser().parse_artifact("p", "a", path)
    return out if isinstance(out, list) else out.atoms


def test_headings_and_lead_ins_are_reject_atoms(tmp_path):
    by = {a.raw_text: a for a in _atoms(tmp_path)}
    for text, kind in (("Scope of Services", "section_heading"),
                       ("A. IT Infrastructure Support", "section_heading"),
                       ("Customer Contact", "section_heading"),
                       ("Provider is responsible for the following:", "list_lead_in")):
        a = by.get(text)
        assert a is not None, (text, list(by))
        assert a.atom_type.value == "deal_metadata"
        assert "chatter" in a.review_flags and a.value["rejected_by"] == kind, (text, a.value)
    # Still structure for the children: the bullets carry the heading path.
    bullet = by["Install and configure network switches at each site."]
    assert "A. IT Infrastructure Support" in bullet.source_refs[0].locator["section_path"]


def test_placeholders_are_open_questions(tmp_path):
    atoms = _atoms(tmp_path)
    ph = [a for a in atoms if "unfilled_placeholder" in a.review_flags]
    texts = [a.raw_text for a in ph]
    assert "Phone: Click here to enter phone number" in texts, [a.raw_text for a in atoms]
    assert "Click here to enter text." in texts
    assert "Site contact email: Click here to enter text." in texts
    assert any("Click or tap here to enter text." in t for t in texts), texts
    phone = next(a for a in ph if a.raw_text.startswith("Phone:"))
    assert phone.atom_type.value == "open_question"
    assert phone.value["field_label"] == "Phone"
    assert phone.value["open_question_reason"] == "unfilled SOW placeholder"
    assert "chatter" not in phone.review_flags


def test_compile_keeps_every_line(tmp_path):
    from app.core.compiler import compile_project

    _build(tmp_path / "SOW v2.docx")
    r = compile_project(tmp_path, project_id="p", allow_errors=True, use_cache=False)
    every = [a.raw_text for a in r.atoms] + [a.raw_text for a in r.suppressed_atoms]
    for line in ("A. IT Infrastructure Support", "Provider is responsible for the following:",
                 "Phone: Click here to enter phone number"):
        assert line in every, (line, every)
