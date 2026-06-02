"""Authority-weighted domain routing over office-doc content.

Grounded in the Yonah deal: a residential TV-install SOW shipped as a
.docx ("STATEMENT OF WORK ... replace approximately 110 existing TVs and
mounts") plus a financial deal-kit .xlsx (margin / revenue / cost, with
incidental infra words like "cat6"). Before this fix the router only read
.txt/.md/.csv, so office-doc deals had no content signal and routed on
incidental filename/customer-name tokens — a TV install got scored as a
cabling/datacenter job. The router must now:

  * read .docx / .xlsx content,
  * match aliases on word boundaries (no "ap" inside "approximately"),
  * weight a statement-of-work above a deal-kit money sheet,

so the authoritative scope drives the pack.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from app.domain.pack_router import (
    _alias_in_text,
    _blob_authority,
    _read_text_preview,
    auto_route_pack,
)

docx = pytest.importorskip("docx")
from openpyxl import Workbook  # noqa: E402


def _write_sow_docx(path: Path) -> None:
    doc = docx.Document()
    doc.add_paragraph("STATEMENT OF WORK (SOW)")
    doc.add_paragraph("PROJECT OVERVIEW")
    doc.add_paragraph(
        "The customer requires onsite field services to replace approximately "
        "110 existing TVs and wall mounts across a resort property. The work "
        "will be performed across 23 dwellings. Mount each television and "
        "verify the display powers on."
    )
    doc.save(str(path))


def _write_deal_kit_xlsx(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.append(["Deal Kit Summary", "Net Margin"])
    ws.append(["Total Deal Revenue", 21560])
    ws.append(["Total Deal Cost", 15660])
    ws.append(["Total Deal Margin", 5900])
    # Incidental infra vocabulary that must NOT drive routing.
    ws.append(["Materials", "cat6 patch, switch spare"])
    wb.save(str(path))


# ── preview extraction ──────────────────────────────────────────────


def test_docx_preview_extracts_scope_text(tmp_path: Path) -> None:
    p = tmp_path / "sow.docx"
    _write_sow_docx(p)
    preview = _read_text_preview(p).lower()
    assert "statement of work" in preview
    assert "tvs" in preview
    assert "wall mount" in preview


def test_xlsx_preview_extracts_cell_text(tmp_path: Path) -> None:
    p = tmp_path / "deal_kit.xlsx"
    _write_deal_kit_xlsx(p)
    preview = _read_text_preview(p).lower()
    assert "deal kit summary" in preview
    assert "21560" in preview


def test_unknown_binary_preview_is_empty(tmp_path: Path) -> None:
    p = tmp_path / "image.png"
    p.write_bytes(b"\x89PNG\r\n\x1a\n not really an image")
    assert _read_text_preview(p) == ""


# ── word-boundary alias matching ────────────────────────────────────


def test_alias_word_boundary_rejects_substring() -> None:
    # "ap" must not match inside "approximately"; "tr" not inside "centre".
    assert not _alias_in_text("ap", "replace approximately 110 units")
    assert not _alias_in_text("tr", "the centre of the room")


def test_alias_word_boundary_accepts_whole_word() -> None:
    assert _alias_in_text("tv", "mount each tv on the wall")
    assert _alias_in_text("wall mount", "install a wall mount per unit")


def test_single_char_alias_is_ignored() -> None:
    assert not _alias_in_text("a", "a building with a door")


# ── source authority ───────────────────────────────────────────────


def test_sow_outweighs_deal_kit_authority() -> None:
    sow = _blob_authority("sow.docx", "STATEMENT OF WORK\nproject overview ...")
    deal = _blob_authority(
        "deal_kit.xlsx", "Deal Kit Summary\nNet Margin\nTotal Deal Revenue"
    )
    plain = _blob_authority("notes.txt", "some neutral note about the job")
    assert sow > plain > deal


# ── end-to-end routing ──────────────────────────────────────────────


def test_tv_install_routes_to_av_not_infra(tmp_path: Path) -> None:
    art = tmp_path / "artifacts"
    art.mkdir()
    _write_sow_docx(art / "Statement of Work.docx")
    _write_deal_kit_xlsx(art / "Deal Kit.xlsx")
    pack, decision = auto_route_pack(tmp_path)
    assert pack.pack_id == "av", (
        f"expected av, got {pack.pack_id} via {decision.source}: "
        f"{decision.rationale} | alts={decision.alternatives}"
    )
    assert decision.source == "content"


def test_money_sheet_alone_does_not_route_to_infra(tmp_path: Path) -> None:
    # Only the deal-kit xlsx (no SOW). Its incidental infra words must not
    # produce a confident non-default pack.
    art = tmp_path / "artifacts"
    art.mkdir()
    _write_deal_kit_xlsx(art / "Deal Kit.xlsx")
    pack, decision = auto_route_pack(tmp_path)
    assert pack.pack_id in {"default_pack"}, (
        f"money sheet alone should fall to default, got {pack.pack_id}: "
        f"{decision.rationale}"
    )
