"""A docx contacts row keeps its EMAIL ADDRESS cell through the compile (010087).

Shape of the live SOW (synthetic names): a plain 3-column table, header
"FULL NAME | JOB TITLE | EMAIL ADDRESS", no merged cells, the email cell
written as two runs ("a@" + "example.com"). The parser read the row whole, but
the stakeholder_table schema atom spoke only "name | title" (the address sat
in its value), and pre_classify_dedup folded both full-row atoms into it, so
the only standing atom for the row lost the email.
"""

from __future__ import annotations

from docx import Document

from app.core.compiler import compile_project
from app.core.schemas import ArtifactType, AtomType, AuthorityClass, EvidenceAtom, ReviewStatus, SourceRef
from app.core.semantic_dedup import _suppress_table_row_blob_doubles
from app.core.table_schema_registry import emit_atoms_for_schema

COLS = ["FULL NAME", "JOB TITLE", "EMAIL ADDRESS"]
ROWS = [
    ("Alex Example", "Executive VP Sales", "a@", "example.com"),
    ("Bo Sample", "Solution Architect", "bo@", "example.com"),
]


def _contacts_docx(path):
    d = Document()
    d.add_heading("SALES CONTACTS", level=1)
    t = d.add_table(rows=3, cols=3)
    for i, h in enumerate(COLS):
        t.cell(0, i).paragraphs[0].add_run(h).bold = True
    for r, (name, title, local, domain) in enumerate(ROWS, 1):
        t.cell(r, 0).text = name
        t.cell(r, 1).text = title
        p = t.cell(r, 2).paragraphs[0]
        p.add_run(local)
        p.add_run(domain)
    assert "gridSpan" not in t._tbl.xml and "vMerge" not in t._tbl.xml
    d.save(path)


def test_stakeholder_schema_atom_states_the_whole_row():
    atoms = emit_atoms_for_schema(
        schema_name="stakeholder_table", columns=COLS,
        row=["Alex Example", "Executive VP Sales", "a@example.com"],
        row_idx=1, table_idx=0, project_id="p", artifact_id="art", filename="sow.docx",
    )
    people = [a for a in atoms if a.atom_type == AtomType.stakeholder]
    assert [a.raw_text for a in people] == ["Alex Example | Executive VP Sales | a@example.com"]
    assert people[0].value["email"] == "a@example.com"


def _atom(aid, atype, text, value, extraction):
    ref = SourceRef(id=f"src_{aid}", artifact_id="art", artifact_type=ArtifactType.docx, filename="sow.docx",
                    locator={"table_index": 0, "row": 1, "extraction": extraction},
                    extraction_method=extraction, parser_version="t")
    return EvidenceAtom(
        id=aid, project_id="p", artifact_id="art", atom_type=atype, raw_text=text,
        normalized_text=text.lower(), value=value, entity_keys=[], source_refs=[ref], receipts=[],
        authority_class=AuthorityClass.contractual_scope, confidence=0.8, review_status=ReviewStatus.auto_accepted,
        review_flags=[], parser_version="t",
    )


def test_row_blob_is_not_folded_into_an_atom_that_holds_the_email_only_in_its_value():
    rich = _atom("atm_rich", AtomType.stakeholder, "Alex Example | Executive VP Sales",
                 {"name": "Alex Example", "email": "a@example.com"}, "table_schema_v49")
    blob = _atom("atm_blob", AtomType.scope_item, "Alex Example | Executive VP Sales | a@example.com",
                 {"kind": "table_row"}, "docx_table_row_v1")
    out = _suppress_table_row_blob_doubles([rich, blob])
    assert [a.id for a in out] == ["atm_rich", "atm_blob"]


def test_contact_row_keeps_its_email_through_the_compile(tmp_path):
    proj = tmp_path / "deal"
    proj.mkdir()
    _contacts_docx(proj / "SOW_v1.docx")
    _contacts_docx(proj / "SOW_v2.docx")
    r = compile_project(proj, project_id="p", allow_errors=True, use_cache=False)
    for name, title, local, domain in ROWS:
        full = f"{name} | {title} | {local}{domain}"
        live = [a for a in r.atoms if name in (a.raw_text or "")]
        assert live, f"no live atom for {name}"
        assert all(f"{local}{domain}" in a.raw_text for a in live), [a.raw_text for a in live]
        assert any(a.raw_text == full for a in live), [a.raw_text for a in live]
