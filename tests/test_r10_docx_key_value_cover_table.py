"""A docx LABEL | VALUE cover table is read as one atom per field.

Live 000132 (structure only; every text here is synthetic): the SOW cover
table has three grid columns. Column 0 holds labels ("Project Name:"), column 1
the values (each inside a w:sdt content control), and column 2 holds cells that
span several rows (w:vMerge) and label themselves ("Requested By:" over a name,
a phone placeholder and an email). Read row by row with row 0 as the header,
every row came out as "Label: | value | Requested By: <name> <phone> <email>"
(the spanning cell repeated on each row it covers), and the email back-scan
over that glued row read the project name's first words as a person, a
stakeholder atom holding only a fragment of the project-name cell.
The PDF twin of this table is #320.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from app.core.compiler import compile_project
from app.parsers.docx_parser import DocxParser

W = nsdecls("w")
BR = "<w:r><w:br/></w:r>"


def _r(text: str, bold: bool = False) -> str:
    rpr = "<w:rPr><w:b/></w:rPr>" if bold else ""
    return f'<w:r>{rpr}<w:t xml:space="preserve">{text}</w:t></w:r>'


def _sdt(alias: str, inner: str, placeholder: bool = False) -> str:
    plc = "<w:showingPlcHdr/>" if placeholder else ""
    return (f'<w:sdt><w:sdtPr><w:alias w:val="{alias}"/><w:tag w:val="{alias}"/>{plc}<w:text/>'
            f"</w:sdtPr><w:sdtContent>{inner}</w:sdtContent></w:sdt>")


def _tc(paras: list[str], vmerge: str | None = None) -> str:
    vm = {"restart": '<w:vMerge w:val="restart"/>', "cont": "<w:vMerge/>"}.get(vmerge or "", "")
    return f'<w:tc><w:tcPr><w:tcW w:w="2000" w:type="dxa"/>{vm}</w:tcPr>' + "".join(
        f"<w:p>{p}</w:p>" for p in paras) + "</w:tc>"


def _tbl(rows: list[list[str]], ncols: int) -> str:
    return (f'<w:tbl {W}><w:tblPr><w:tblW w:w="0" w:type="auto"/></w:tblPr><w:tblGrid>'
            + '<w:gridCol w:w="2000"/>' * ncols + "</w:tblGrid>"
            + "".join("<w:tr>" + "".join(r) + "</w:tr>" for r in rows) + "</w:tbl>")


def _cover_table() -> str:
    """The real table's shape: label | sdt value | spanning self-labelled cell."""
    return _tbl([
        [_tc(['<w:r><w:br w:type="page"/></w:r>' + _r("Project Name:")]),
         _tc([_sdt("Project Name", _r("Harbor Point") + _r(" Fiber Upgrade"))]),
         _tc([_r("Requested By:"), _sdt("Rep Name", _r("Avery Quinn")) + BR
              + _sdt("Rep Phone", _r("Click here to enter phone number."), True) + BR
              + _sdt("Rep Email", _r("avery.quinn@example.com"))], "restart")],
        [_tc([_r("Customer Name:")]), _tc([_sdt("Customer", _r("Example Holdings"))]), _tc([""], "cont")],
        [_tc([_r("Provider Name:")]), _tc([_sdt("Provider", _r("Sample Services"))]), _tc([""], "cont")],
        [_tc([_r("Affiliate:")]), _tc([_sdt("Affiliate", _r("Example Partners, Inc."))]),
         _tc([_r("Submitted By:"), _sdt("Author", _r("Sample Services"))], "restart")],
        [_tc([_r("Effective Date:")]), _tc([_sdt("Date", _r("March 3, 2026"))]), _tc([""], "cont")],
        [_tc([_r("Version:")]), _tc([_sdt("Version", _r("V1"))]), _tc([""], "cont")],
    ], 3)


def _save(tmp_path: Path, *tables: str, name: str = "SOW.docx") -> Path:
    d = Document()
    d.add_paragraph("Statement of Work")
    body = d.element.body
    for t in tables:
        body.insert(len(body) - 1, parse_xml(t))
    d.add_heading("Project Scope", 1)
    d.add_paragraph("Provider will deliver recurring on site support services for every listed office.")
    path = tmp_path / name
    d.save(path)
    return path


def _rows(path: Path) -> list[str]:
    atoms = DocxParser().parse_artifact("p", "a", path)
    return [a.raw_text for a in atoms if a.source_refs[0].locator.get("table_index") is not None
            and a.source_refs[0].extraction_method == "docx_table_row_v1"]


def test_cover_table_is_one_atom_per_field(tmp_path):
    assert _rows(_save(tmp_path, _cover_table())) == [
        "Project Name: Harbor Point Fiber Upgrade",
        "Requested By:\nAvery Quinn\nClick here to enter phone number.\navery.quinn@example.com",
        "Customer Name: Example Holdings",
        "Provider Name: Sample Services",
        "Affiliate: Example Partners, Inc.",
        "Submitted By:\nSample Services",
        "Effective Date: March 3, 2026",
        "Version: V1",
    ]


def test_cover_table_compiles_without_glued_rows_or_name_fragments(tmp_path):
    proj = tmp_path / "proj"
    proj.mkdir()
    _save(proj, _cover_table())
    r = compile_project(proj, project_id="p", allow_errors=True, use_cache=False)
    texts = [a.raw_text for a in r.atoms]
    assert "Project Name: Harbor Point Fiber Upgrade" in texts
    assert "Customer Name: Example Holdings" in texts
    # the spanning cell is one atom, never repeated onto the rows it covers
    assert sum("Avery Quinn" in t and "@" in t for t in texts) == 1
    assert not any("|" in t and "Requested By" in t for t in texts)
    # the project-name cell is never cut into a person
    assert "Harbor Point" not in texts
    assert not any(t.startswith("Harbor Point") and t != "Harbor Point Fiber Upgrade" for t in texts)


def test_two_column_label_grid_keeps_its_first_row(tmp_path):
    # Row 0 has no digit, so it used to read as a header row and vanish.
    t = _tbl([
        [_tc([_r("Project Name:", True)]), _tc([_r("Harbor Point Fiber Upgrade")])],
        [_tc([_r("Customer Name:", True)]), _tc([_r("Example Holdings")])],
    ], 2)
    assert _rows(_save(tmp_path, t)) == [
        "Project Name: Harbor Point Fiber Upgrade",
        "Customer Name: Example Holdings",
    ]


def test_a_headed_table_is_not_a_label_grid(tmp_path):
    # Bold header over plain data rows: column 0's data cells are not labels.
    t = _tbl([
        [_tc([_r("Role", True)]), _tc([_r("Standard Rate", True)])],
        [_tc([_r("Field Technician")]), _tc([_r("$95.00/hr")])],
        [_tc([_r("Project Manager")]), _tc([_r("$125.00/hr")])],
    ], 2)
    assert _rows(_save(tmp_path, t)) == ["Field Technician | $95.00/hr", "Project Manager | $125.00/hr"]


def test_a_third_column_that_does_not_span_is_a_data_column(tmp_path):
    t = _tbl([
        [_tc([_r("Phase:")]), _tc([_r("Survey")]), _tc([_r("Week 1")])],
        [_tc([_r("Phase:")]), _tc([_r("Install")]), _tc([_r("Week 2")])],
    ], 3)
    rows = _rows(_save(tmp_path, t))
    assert not any(r.startswith("Phase: Survey") for r in rows)
