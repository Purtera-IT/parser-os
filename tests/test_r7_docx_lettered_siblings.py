"""R7 (000132): typed lettered lead lines under one heading are siblings.

000132's SOW drafts put "A." .. "I." group lead lines on Normal, all bold,
each over its own bullet list. Most carried ``w:outlineLvl 3`` (depth 4) and
two did not (bold only, depth 3), so the heading stack nested D under C and
F..I under E. A line that continues the same typed sequence as an open lead
line is its peer, whatever depth its formatting suggested.
"""

from __future__ import annotations

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.parsers.docx_parser import DocxParser, _continues_enumeration

LETTERS = "ABCDEFG"
NO_OUTLINE = "CE"  # these two rank one level shallower than their peers
NUM_IDS = {"A": 15, "B": 16, "C": 16, "D": 17, "E": 18, "F": 19, "G": 19}


def _add_num(doc, num_id: int) -> None:
    numbering = doc.part.numbering_part.element
    abs_id = numbering.find(qn("w:abstractNum")).get(qn("w:abstractNumId"))
    n = OxmlElement("w:num")
    n.set(qn("w:numId"), str(num_id))
    a = OxmlElement("w:abstractNumId")
    a.set(qn("w:val"), abs_id)
    n.append(a)
    numbering.append(n)


def _bullet(doc, text: str, num_id: int) -> None:
    p = doc.add_paragraph(text)
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    nid = OxmlElement("w:numId")
    nid.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(nid)
    p._p.get_or_add_pPr().append(num_pr)


def _lead(doc, text: str, outline: int | None) -> None:
    p = doc.add_paragraph()
    if outline is not None:
        o = OxmlElement("w:outlineLvl")
        o.set(qn("w:val"), str(outline))
        p._p.get_or_add_pPr().append(o)
    p.add_run(text[:3]).bold = True
    p.add_run(text[3:]).bold = True


def _child(letter: str, k: int) -> str:
    return f"Group {letter} duty {k}: maintain and monitor the managed equipment"


def _build(path) -> None:
    doc = Document()
    for num_id in sorted(set(NUM_IDS.values())):
        _add_num(doc, num_id)
    doc.add_heading("Project Scope", level=1)
    doc.add_paragraph("Provider will deliver recurring on site support services for the customer.")
    doc.add_paragraph("Services will be provided on a time and materials basis each month.")
    doc.add_paragraph("Provider may perform the services in the groups listed below.")
    for letter in LETTERS:
        if letter == LETTERS[-1]:
            for _ in range(3):
                doc.add_paragraph("")
        _lead(doc, f"{letter}. Support Group {letter}", None if letter in NO_OUTLINE else 3)
        for k in (1, 2):
            _bullet(doc, _child(letter, k), NUM_IDS[letter])
    doc.add_heading("Assumptions", level=1)
    doc.add_paragraph("Customer will provide site access during business hours.")
    doc.save(path)


def _path(a):
    return list(a.source_refs[0].locator.get("section_path") or [])


def test_lettered_lead_lines_are_siblings(tmp_path):
    path = tmp_path / "sow.docx"
    _build(path)
    atoms = DocxParser().parse_artifact("p", "a", path)
    for letter in LETTERS:
        lead = f"{letter}. Support Group {letter}"
        for k in (1, 2):
            hits = [a for a in atoms if a.raw_text == _child(letter, k)]
            assert hits, _child(letter, k)
            assert _path(hits[0]) == ["Project Scope", lead], (letter, _path(hits[0]))
        heads = [a for a in atoms if a.raw_text == lead]
        assert heads and _path(heads[0]) == ["Project Scope", lead]


def test_different_enumerator_family_still_nests(tmp_path):
    """"1." under "A." is a sub-group, not a sibling."""
    doc = Document()
    _add_num(doc, 15)
    doc.add_heading("Project Scope", level=1)
    _lead(doc, "A. Support Group A", None)
    _lead(doc, "1. Network Tasks", 3)
    _bullet(doc, _child("A", 1), 15)
    path = tmp_path / "nested.docx"
    doc.save(path)
    atoms = DocxParser().parse_artifact("p", "a", path)
    hit = [a for a in atoms if a.raw_text == _child("A", 1)][0]
    assert _path(hit) == ["Project Scope", "A. Support Group A", "1. Network Tasks"]


def test_continues_enumeration():
    assert _continues_enumeration("C. Server", "D. VMware")
    assert _continues_enumeration("H. Weekly", "I. General")
    assert _continues_enumeration("I. Scope", "II. Fees")
    assert _continues_enumeration("2. Scope", "3. Fees")
    assert not _continues_enumeration("A. Scope", "1. Fees")
    assert not _continues_enumeration("B. Scope", "A. Fees")
    assert not _continues_enumeration("a) Scope", "B. Fees")


# The document view (``structured.json`` / ``structured.md``, the envelope's
# per-document projection) walks the paragraphs on its own. It opened sections
# only on outlined lines, so "C." and "E." were plain lines inside B's and D's
# sections there even after the atoms' section paths were fixed.

def _outline(path):
    import docx as _docx
    sd = DocxParser()._build_structured_doc(filename="sow.docx", document=_docx.Document(str(path)))
    return [(o["level"], o["heading"], o["block_count"]) for o in sd["pages"][0]["outline"]]


def test_document_view_lettered_lead_lines_are_siblings(tmp_path):
    path = tmp_path / "sow.docx"
    _build(path)
    groups = [o for o in _outline(path) if o[1].endswith(tuple(f"Support Group {x}" for x in LETTERS))]
    assert [g[1] for g in groups] == [f"{x}. Support Group {x}" for x in LETTERS]
    assert {g[0] for g in groups} == {5}  # outlineLvl 3 -> level 5, every peer
    assert all(g[2] == 2 for g in groups), groups  # each holds only its own two lines


def test_document_view_numbered_peer_with_deeper_outline_is_sibling(tmp_path):
    doc = Document()
    _add_num(doc, 15)
    doc.add_heading("Project Scope", level=1)
    _lead(doc, "1. Network Tasks", 2)
    _bullet(doc, _child("A", 1), 15)
    _lead(doc, "2. Server Tasks", 3)
    _bullet(doc, _child("B", 1), 15)
    path = tmp_path / "num.docx"
    doc.save(path)
    levels = {h: lv for lv, h, _ in _outline(path)}
    assert levels["1. Network Tasks"] == levels["2. Server Tasks"] == 4


def test_document_view_keeps_other_family_nested_and_body_lines_plain(tmp_path):
    doc = Document()
    _add_num(doc, 15)
    doc.add_heading("Project Scope", level=1)
    _lead(doc, "A. Support Group A", 3)
    _lead(doc, "1. Network Tasks", 4)
    _bullet(doc, _child("A", 1), 15)
    # Not bold, and a sentence: body text that happens to start with "B.".
    doc.add_paragraph("B. The provider will also report on these tasks every month.")
    path = tmp_path / "mixed.docx"
    doc.save(path)
    out = _outline(path)
    assert [(lv, h) for lv, h, _ in out if h[:2] in ("A.", "1.", "B.")] == [
        (5, "A. Support Group A"), (6, "1. Network Tasks")]
