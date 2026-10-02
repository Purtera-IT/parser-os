"""A bullet's marker is list metadata, never part of the atom text.

Live 000132: one HubSpot note typed its scope one "- " bullet per line and its
atoms kept the "- " ("- 24/7 on-call availability"); a sibling note flattened
the same list onto one " - " line and its atoms came out bare. The same bullet
had two texts, two keys and two highlights. The notes below are synthetic but
mirror the real shapes: header, a truncated title repeated as the first body
line, a lead line, then either per-line "- " bullets or one " - " run with a
double-escaped "&amp;amp;".
"""
from __future__ import annotations

import docx
import pytest
from pathlib import Path

from app.core.sentences import strip_list_marker
from app.parsers.docx_parser import DocxParser
from app.parsers.email_parser import EmailParser
from app.parsers.hubspot_note_parser import HubspotNoteParser

SHARED = [
    "Support for branch wireless access points",
    "Support for the file server and backup appliance",
    "Troubleshooting and incident response",
    "24/7 on-call availability",
]
LINE_NOTE = (
    "HubSpot Note: Managed support of the branch office network and servers…\n"
    "HubSpot Note ID: 900000000001\nDate: 2026-05-29T18:58:26.171Z\nAuthor: Pat Example\n"
    "Author-Email: pat@example.com\n\n"
    "Managed support of the branch office network and servers…\n\n"
    "Managed support of the branch office network and servers, storage, and backup\n"
    + "".join(f"- {item}\n" for item in SHARED)
    + "- Monthly on-site visits (2–4 hours as needed)"
)
INLINE_NOTE = (
    "HubSpot Note: Managed support of the branch office network…\n"
    "HubSpot Note ID: 900000000002\nDate: 2026-05-29T18:58:14.007Z\nAuthor: Pat Example\n\n"
    "Managed support of the branch office network…\n\n"
    "Managed support of the branch office network, servers &amp;amp; backup - "
    + " - ".join(SHARED)
    + " - Monthly on-site visits (2–4 hours as needed) Locations Springfield, OH Shelbyville, IN"
)


def _note_atoms(tmp_path: Path, name: str, text: str):
    p = tmp_path / name
    p.write_text(text, encoding="utf-8")
    return p.read_text(encoding="utf-8").splitlines(), HubspotNoteParser().parse_artifact("p", "art", p)


def _highlight(lines: list[str], atom) -> str | None:
    loc = atom.source_refs[0].locator
    if loc.get("char_end") is None:
        return None
    return lines[loc["line_start"] - 1][loc["char_start"]:loc["char_end"]]


def test_line_bullets_and_inline_bullets_mint_the_same_item(tmp_path: Path):
    lines_437, atoms_437 = _note_atoms(tmp_path, "000132-hs-note-900000000001.txt", LINE_NOTE)
    _, atoms_233 = _note_atoms(tmp_path, "000132-hs-note-900000000002.txt", INLINE_NOTE)
    texts_437 = {a.raw_text for a in atoms_437}
    texts_233 = {a.raw_text for a in atoms_233}
    assert not any(t.startswith("- ") for t in texts_437 | texts_233)
    for item in SHARED:
        assert item in texts_437, item
        assert item in texts_233, item
    for a in atoms_437:
        if a.raw_text in SHARED:
            assert a.value.get("list_marker") == "-"
            # The highlight is the item, not the marker.
            assert _highlight(lines_437, a) == a.raw_text
            assert a.source_refs[0].locator["char_start"] == 2
    by_text_233 = {a.raw_text: a for a in atoms_233}
    assert by_text_233["24/7 on-call availability"].value.get("list_marker") == "-"
    # The inline item's highlight lands on the item in the entity-decoded line.
    from app.core.textio import decode_html_entities

    lines_233 = decode_html_entities(INLINE_NOTE).splitlines()
    assert _highlight(lines_233, by_text_233["24/7 on-call availability"]) == "24/7 on-call availability"


@pytest.mark.parametrize(
    ("line", "marker", "char_start"),
    [
        ("- 24/7 on-call availability", "-", 2),
        ("* 24/7 on-call availability", "*", 2),
        ("• 24/7 on-call availability", "•", 2),
        ("◦ 24/7 on-call availability", "◦", 2),
        ("\uf0b7 24/7 on-call availability", "\uf0b7", 2),
        ("– 24/7 on-call availability", "–", 2),
        ("— 24/7 on-call availability", "—", 2),
        ("o 24/7 on-call availability", "o", 2),
        ("1. 24/7 on-call availability", "1.", 3),
        ("1) 24/7 on-call availability", "1)", 3),
        ("a) 24/7 on-call availability", "a)", 3),
        ("  - 24/7 on-call availability", "-", 4),
        ("   -\t24/7 on-call availability", "-", 5),
    ],
)
def test_note_line_bullet_shapes(tmp_path: Path, line: str, marker: str, char_start: int):
    body = f"Scope of the support contract:\n{line}\n- Troubleshooting and incident response\n"
    text = (
        "HubSpot Note: Scope\nHubSpot Note ID: 1\nDate: 2026-05-29T18:58:26.171Z\nAuthor: Pat Example\n\n"
        + body
    )
    lines, atoms = _note_atoms(tmp_path, "000132-hs-note-1.txt", text)
    hit = [a for a in atoms if a.raw_text == "24/7 on-call availability"]
    assert hit, [a.raw_text for a in atoms]
    assert hit[0].value.get("list_marker") == marker
    assert hit[0].source_refs[0].locator["char_start"] == char_start
    assert _highlight(lines, hit[0]) == "24/7 on-call availability"


def test_note_html_list_items(tmp_path: Path):
    text = (
        "HubSpot Note: Scope\nHubSpot Note ID: 1\nDate: 2026-05-29T18:58:26.171Z\nAuthor: Pat Example\n\n"
        "<ul><li>Support for physical network components (LAN, WLAN)</li><li>24/7 on-call availability</li></ul>\n"
    )
    lines, atoms = _note_atoms(tmp_path, "000132-hs-note-2.txt", text)
    by_text = {a.raw_text: a for a in atoms}
    a = by_text["24/7 on-call availability"]
    assert a.value.get("list_marker") == "li"
    assert _highlight(lines, a) == "24/7 on-call availability"
    assert not any("<li>" in t for t in by_text)


def test_note_field_list_items_drop_markers(tmp_path: Path):
    text = (
        "HubSpot Note: Scope\nHubSpot Note ID: 1\nDate: 2026-05-29T18:58:26.171Z\nAuthor: Pat Example\n\n"
        "Address: 123 Main St, Springfield, OH 45501\nScope of work:\n"
        "- Support for physical network components (LAN, WLAN)\n• 24/7 on-call availability\nDate: 6/1\n"
    )
    _, atoms = _note_atoms(tmp_path, "000132-hs-note-3.txt", text)
    items = {a.raw_text for a in atoms if a.value.get("kind") == "note_field_item"}
    assert "24/7 on-call availability" in items
    assert "Support for physical network components (LAN, WLAN)" in items


def test_minus_numbers_and_hyphenated_words_are_not_markers():
    for text in ("-5 degrees at the dock", "-based pricing", "1.5 hours on site", "o365 licences",
                 "on-call availability", "a)b"):
        assert strip_list_marker(text) == ("", text), text
    assert strip_list_marker("- -5 degrees") == ("-", "-5 degrees")


def test_email_symbol_bullet_and_letter_marker(tmp_path: Path):
    p = tmp_path / "scope.eml"
    p.write_bytes(
        (
            "From: Customer Ops <ops@customer.com>\nTo: pat@example.com\nSubject: Scope\n"
            "Date: Mon, 1 Jun 2026 10:00:00 -0400\nContent-Type: text/plain; charset=utf-8\n\n"
            "Hi Pat,\n\nScope is below:\n\uf0b7 Support for physical network components (LAN, WLAN)\n"
            "a) 24/7 on-call availability\n- Troubleshooting and incident response\n\nThanks,\nOps\n"
        ).encode("utf-8")
    )
    texts = {a.raw_text for a in EmailParser().parse_artifact(project_id="p", artifact_id="m", path=p)}
    assert "Support for physical network components (LAN, WLAN)" in texts
    assert "24/7 on-call availability" in texts
    assert "Troubleshooting and incident response" in texts


def test_docx_typed_bullets_drop_markers(tmp_path: Path):
    p = tmp_path / "scope.docx"
    d = docx.Document()
    d.add_heading("1. Scope of Work", 1)
    d.add_paragraph("The contractor shall provide the following services:")
    for t in ("- 24/7 on-call availability for all sites", "\uf0b7 Troubleshooting and incident response at site",
              "a) Coordination with local stakeholders on site"):
        d.add_paragraph(t)
    d.save(p)
    atoms = DocxParser().parse_artifact(project_id="p", artifact_id="d", path=p)
    by_text = {a.raw_text: a for a in atoms}
    assert by_text["24/7 on-call availability for all sites"].value.get("list_marker") == "-"
    assert "Troubleshooting and incident response at site" in by_text
    # A typed ordinal in a SOW is the clause's reference and stays.
    assert "a) Coordination with local stakeholders on site" in by_text
    assert not any(t.startswith(("- ", "\uf0b7")) for t in by_text)
