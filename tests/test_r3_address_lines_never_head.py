"""An address box never becomes a section heading (010003).

"40 10TH AVE FL 4" -- the street line of a ship-to box, bold or all caps --
was promoted to a section heading and headed every atom after it (13 on the
live deal); the caption "SHIP TO:" and the box's company line did the same.
They are the box's content, in a .docx as in a PDF.
"""
from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from app.core.address_parse import is_address_block_line
from app.parsers.docx_parser import DocxParser
from app.parsers.orbitbrief_pdf import _looks_like_section_heading

BOX = ["SHIP TO:", "ACME CORPORATION", "40 10TH AVE FL 4", "NEW YORK, NY 10014-1066"]


@pytest.fixture(autouse=True)
def _no_llm(monkeypatch):
    monkeypatch.setenv("SOWSMITH_DISABLE_LLM", "1")


def test_address_lines_are_address_lines():
    for line in ("SHIP TO:", "Bill To", "SHIPPING ADDRESS", "40 10TH AVE FL 4", "PO BOX 1234",
                 "NEW YORK, NY 10014-1066"):
        assert is_address_block_line(line), line
        assert not _looks_like_section_heading(line), line
    for line in ("SCOPE OF WORK", "INSTALLATION", "Project Management", "3 PHASE POWER"):
        assert not is_address_block_line(line), line
    assert _looks_like_section_heading("SCOPE OF WORK")


def test_a_bold_ship_to_box_in_a_docx_heads_nothing(tmp_path: Path):
    d = Document()
    d.add_heading("Project Details", 1)
    for line in BOX:
        d.add_paragraph().add_run(line).bold = True
    d.add_paragraph("Install 12 displays on the north wall of the lobby.")
    path = tmp_path / "SOW.docx"
    d.save(path)
    atoms = DocxParser().parse_artifact("p", "a", path)
    atoms = atoms if isinstance(atoms, list) else atoms.atoms
    by = {a.raw_text: a for a in atoms}
    work = by["Install 12 displays on the north wall of the lobby."]
    assert work.source_refs[0].locator["section_path"] == ["Project Details"]
    for line in BOX:
        assert line in by, (line, list(by))
        assert "section_heading" not in by[line].review_flags, line
