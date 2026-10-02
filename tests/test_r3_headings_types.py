"""R3 (010003 f/h/i, 010087): heading-shaped lines open sections; amounts,
PO numbers and fee lead-ins are never chatter; Gantt rows and PO lines are
never commercial_total.

* "Invoicing Procedures:" (a Title Case colon label with prose, not bullets,
  beneath it) went to the labeler as a chatter line and opened no section.
* "PO #4500123456" / "Total: $1,622.00" are one-word labels with a figure;
  the short-fragment rule wanted two words, so both were chatter.
* "Services Fees hereunder are fixed fees totaling $1,622.00 and will be
  invoiced as follows:" is a list lead-in, so it became a structure (chatter)
  atom and the price was hidden.
* A Gantt sheet that also carries deal economics typed every schedule row
  commercial_total; a PO line typed commercial_total stayed one.
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path

import pytest
from docx import Document

from app.core.schemas import (
    ArtifactType, AtomType, AuthorityClass, EvidenceAtom, ReviewStatus, SourceRef,
)
from app.parsers.docx_parser import DocxParser


@pytest.fixture(autouse=True)
def _no_llm(monkeypatch):
    monkeypatch.setenv("SOWSMITH_DISABLE_LLM", "1")


def _atoms(path: Path):
    return DocxParser().parse_artifact("p", "a", path)


def _one(atoms, text):
    hits = [a for a in atoms if a.raw_text == text]
    assert hits, (text, [a.raw_text for a in atoms])
    return hits[0]


def _chatter(a) -> bool:
    return "chatter" in (a.review_flags or []) or bool((a.value or {}).get("chatter"))


def _path(a):
    return list(a.source_refs[0].locator.get("section_path") or [])


def test_amount_and_po_lines_are_not_chatter(tmp_path: Path) -> None:
    doc = Document()
    doc.add_heading("Pricing", level=1)
    doc.add_paragraph("PO #4500123456")
    doc.add_paragraph("Total: $1,622.00")
    doc.add_paragraph("Quote No. Q-1182")
    doc.add_paragraph("Page 3")
    doc.save(tmp_path / "SOW.docx")
    atoms = _atoms(tmp_path / "SOW.docx")
    for t in ("PO #4500123456", "Total: $1,622.00", "Quote No. Q-1182"):
        assert not _chatter(_one(atoms, t)), t
    # A one-word label with a bare number and no money / reference is still
    # too thin to be a fact.
    assert _chatter(_one(atoms, "Page 3"))


def test_fee_lead_in_with_an_amount_is_not_chatter(tmp_path: Path) -> None:
    fee = "Services Fees hereunder are fixed fees totaling $1,622.00 and will be invoiced as follows:"
    doc = Document()
    doc.add_heading("Pricing", level=1)
    doc.add_paragraph(fee)
    doc.add_paragraph("50% upon signature of this SOW", style="List Bullet")
    doc.add_paragraph("50% upon completion of the Services", style="List Bullet")
    doc.add_paragraph("The following will apply to all invoices:")
    doc.add_paragraph("Invoices are due net 30", style="List Bullet")
    doc.save(tmp_path / "SOW.docx")
    atoms = _atoms(tmp_path / "SOW.docx")
    a = _one(atoms, fee)
    assert not _chatter(a), a.review_flags
    assert getattr(a.atom_type, "value", a.atom_type) != "deal_metadata"
    # A lead-in with no figure is still the reject-able structure line.
    lead = _one(atoms, "The following will apply to all invoices:")
    assert _chatter(lead) and lead.value.get("rejected_by") == "list_lead_in"


def test_title_case_colon_label_heads_the_prose_beneath_it(tmp_path: Path) -> None:
    doc = Document()
    doc.add_heading("Out of Scope", level=1)
    doc.add_paragraph("Cabling of new drops is not included in this project", style="List Bullet")
    doc.add_paragraph("Invoicing Procedures:")
    doc.add_paragraph("PurTera will invoice the customer upon completion of the work.")
    doc.add_paragraph("Customer Name:")
    doc.add_paragraph("Acme Corporation of America")
    doc.save(tmp_path / "SOW.docx")
    atoms = _atoms(tmp_path / "SOW.docx")
    head = _one(atoms, "Invoicing Procedures:")
    assert head.value.get("rejected_by") == "section_heading", head.value
    body = _one(atoms, "PurTera will invoice the customer upon completion of the work.")
    assert _path(body)[-1] == "Invoicing Procedures", _path(body)
    # An invoicing section ends the exclusions section above it.
    assert getattr(body.atom_type, "value", body.atom_type) != "exclusion"
    # A form field label is not a heading.
    assert _one(atoms, "Customer Name:").value.get("rejected_by") != "section_heading"


def _mk(aid: str, atype: AtomType, text: str) -> EvidenceAtom:
    return EvidenceAtom(
        id=aid, project_id="p", artifact_id="x", atom_type=atype, raw_text=text,
        normalized_text=text.lower(), value={"text": text}, entity_keys=[],
        source_refs=[SourceRef(id="s" + aid, artifact_id="x", artifact_type=ArtifactType.xlsx,
                               filename="x.xlsx", extraction_method="x", parser_version="x", locator={})],
        authority_class=AuthorityClass.vendor_quote, confidence=0.8,
        review_status=ReviewStatus.needs_review, parser_version="x",
    )


def test_gantt_rows_and_po_lines_are_not_commercial_totals() -> None:
    from app.core.atom_type_sanity import retype_schedule_reference_rows

    rows = [
        ("g1", AtomType.commercial_total,
         "Task: Install displays | Start Date: 2026-08-17 00:00:00 | End Date: 2026-08-21 00:00:00 | Days: 5 | Cost: 4200"),
        ("g2", AtomType.pricing_assumption, "TV Delivery to NYC Office | 8/10/2026 | 8/12/2026 | 3 | 0"),
        ("g3", AtomType.commercial_total, "Project kickoff | Start: Aug 3, 2026 | $250.00"),
        ("po", AtomType.commercial_total, "PO #4500123456 - $18,207.48"),
        ("pr", AtomType.commercial_total, "PO # 4500123456: 18207.48 | Acme"),
        ("tt", AtomType.commercial_total, "Total | 9 | $4,450.00"),
        ("ft", AtomType.commercial_total, "Total Deal Revenue | $18,207.48"),
        ("pp", AtomType.commercial_total, "Fixed fee of $1,622.00 invoiced on 2026-08-21 and 2026-09-21."),
    ]
    atoms = [_mk(i, t, x) for i, t, x in rows]
    retype_schedule_reference_rows(atoms)
    got = {a.id: getattr(a.atom_type, "value", a.atom_type) for a in atoms}
    assert got == {
        "g1": "task", "g2": "task", "g3": "task",
        "po": "deal_metadata", "pr": "deal_metadata",
        "tt": "commercial_total", "ft": "commercial_total", "pp": "commercial_total",
    }, got


def test_gantt_sheet_with_deal_economics_after_compile(tmp_path: Path) -> None:
    openpyxl = pytest.importorskip("openpyxl")
    from app.core.compiler import compile_project

    d = tmp_path / "deal"
    d.mkdir()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Gantt"
    ws.append(["Oppty #", "010003"])
    ws.append(["Sales Rep", "Pat Doe"])
    ws.append(["Total Deal Revenue", 18207.48])
    ws.append(["Margin %", 0.31])
    ws.append(["PO # 4500123456", 18207.48])
    ws.append([])
    ws.append(["Task", "Start Date", "End Date", "Days", "Cost"])
    ws.append(["Project kickoff", dt.date(2026, 8, 3), dt.date(2026, 8, 3), 1, 250])
    ws.append(["TV Delivery to NYC Office", dt.date(2026, 8, 10), dt.date(2026, 8, 12), 3, 0])
    ws.append(["Install displays", dt.date(2026, 8, 17), dt.date(2026, 8, 21), 5, 4200])
    ws.append(["Total", None, None, 9, 4450])
    wb.save(d / "Gantt.xlsx")
    r = compile_project(d, project_id="p", allow_errors=True, use_cache=False)
    for name in ("Project kickoff", "TV Delivery", "Install displays", "PO # 4500123456"):
        hits = [a for a in r.atoms if name in (a.raw_text or "")]
        assert hits, (name, [a.raw_text for a in r.atoms])
        for a in hits:
            assert getattr(a.atom_type, "value", a.atom_type) not in ("commercial_total", "pricing_assumption"), (
                a.raw_text, a.atom_type)


def test_plain_title_case_lines_over_bullets_are_sibling_headings(tmp_path: Path) -> None:
    """010087: "Out of Scope" and "PurTera Responsibilities" left on Normal
    (not bold, no colon) over bullet lists. Unrecognised, the first was an
    exclusion atom of its own, the second a chatter line, and the PMO duties
    had no section. Each heads its own list; neither nests in the other."""
    doc = Document()
    doc.add_paragraph("Out of Scope")
    doc.add_paragraph("Cabling of new drops", style="List Bullet")
    doc.add_paragraph("PurTera Responsibilities")
    doc.add_paragraph("Develop schedule and coordinate resources", style="List Bullet")
    doc.save(tmp_path / "SOW.docx")
    atoms = _atoms(tmp_path / "SOW.docx")
    for h in ("Out of Scope", "PurTera Responsibilities"):
        a = _one(atoms, h)
        assert a.value.get("rejected_by") == "section_heading", (h, a.atom_type, a.value)
        assert _path(a) == [h], _path(a)
    ex = _one(atoms, "Cabling of new drops")
    assert _path(ex) == ["Out of Scope"]
    assert getattr(ex.atom_type, "value", ex.atom_type) == "exclusion"
    pmo = _one(atoms, "Develop schedule and coordinate resources")
    assert _path(pmo) == ["PurTera Responsibilities"], _path(pmo)
    assert getattr(pmo.atom_type, "value", pmo.atom_type) != "exclusion"
