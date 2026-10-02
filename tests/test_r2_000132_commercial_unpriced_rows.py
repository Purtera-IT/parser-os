"""A priced sheet keeps the rows that carry no price.

Found while reproducing deal 000132's Deal Kit v2: the commercial-sheet path
emitted priced rows and label boxes only. A "SELL RATES" title, an "After
Hours | After-hours work is billed at 150% of the standard rate" line, a
Notes block and a site table on the Deal Kit tab vanished without a trace.
Rows above the data block are now chatter atoms (rejected_by
commercial_sheet_scaffolding); unpriced rows inside it are pricing_assumption
atoms flagged unpriced_sheet_row; tick-box cells are split off.
"""
from __future__ import annotations

from pathlib import Path

import openpyxl

from app.parsers.xlsx_parser import XlsxParser


def _workbook(path: Path) -> list[list]:
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Deal Kit"
    ws.append(["Customer", "Acme Foods"]); ws.append(["OPPTY #", "000132"]); ws.append([])
    ws.append(["Site Name", "Address", "Service Type"])
    ws.append(["Delphos, OH", "123 Main St", "☐ Assessment ☐ Configuration ☒ Installation"])
    ws2 = wb.create_sheet("Sell Rates")
    ws2.append(["SELL RATES"])
    ws2.append(["Service", "Description", "Sell Rate", "Unit"])
    ws2.append(["Site Survey", "Site survey, 2 hour minimum", 96, "per hour"])
    ws2.append(["L1 Technician", "Onsite technician, 2 hour minimum", 65, "per hour"])
    ws2.append(["After Hours", "After-hours work is billed at 150% of the standard rate", None, None])
    ws2.append([])
    ws2.append(["Notes:"])
    ws2.append(["Travel billed at cost."])
    wb.save(path)
    return [[c for c in r if c not in (None, "")] for s in wb.worksheets for r in s.iter_rows(values_only=True)]


def test_every_cell_of_a_priced_workbook_is_in_an_atom(tmp_path: Path):
    rows = _workbook(tmp_path / "Deal Kit v2.xlsx")
    atoms = XlsxParser().parse(tmp_path / "Deal Kit v2.xlsx")
    blob = "\n".join(a.raw_text for a in atoms)
    for row in rows:
        for cell in row:
            assert str(cell) in blob, (cell, blob)


def test_unpriced_rows_are_typed_for_review(tmp_path: Path):
    _workbook(tmp_path / "Deal Kit v2.xlsx")
    by = {a.raw_text: a for a in XlsxParser().parse(tmp_path / "Deal Kit v2.xlsx")}
    title = by["SELL RATES"]
    assert "chatter" in title.review_flags and title.value["rejected_by"] == "commercial_sheet_scaffolding"
    ah = by["After Hours | After-hours work is billed at 150% of the standard rate"]
    assert ah.atom_type.value == "pricing_assumption" and "unpriced_sheet_row" in ah.review_flags
    assert by["Travel billed at cost."].atom_type.value == "pricing_assumption"
    # the tick boxes are their own atom, not part of the site row
    assert "Delphos, OH | 123 Main St" in by
    box = by["☐ Assessment ☐ Configuration ☒ Installation"]
    assert box.value["selected"] == ["Installation"] and box.value["subject"] == "Delphos, OH"
    # priced rows unchanged
    assert any(t.startswith("Sell Rates | Service: Site Survey") for t in by)
