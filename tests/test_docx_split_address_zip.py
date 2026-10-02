"""A site address whose ZIP sits in its own form field becomes ONE physical_site.

Deal 010353's SOW typed "15733 US-224, Findlay, OH" and put the ZIP 45840 in a
separate form field, so street/city/state and ZIP were separate runs / fields /
cells. The address never became an atom: a ``w:fldSimple`` field's runs were
invisible to python-docx, a paragraph form ("Site Address: ... Zip: ...") was
never read as a site, and a table block dropped the city or the ZIP from the
site's text. Every layout below must yield one physical_site carrying the full
address.
"""

from __future__ import annotations

import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from app.parsers.docx_parser import DocxParser

W = nsdecls("w")
FULL = "15733 US-224, Findlay, OH 45840"


def _r(t: str) -> str:
    return f'<w:r><w:t xml:space="preserve">{t}</w:t></w:r>'


def _formtext(v: str) -> str:
    return (
        '<w:r><w:fldChar w:fldCharType="begin"><w:ffData><w:name w:val="Zip"/><w:textInput/>'
        '</w:ffData></w:fldChar></w:r><w:r><w:instrText xml:space="preserve"> FORMTEXT </w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r>' + _r(v) + '<w:r><w:fldChar w:fldCharType="end"/></w:r>'
    )


def _para(xml: str):
    return lambda doc: doc.element.body[-1].addprevious(parse_xml(f"<w:p {W}>{xml}</w:p>"))


def _table(*rows: list[str]):
    def build(doc):
        t = doc.add_table(rows=len(rows), cols=len(rows[0]))
        for i, row in enumerate(rows):
            for j, v in enumerate(row):
                t.cell(i, j).text = v
    return build


LAYOUTS = {
    "content_control": _para(
        _r("Site Address: 15733 US-224, Findlay, OH ")
        + f"<w:sdt><w:sdtPr/><w:sdtContent>{_r('45840')}</w:sdtContent></w:sdt>"
    ),
    "fld_simple": _para(_r("Site Address: 15733 US-224, Findlay, OH ")
                        + f'<w:fldSimple w:instr=" FORMTEXT ">{_r("45840")}</w:fldSimple>'),
    "smart_tag": _para(
        _r("Site Address: 15733 US-224, ")
        + f'<w:smartTag w:uri="urn:schemas-microsoft-com:office:smarttags" w:element="City">{_r("Findlay")}</w:smartTag>'
        + _r(", OH ")
        + f'<w:smartTag w:uri="urn:schemas-microsoft-com:office:smarttags" w:element="PostalCode">{_r("45840")}</w:smartTag>'
    ),
    "legacy_form_fields": _para(_r("Site Address: ") + _formtext("15733 US-224, Findlay, OH")
                                + _r(" Zip: ") + _formtext("45840")),
    "tab_stops": _para(_r("Site Address: 15733 US-224, Findlay, OH") + "<w:r><w:tab/></w:r>"
                       + _r("Zip:") + "<w:r><w:tab/></w:r>" + _r("45840")),
    "next_paragraph": lambda doc: (doc.add_paragraph("Site Address: 15733 US-224, Findlay, OH"),
                                   doc.add_paragraph("Zip: 45840")),
    "label_cells": _table(["Site Address:", "15733 US-224, Findlay, OH", "Zip:", "45840"]),
    "split_cells": _table(["Site Address", "15733 US-224", "Findlay, OH", "45840"]),
    "zip_row": _table(["Site Address", "15733 US-224, Findlay, OH"], ["Zip Code", "45840"]),
    "header_row": _table(["Street", "City", "State", "Zip"], ["15733 US-224", "Findlay", "OH", "45840"]),
    "labelled_cells": _table(["Site Address: 15733 US-224, Findlay, OH", "Zip: 45840"]),
}


@pytest.mark.parametrize("layout", sorted(LAYOUTS))
def test_split_zip_joins_into_one_physical_site(tmp_path, layout) -> None:
    doc = Document()
    doc.add_heading("Statement of Work", level=1)
    doc.add_paragraph("Contractor shall install network equipment at the customer site listed below.")
    LAYOUTS[layout](doc)
    doc.add_paragraph("All work to be performed during normal business hours.")
    path = tmp_path / "SOW.docx"
    doc.save(path)

    out = DocxParser().parse_artifact("p", "a", path)
    atoms = out if isinstance(out, list) else out.atoms
    sites = [a for a in atoms if a.atom_type.value == "physical_site"]
    assert len(sites) == 1, [a.raw_text for a in sites]
    site = sites[0]
    assert FULL in site.raw_text
    assert site.value["full_address"] == FULL
    assert (site.value["address"], site.value["city"], site.value["state"], site.value["zip"]) == (
        "15733 US-224", "Findlay", "OH", "45840")


def test_vendor_office_line_is_not_a_site(tmp_path) -> None:
    doc = Document()
    doc.add_paragraph("Address: 11720 Amber Park Dr, Alpharetta, GA\tZip: 30009")
    path = tmp_path / "SOW.docx"
    doc.save(path)
    out = DocxParser().parse_artifact("p", "a", path)
    atoms = out if isinstance(out, list) else out.atoms
    assert not [a for a in atoms if a.atom_type.value == "physical_site"]
