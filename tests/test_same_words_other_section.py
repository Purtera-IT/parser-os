"""The same sentence under two sections of one document is two lines.

A SOW states a sentence once in its scope section and again inside a field of
a later section (there, the reason a field gives). The intra-document
collapse folded the second into the first, so the field lost its reason.
Identical text under a different parent is its own fact, never a repeat.
The same sentence twice under ONE section still folds.

Synthetic document only.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document

from app.core.entity_resolution import collapse_duplicate_atoms
from app.parsers.docx_parser import DocxParser

SENTENCE = "All work under this agreement will be performed at the customer premises."


def _docx(tmp: Path, *, second_heading: str) -> Path:
    doc = Document()
    doc.add_heading("Work Proposal", level=1)
    doc.add_heading("Scope", level=2)
    doc.add_paragraph(
        "Services will be provided on a time and materials basis for the whole term. "
        "Provider will perform weekly visits at each supported location, as requested and scheduled "
        "by the customer, with each visit lasting one full working day. " + SENTENCE
    )
    doc.add_heading(second_heading, level=2)
    p = doc.add_paragraph()
    p.add_run("Off-site Support:").bold = True
    p.add_run(" Not included. " + SENTENCE.rstrip("."))
    p.add_run().add_break()
    p.add_run("Scheduling:").bold = True
    p.add_run(" Subject to technician availability")
    path = tmp / "proposal.docx"
    doc.save(path)
    return path


def _copies(atoms):
    return [a for a in atoms if a.raw_text.strip().rstrip(".") == SENTENCE.rstrip(".")]


def _parse(path: Path):
    return DocxParser().parse_artifact_full(project_id="p", artifact_id="art_doc", path=path).atoms


def test_same_sentence_under_another_section_is_kept(tmp_path):
    atoms = _parse(_docx(tmp_path, second_heading="Service Model"))
    copies = _copies(atoms)
    assert len(copies) == 2, "fixture: the sentence once per section"
    sections = {tuple(a.source_refs[0].locator.get("section_path") or ()) for a in copies}
    assert len(sections) == 2, "fixture: the two copies sit under different sections"
    out = collapse_duplicate_atoms(atoms)
    assert len(_copies(out)) == 2


def test_same_sentence_twice_under_one_section_still_folds(tmp_path):
    atoms = _parse(_docx(tmp_path, second_heading="Service Model"))
    first = _copies(atoms)[0]
    twin = first.model_copy(deep=True)
    twin.id = "atm_twin_same_section"
    alone = collapse_duplicate_atoms([a.model_copy(deep=True) for a in atoms])
    out = collapse_duplicate_atoms(atoms + [twin])
    assert len(_copies(out)) == len(_copies(alone))
