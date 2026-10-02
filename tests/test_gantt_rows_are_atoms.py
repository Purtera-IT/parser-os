"""A Gantt task row is one atom, read left to right, under its phase.

Deal 010246's Deal Kit Gantt came out as atoms that glued task rows
together column by column: "Cable pulls: Rack and stack | 120: 40 |
<start>: <start> | <end>: <end>". The phase rows ("Planning", "Install",
"Closeout") cut the table into inline-titled segments, each segment had no
header of its own, and its first task row -- a label, a number and two
dates -- passed for a header, so every other task in the phase was bound to
it. Where the phase sits in a column merged down its tasks, only the first
task carried it.

Now: the phase rows stay inside the table as context for the rows under
them, a data row is never taken for a header, a merged phase cell is on
every row it spans, and each task is its own atom.
"""

from __future__ import annotations

import datetime as dt

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.core.schemas import AtomType
from app.parsers.xlsx_blocks import sheet_blocks
from app.parsers.xlsx_parser import XlsxParser

_BANNER = PatternFill("solid", fgColor="FF1F3864")
_BOX = PatternFill("solid", fgColor="FFDDEBF7")
_BAR = PatternFill("solid", fgColor="FF92D050")
_WEEKS = [dt.datetime(2025, 3, 3) + dt.timedelta(days=7 * i) for i in range(6)]
_PHASES = [
    ("Planning", [("Kickoff meeting", 4, 0, 0), ("Site survey walk", 16, 0, 1),
                  ("Design review", 8, 1, 1), ("Permit filing", 6, 1, 2)]),
    ("Install", [("Cable pulls", 120, 2, 3), ("Rack and stack", 40, 3, 3),
                 ("Terminate and test", 60, 3, 4), ("Label ports", 8, 4, 4)]),
    ("Closeout", [("As-built drawings", 12, 5, 5), ("Customer walkthrough", 4, 5, 5),
                  ("Punch list", 6, 5, 5), ("Final signoff", 2, 5, 5)]),
]
_TASKS = {t: (ph, h) for ph, ts in _PHASES for t, h, _s, _e in ts}


def _gantt(path, *, phase_layout: str, header: bool = True) -> None:
    """A Deal Kit-shaped Gantt tab.

    ``phase_layout``: "merged" puts the phase in column A merged down its
    tasks; "rows" puts each phase on its own row above its tasks.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Gantt"
    r = 1

    def put(vals, fill=None, bold=False):
        nonlocal r
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            if fill is not None:
                cell.fill = fill
            if bold:
                cell.font = Font(bold=True)
        r += 1

    put(["OxBlue Pumphouse - Project Gantt"], _BANNER, bold=True)
    r += 1
    for k, v in (("Customer", "OxBlue"), ("OPPTY #", "010246"), ("Project Manager", "PK")):
        put([k, v], _BOX)
    r += 1
    lead = ["Phase"] if phase_layout == "merged" else []
    off = len(lead)
    if header:
        put(lead + ["Task", "Hours", "Start", "End"] + _WEEKS, _BANNER, bold=True)
    for phase, tasks in _PHASES:
        if phase_layout == "rows":
            put([phase], _BOX, bold=True)
        first = r
        for i, (task, hours, s, e) in enumerate(tasks):
            pc = [phase if i == 0 else None] if phase_layout == "merged" else []
            put(pc + [task, hours, _WEEKS[s], _WEEKS[e] + dt.timedelta(days=4)])
            for k in range(s, e + 1):   # the Gantt bar: filled week cells
                ws.cell(row=r - 1, column=off + 5 + k).fill = _BAR
        if phase_layout == "merged":
            ws.merge_cells(start_row=first, start_column=1, end_row=r - 1, end_column=1)
    wb.save(path)


def _rows(path):
    out = XlsxParser().parse_artifact("proj", "art", path)
    atoms = out if isinstance(out, list) else out.atoms
    return [a for a in atoms if a.atom_type == AtomType.raw_table_row
            and any(t in a.raw_text for t in _TASKS)]


def _assert_one_atom_per_task(rows, *, header: bool = True) -> None:
    by_task: dict[str, list] = {t: [] for t in _TASKS}
    for a in rows:
        named = [t for t in _TASKS if t in a.raw_text]
        # Never two task rows in one atom.
        assert len(named) == 1, a.raw_text
        by_task[named[0]].append(a)
    for task, hits in by_task.items():
        assert len(hits) == 1, (task, [h.raw_text for h in hits])
        text = hits[0].raw_text
        phase, hours = _TASKS[task]
        # The phase rides on the row as context, ahead of the task.
        assert phase in text, text
        assert text.index(phase) < text.index(task)
        # Left to right: task, then its hours, then start before end.
        assert text.index(task) < text.index(str(hours))
        if header:
            assert f"Task: {task}" in text, text
            assert f"Hours: {hours}" in text, text
            assert text.index("Start:") < text.index("End:")


def test_phase_rows_divide_one_table(tmp_path) -> None:
    path = tmp_path / "Deal_Kit.xlsx"
    _gantt(path, phase_layout="rows")
    rows = _rows(path)
    _assert_one_atom_per_task(rows)
    # The phase extends the row's path as well as leading its text.
    cable = next(a for a in rows if "Cable pulls" in a.raw_text)
    assert cable.raw_text.startswith("Install | Task: Cable pulls | Hours: 120")
    assert cable.value["section"] == "Install"
    assert cable.source_refs[0].locator["section_path"][-1] == "Install"
    # Each row cites its own worksheet row.
    assert len({a.source_refs[0].locator["row"] for a in rows}) == len(_TASKS)


def test_merged_phase_column_is_on_every_task(tmp_path) -> None:
    path = tmp_path / "Deal_Kit.xlsx"
    _gantt(path, phase_layout="merged")
    rows = _rows(path)
    _assert_one_atom_per_task(rows)
    rack = next(a for a in rows if "Rack and stack" in a.raw_text)
    assert rack.raw_text.startswith("Phase: Install | Task: Rack and stack | Hours: 40")


def test_no_header_row_never_promotes_a_task(tmp_path) -> None:
    # With no header row at all, the first task must not become one.
    for layout in ("rows", "merged"):
        path = tmp_path / f"Deal_Kit_{layout}.xlsx"
        _gantt(path, phase_layout=layout, header=False)
        out = XlsxParser().parse_artifact("proj", "art", path)
        atoms = out if isinstance(out, list) else out.atoms
        for a in atoms:
            named = [t for t in _TASKS if t in (a.raw_text or "")]
            assert len(named) <= 1, (layout, a.raw_text)
        for task in _TASKS:
            assert any(task in (a.raw_text or "") for a in atoms), (layout, task)


def test_stacked_label_value_boxes_still_split() -> None:
    # The inline-title split exists for a right rail of separate label/value
    # boxes stacked with no blank row between them; a table above them must
    # not swallow them.
    rows = [
        ["Task Category", "Labor Hours", "Rate", "Total"],
        ["Install", 40, 85, 3400],
        ["Programming", 10, 105, 1050],
        ["Key Unit Metrics", None, None, None],
        ["Cost per drop", 125, None, None],
        ["Hours per drop", 1.5, None, None],
        ["Gross Margin Deal Kit", None, None, None],
        ["Margin", 0.32, None, None],
        ["Net", 1200, None, None],
    ]
    blocks = sheet_blocks(rows)
    titles = [b.get("title") for b in blocks]
    assert "Key Unit Metrics" in titles
    assert "Gross Margin Deal Kit" in titles
    table = next(b for b in blocks if b["kind"] == "table")
    assert table["header"][:2] == ["Task Category", "Labor Hours"]
    assert [r[0] for r in table["rows"]] == ["Install", "Programming"]
