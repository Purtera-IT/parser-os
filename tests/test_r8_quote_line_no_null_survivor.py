"""quote_line_head never removes a line with no survivor (000132, round 8).

Two drafts of one SOW carried a byte-identical "<duties>:" list (a Word
numbered list inside nested content controls: lead-in at ilvl 0, four items at
ilvl 2, the last item with no paragraph-mark rPr). The parser typed one
draft's items ``task`` and the other's ``scope_item``. In the ``task`` draft
quote_line_head read the last item's wording as a PMO/admin step and dropped
it -- a ledger entry with survivor null -- while the ``scope_item`` draft kept
it. The same line must get the same structural outcome whatever its type:
the head may keep it off the quote, but the atom stands. Synthetic text, real
XML shape.
"""
from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from app.core.schemas import ArtifactType, AtomType, AuthorityClass, EvidenceAtom, ReviewStatus, SourceRef

W = nsdecls("w")
LEAD = "Admin Duties:"
ITEMS = [
    "Develop schedule from site windows",
    "Deploy crew(s)",
    "Validate handoff items with Customer",
    "Complete billing paperwork at closeout",  # the PMO/admin wording
]
_MARK = '<w:rPr><w:rFonts w:asciiTheme="majorHAnsi" w:hAnsiTheme="majorHAnsi" w:cstheme="majorHAnsi"/></w:rPr>'
_RUN = ('<w:rPr><w:rFonts w:asciiTheme="majorHAnsi" w:hAnsiTheme="majorHAnsi" w:cstheme="majorHAnsi"/>'
        '<w:color w:val="000000"/></w:rPr>')


@pytest.fixture(autouse=True)
def _no_llm(monkeypatch):
    monkeypatch.setenv("SOWSMITH_DISABLE_LLM", "1")


def _p(text: str, ilvl: int, mark_rpr: bool = True) -> str:
    ppr = (f'<w:pPr><w:pStyle w:val="BodyText"/><w:numPr><w:ilvl w:val="{ilvl}"/><w:numId w:val="13"/></w:numPr>'
           '<w:autoSpaceDE w:val="0"/><w:autoSpaceDN w:val="0"/>'
           '<w:spacing w:after="0" w:line="240" w:lineRule="auto"/><w:contextualSpacing/>'
           + (_MARK if mark_rpr else "") + "</w:pPr>")
    return f'<w:p>{ppr}<w:r>{_RUN}<w:t xml:space="preserve">{text}</w:t></w:r></w:p>'


def _sow(path: Path) -> None:
    d = Document()
    num = d.part.numbering_part.element
    lvls = "".join(f'<w:lvl w:ilvl="{i}"><w:start w:val="1"/><w:numFmt w:val="{f}"/><w:lvlText w:val="{t}"/></w:lvl>'
                   for i, (f, t) in enumerate([("decimal", "%1."), ("lowerLetter", "%2."), ("bullet", "o")]))
    num.append(parse_xml(f'<w:abstractNum {W} w:abstractNumId="77">{lvls}</w:abstractNum>'))
    num.append(parse_xml(f'<w:num {W} w:numId="13"><w:abstractNumId w:val="77"/></w:num>'))
    d.add_heading("Services Proposal", 1)
    d.add_heading("Project Scope", 2)
    d.add_heading("Provider Responsibilities", 3)
    d.add_paragraph("Provider is responsible for the following:")
    d.add_paragraph("Install and configure network switches at each site.", style="List Bullet")
    body = _p(LEAD, 0) + "".join(_p(t, 2, mark_rpr=i < len(ITEMS) - 1) for i, t in enumerate(ITEMS))
    d.element.body.append(parse_xml(
        f"<w:sdt {W}><w:sdtPr/><w:sdtContent><w:sdt><w:sdtPr/><w:sdtContent>{body}"
        "</w:sdtContent></w:sdt></w:sdtContent></w:sdt>"))
    d.add_heading("Customer Responsibilities", 3)
    d.add_paragraph("Provide site access during normal business hours.", style="List Bullet")
    d.save(path)


def _compile(tmp_path: Path, monkeypatch, item_type: AtomType | None):
    """Compile the SOW; ``item_type`` stands in for the parser's type guess."""
    from app.core.compiler import compile_project
    from app.parsers.docx_parser import DocxParser

    if item_type is not None:
        orig = DocxParser.parse_artifact_full

        def _typed(self, *a, **k):
            out = orig(self, *a, **k)
            for atom in out.atoms:
                if (atom.raw_text or "").strip() in ITEMS:
                    atom.atom_type = item_type
            return out

        monkeypatch.setattr(DocxParser, "parse_artifact_full", _typed)
    d = tmp_path / "deal"
    d.mkdir()
    _sow(d / "SOW v1.docx")
    return compile_project(d, project_id="p", allow_errors=True, use_cache=False)


@pytest.mark.parametrize("item_type", [AtomType.task, AtomType.scope_item])
def test_every_list_item_stands_whatever_its_type(tmp_path, monkeypatch, item_type) -> None:
    r = _compile(tmp_path, monkeypatch, item_type)
    for text in [LEAD] + ITEMS:
        kept = [a for a in r.atoms if (a.raw_text or "").strip() == text]
        assert len(kept) == 1, (item_type, text, [a.raw_text for a in r.atoms])
    lost = [a for a in r.suppressed_atoms if (a.raw_text or "").strip() in ITEMS]
    assert not lost, [(a.raw_text, (a.value or {}).get("_suppression")) for a in lost]
    # Type is the parser's call, untouched here.
    for text in ITEMS:
        a = next(a for a in r.atoms if (a.raw_text or "").strip() == text)
        assert a.atom_type == item_type


def test_a_pmo_task_stays_an_atom_off_the_quote(tmp_path, monkeypatch) -> None:
    r = _compile(tmp_path, monkeypatch, AtomType.task)
    a = next(a for a in r.atoms if (a.raw_text or "").strip() == ITEMS[-1])
    assert a.value.get("is_quote_line") is False
    assert a.value["quote_line"]["source"] == "pmo_filtered"
    assert not any(str(f).startswith("suppressed:") for f in a.review_flags)
    from app.core.task_tier_classifier import is_quote_line_task_atom

    assert not is_quote_line_task_atom(a)


def _task(text: str, n: int) -> EvidenceAtom:
    ref = SourceRef(id=f"src_{n}", artifact_id="art", artifact_type=ArtifactType.docx, filename="SOW.docx",
                    locator={"paragraph_index": n}, extraction_method="test", parser_version="t")
    return EvidenceAtom(
        id=f"atm_{n:04d}", project_id="p", artifact_id="art", atom_type=AtomType.task,
        raw_text=text, normalized_text=text.lower(),
        value={"text": text, "task_tier": "parent", "is_quote_line": True}, entity_keys=[],
        source_refs=[ref], authority_class=AuthorityClass.machine_extractor, confidence=0.9,
        review_flags=[], review_status=ReviewStatus.auto_accepted, parser_version="test",
    )


def test_head_removes_no_atom_without_an_umbrella_that_stands() -> None:
    from app.core.quote_line_head import consolidate_quote_line_tasks
    from app.core.suppression_ledger import take_folds

    atoms = [_task(t, i) for i, t in enumerate(ITEMS)]
    take_folds()
    out, _ = consolidate_quote_line_tasks(list(atoms), project_id="p")
    folds = take_folds()
    out_ids = {id(a) for a in out}
    for a in atoms:
        if id(a) in out_ids:
            continue
        # Anything removed was folded, into an umbrella that is in the output.
        assert id(a) in folds, a.raw_text
        assert id(folds[id(a)][1]) in out_ids, a.raw_text
    assert any(a.raw_text == ITEMS[-1] for a in out)
