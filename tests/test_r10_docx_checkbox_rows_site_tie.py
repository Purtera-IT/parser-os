"""A site list whose rows carry several checkbox cells each.

    Site Locations                            (a heading over the table)
    | Location(s)     | Service(s)  (spans three columns)                 |
    | Springfield, IL | ☐ Alpha  | ☐ Gamma  | ☑ Epsilon |
    |                 | ☐ Beta   | ☐ Delta  | ☑ Zeta    |
    | Salem, OR       | (the same three cells, word for word)            |

The parser emitted one atom per checkbox cell ("☐ Alpha ☐ Beta", three a row,
every row alike), each tied to its site only by structure. Now each grid row
reads as two atoms: the site cell, exactly the atom it always was (same id,
text and value, so the site it names and a label saved on it stay put), then
ONE atom for the row's boxes: the site first, then only the ticked options
("Springfield, IL: Epsilon, Zeta"). Unticked options are never atoms and never
in the text; they stay on the row atom as structure, with the column header
each came from. The heading over the grid and the grid's header row
("Location(s) | Service(s)") are pointer atoms every row atom names
(heading_atom_id, header_atom_id).

R12 (000132): folding the site cell INTO the row atom ("Springfield, IL:
Epsilon, Zeta" as the site) lost the table's own site line -- a label keyed on
"Springfield, IL" moved to another line with the same words -- and the row
stood in for the site in dedup, where a model-typed copy of it knocked the
email's quoted site lines out of the site list.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from app.core.ids import stable_id
from app.core.orbitbrief_envelope import _tie_checkbox_cells_to_site
from app.parsers import checkbox_cells
from app.parsers.docx_parser import DocxParser

SITES = ["Springfield, IL", "Salem, OR", "Dover, NH"]
COLS = [
    [(0, "Alpha"), (0, "Beta")],
    [(0, "Gamma"), (0, "Delta")],
    [(1, "Epsilon"), (1, "Zeta"), (0, "Eta")],
]
NONE_ROW = "Concord, NH"  # a row with no box ticked


def _r(t: str) -> str:
    return f"<w:r><w:t xml:space='preserve'>{t}</w:t></w:r>"


def _tc(paras: list[str], span: int | None = None) -> str:
    s = f"<w:gridSpan w:val='{span}'/>" if span else ""
    return f"<w:tc><w:tcPr><w:tcW w:w='2000' w:type='dxa'/>{s}</w:tcPr>" + "".join(
        f"<w:p>{p}</w:p>" for p in paras) + "</w:tc>"


def _row(site: str, cols) -> str:
    cells = "".join(_tc([_r(("☑" if c else "☐")) + _r(" " + l) for c, l in col]) for col in cols)
    return "<w:tr>" + _tc([_r(site)]) + cells + "</w:tr>"


def _doc(path: Path) -> Path:
    rows = ["<w:tr>" + _tc([_r("Location(s)")]) + _tc([_r("Service(s)")], 3) + "</w:tr>"]
    rows += [_row(site, COLS) for site in SITES]
    rows.append(_row(NONE_ROW, [[(0, l) for _c, l in col] for col in COLS]))
    tbl = (f"<w:tbl {nsdecls('w')}><w:tblPr><w:tblW w:w='8000' w:type='dxa'/></w:tblPr>"
           "<w:tblGrid>" + "<w:gridCol w:w='2000'/>" * 4 + "</w:tblGrid>" + "".join(rows) + "</w:tbl>")
    d = Document()
    d.add_heading("Site Locations", 2)
    d.element.body.insert(len(d.element.body) - 1, parse_xml(tbl))
    d.add_paragraph("Services follow the table above.")
    # An ordinary table beside it: untouched.
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Item", "Owner"
    t.cell(1, 0).text, t.cell(1, 1).text = "Rack the switch", "Provider"
    d.save(path)
    return path


def _parse(tmp_path: Path):
    out = DocxParser().parse_artifact("p", "art_x", _doc(tmp_path / "sow.docx"))
    return out if isinstance(out, list) else out.atoms


def _loc(a) -> dict:
    return a.source_refs[0].locator


def _site_atom(atoms, row: int, table: int = 0):
    return next(a for a in atoms
                if _loc(a).get("extraction") == "docx_table_row_v1"
                and _loc(a).get("table_index") == table and _loc(a).get("row") == row)


def _row_atom(atoms, row: int, table: int = 0):
    return next(a for a in atoms
                if a.source_refs[0].extraction_method == "docx_checkbox_grid_row_v1"
                and _loc(a).get("table_index") == table and _loc(a).get("row") == row)


TICKED = [l for col in COLS for c, l in col if c]
UNTICKED = [l for col in COLS for c, l in col if not c]


def test_site_cell_is_its_own_atom_then_one_row_atom(tmp_path):
    atoms = _parse(tmp_path)
    # No per-cell checkbox atom, and no box glyph in any atom's text.
    assert not [a for a in atoms if "checkbox_cell" in (a.review_flags or [])]
    assert not [a for a in atoms if "☐" in a.raw_text or "☑" in a.raw_text]
    for row, site in enumerate(SITES, start=1):
        s, a = _site_atom(atoms, row), _row_atom(atoms, row)
        # The site cell's atom, as before the boxes were split off: its own
        # words, its own id, its own cell.
        assert s.raw_text == site
        assert s.id == stable_id("atm", "art_x", "docx_row", 0, row, site)
        assert _loc(s)["cells"] == [0] and "checkbox_row" not in s.value
        # Then the row's ONE checkbox atom, right after it.
        assert atoms.index(a) == atoms.index(s) + 1
        assert a.raw_text == f"{site}: {', '.join(TICKED)}"
        assert a.value["row_atom_id"] == s.id
        assert a.atom_type.value == "site_attribute"
        assert not [x for x in atoms if _loc(x).get("table_index") == 0 and _loc(x).get("row") == row
                    and x.atom_type.value != "raw_table_row" and x not in (a, s)]
    assert _row_atom(atoms, len(SITES) + 1).raw_text == f"{NONE_ROW}: {checkbox_cells.NONE_SELECTED}"


def test_row_atom_names_every_cell_and_keeps_unticked_as_structure(tmp_path):
    atoms = _parse(tmp_path)
    a = _row_atom(atoms, 2)
    loc, v = _loc(a), a.value
    assert loc["table_index"] == 0 and loc["row"] == 2
    assert loc["cells"] == [0, 1, 2, 3] and loc["row_key"] == "t0 r2"
    assert v["checkbox_row"] and v["site"] == v["subject"] == SITES[1]
    assert v["site_cell"] == 0 and v["site_column"] == "Location(s)"
    assert v["selected"] == TICKED and v["not_selected"] == UNTICKED
    assert v["options"] == [
        {"label": l, "checked": bool(c), "column": "Service(s)", "cell": i}
        for i, col in enumerate(COLS, start=1) for c, l in col
    ]
    assert v["label_column"] == "Location(s)" and v["option_columns"] == ["Service(s)"]
    assert v["heading"] == "Site Locations"
    # The column names stay on the site cell's atom, as they always were.
    assert _site_atom(atoms, 2).value["columns"][0] == "Location(s)"
    none = _row_atom(atoms, len(SITES) + 1).value
    assert none["selected"] == []
    assert none["not_selected"] == [l for col in COLS for _c, l in col]


def test_grid_heading_and_header_row_are_pointer_atoms_the_rows_name(tmp_path):
    atoms = _parse(tmp_path)
    heads = [a for a in atoms if a.raw_text == "Site Locations"]
    assert len(heads) == 1, [a.raw_text for a in atoms][:5]
    h = heads[0]
    # Kept (not pre-suppressed): a heading that leads a table leads child lines.
    assert _loc(h).get("block_kind") == "heading"
    assert not [f for f in h.review_flags if f.startswith("suppressed:")]
    # The header row names the columns: one pointer atom, merged cells once,
    # read after the heading and before the first row.
    hdrs = [a for a in atoms if a.raw_text == "Location(s) | Service(s)"]
    assert len(hdrs) == 1, [a.raw_text for a in atoms][:8]
    hr = hdrs[0]
    assert _loc(hr)["block_kind"] == "table_header"
    assert _loc(hr)["table_index"] == 0 and _loc(hr)["row"] == 0 and _loc(hr)["cells"] == [0, 1]
    assert hr.value["kind"] == "table_header" and hr.value["structure"]
    assert hr.value["heading_atom_id"] == h.id
    assert atoms.index(h) < atoms.index(hr) < atoms.index(_site_atom(atoms, 1))
    for row in range(1, len(SITES) + 2):
        v = _row_atom(atoms, row).value
        assert v["header_atom_id"] == hr.id and v["heading_atom_id"] == h.id
    # An ordinary table's header row is a pointer too (r13), never a row.
    items = [a for a in atoms if a.raw_text.startswith("Item")]
    assert items and all(_loc(a)["block_kind"] == "table_header" for a in items)


def test_header_and_heading_are_held_out_of_dedup_like_headings():
    from app.core.compiler import _is_heading_atom
    from app.core.schemas import ArtifactType, AtomType, AuthorityClass, EvidenceAtom, ReviewStatus, SourceRef

    def atom(kind):
        return EvidenceAtom(
            id="a", project_id="p", artifact_id="x", atom_type=AtomType.deal_metadata,
            raw_text="Location(s) | Service(s)", normalized_text="location(s) | service(s)",
            value={"structure": True}, entity_keys=[],
            source_refs=[SourceRef(id="s", artifact_id="x", artifact_type=ArtifactType.docx,
                                   filename="f.docx", locator={"block_kind": kind},
                                   extraction_method="t", parser_version="t")],
            authority_class=AuthorityClass.contractual_scope, confidence=0.1,
            review_status=ReviewStatus.needs_review, review_flags=[], parser_version="t",
        )

    assert _is_heading_atom(atom("heading")) and _is_heading_atom(atom("table_header"))
    assert not _is_heading_atom(atom("table_row"))


def test_heading_over_a_table_leads_it_whatever_its_index():
    """`_headings_with_children` looked a TABLE's index up as a paragraph
    index: a table right under a heading read as a sibling heading whenever
    paragraph N (N = the table's index) was a heading, and the heading was
    suppressed."""
    from docx import Document as _D

    d = _D()
    d.add_heading("Scope", 1)          # p0
    d.add_heading("Locations", 2)      # p1, then table 0
    d.add_table(rows=1, cols=1).cell(0, 0).text = "Springfield, IL"
    d.add_heading("Fees", 2)           # p2, then table 1
    d.add_table(rows=1, cols=1).cell(0, 0).text = "$100"
    import tempfile

    with tempfile.TemporaryDirectory() as td:
        path = Path(td) / "h.docx"
        d.save(path)
        out = DocxParser().parse_artifact("p", "art_h", path)
    atoms = out if isinstance(out, list) else out.atoms
    for text in ("Locations", "Fees"):
        h = next(a for a in atoms if a.raw_text == text)
        assert _loc(h).get("block_kind") == "heading", text


def test_drop_unticked_fully_is_one_switch(tmp_path, monkeypatch):
    monkeypatch.setattr(checkbox_cells, "DROP_UNTICKED_FULLY", True)
    atoms = _parse(tmp_path)
    v = _row_atom(atoms, 1).value
    assert v["selected"] == TICKED and "not_selected" not in v
    assert all(o["checked"] for o in v["options"])
    assert not [t for t in _site_atom(atoms, 1).value["cells"].values() if "☐" in t or "☑" in t]
    assert _row_atom(atoms, 1).raw_text == f"{SITES[0]}: {', '.join(TICKED)}"


def test_ordinary_table_rows_are_untouched(tmp_path):
    atoms = _parse(tmp_path)
    a = _site_atom(atoms, 1, table=1)
    assert a.raw_text == "Rack the switch | Provider"
    assert "checkbox_row" not in a.value and "row_key" not in _loc(a)


def test_envelope_names_the_kept_site_atom_for_each_row():
    loc = lambda row, **kw: {"table_index": 1, "row": row, **kw}  # noqa: E731
    rows = [
        {"id": "site1", "artifact_id": "a", "atom_type": "physical_site", "locator": loc(1), "structured": {}},
        {"id": "row1", "artifact_id": "a", "atom_type": "site_attribute", "locator": loc(1),
         "structured": {"kind": "checkbox_selection", "checkbox_row": True, "row_key": "t1 r1",
                        "row_atom_id": "site1"}},
        {"id": "label2", "artifact_id": "a", "atom_type": "scope_item", "locator": loc(2), "structured": {}},
        {"id": "row2", "artifact_id": "a", "atom_type": "site_attribute", "locator": loc(2),
         "structured": {"kind": "checkbox_selection", "checkbox_row": True, "row_key": "t1 r2",
                        "row_atom_id": "label2"}},
        {"id": "row3", "artifact_id": "a", "atom_type": "scope_item", "locator": loc(3), "structured": {}},
        {"id": "cb3", "artifact_id": "b", "atom_type": "site_attribute", "locator": loc(1),
         "structured": {"kind": "checkbox_selection", "row_key": "t1 r1"}},
    ]
    out = {r["id"]: r["structured"] for r in _tie_checkbox_cells_to_site(rows)}
    assert out["row1"]["site_atom_id"] == "site1"  # the physical_site wins
    assert out["row2"]["site_atom_id"] == "label2"  # else the row's own site cell
    assert "site_atom_id" not in out["row3"]
    assert "site_atom_id" not in out["cb3"]  # never across documents


def _compile(proj: Path, monkeypatch=None):
    import os

    from app.core.compiler import compile_project

    os.environ.setdefault("SOWSMITH_DISABLE_LLM", "1")
    return compile_project(proj, project_id="p", allow_errors=True, use_cache=False)


def test_sites_still_found_and_read_from_the_site_cell(tmp_path):
    proj = tmp_path / "proj"
    proj.mkdir()
    _doc(proj / "sow.docx")
    r = _compile(proj)
    sites = [a for a in r.atoms if a.atom_type.value == "physical_site"]
    names = {a.raw_text for a in sites}
    for site in SITES:
        assert site in names, (site, names)
    for a in sites:
        if a.raw_text in SITES:
            assert a.source_refs[0].locator.get("cells") == [0]
    # The checkbox row is never a site of its own.
    assert not [a for a in sites if ":" in a.raw_text]


EML = """From: Pat Customer <pat@customer.example>
To: Sam Vendor <sam@vendor.example>
Subject: RE: Multi site support
Date: Tue, 22 Sep 2026 21:20:19 +0000
Message-ID: <r12@customer.example>
MIME-Version: 1.0
Content-Type: text/plain; charset=utf-8

Any update on this one?
Pat

From: Sam Vendor <sam@vendor.example>
Sent: Thursday, June 4, 2026 6:07 PM
To: Pat Customer <pat@customer.example>
Subject: Multi site support

Below is the original request.
Locations
Springfield, IL
Salem, OR
Dover, NH
Concord, NH
Sam
"""


def test_a_model_typed_grid_row_never_displaces_the_email_sites(tmp_path, monkeypatch):
    """The live compile's model typed "Riverton, WY: Install, Support"
    a physical_site (name = address = site_id = "Riverton, WY"); folded into the
    email's quoted "Riverton, WY", that value made the email's site read as a
    hallucinated one and it was dropped, with the site count. The row atom is
    the checkbox row now, not a promotable line; the site cell the model
    reads is the bare site it always read."""
    import re

    import app.core.typed_atom_classifier as tac
    from app.core.schemas import AtomType

    city = re.compile(r"^([A-Z][A-Za-z .]+, [A-Z]{2})(?::|$)")

    def model(atoms):
        n = 0
        for a in atoms:
            if tac._atom_type_str(a) not in tac._PROMOTABLE_FROM:
                continue
            m = city.match((a.raw_text or "").strip())
            if not m:
                continue
            a.atom_type = AtomType.physical_site
            if ":" in a.raw_text:
                c = m.group(1)
                a.value = {**(a.value or {}), "site_id": c, "name": c, "address": c}
            n += 1
        return n

    monkeypatch.setattr(tac, "classify_atoms", model)
    proj = tmp_path / "proj"
    proj.mkdir()
    _doc(proj / "sow.docx")
    (proj / "thread.eml").write_text(EML)
    r = _compile(proj)
    mail = {a.raw_text for a in r.atoms
            if a.atom_type.value == "physical_site" and a.source_refs[0].filename == "thread.eml"}
    assert set(SITES) <= mail, mail
