"""A colon list item governs only the items indented under it (010003).

The signed SOW's list has "Install the display:" (a top-level item) with one
"o" sub-step under it, then more top-level steps. In the Word draft the colon
item opened a sub-section that only a non-list paragraph closed, so every
later step of the list, siblings included, carried "Install the display" in
its section_path. The sub-section now closes at the next item of the list at
the colon item's own w:ilvl or shallower.
"""
from __future__ import annotations

from pathlib import Path

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.parsers.docx_parser import DocxParser

ITEMS = [
    (0, "Send one field technician to the office"),
    (0, "Unpack and check the three meeting room panels"),
    (0, "Mount the panel:"),
    (1, "Brackets are already fixed to the wall and ready for the panels"),
    (0, "Tighten and square the panel so it hangs safely"),
    (0, "Route the cables along the surface only"),
    (0, "Tidy the room when the work is done"),
]


def _numbered(paragraph, ilvl: int) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    numPr = OxmlElement("w:numPr")
    lvl = OxmlElement("w:ilvl")
    lvl.set(qn("w:val"), str(ilvl))
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), "1")
    numPr.append(lvl)
    numPr.append(num)
    pPr.append(numPr)


def _docx(tmp: Path) -> Path:
    d = docx.Document()
    d.add_heading("PROJECT SCOPE", 1)
    d.add_paragraph("The provider will do the following:")
    for ilvl, text in ITEMS:
        _numbered(d.add_paragraph(text), ilvl)
    d.add_paragraph("Work will take place during regular business hours.")
    out = tmp / "sow.docx"
    d.save(str(out))
    return out


def _paths(tmp: Path) -> dict[str, list[str]]:
    atoms = DocxParser().parse_artifact("p", "art", _docx(tmp))
    return {a.raw_text: list(a.source_refs[0].locator.get("section_path") or []) for a in atoms}


def test_sub_step_sits_under_its_colon_item(tmp_path: Path) -> None:
    got = _paths(tmp_path)
    assert got[ITEMS[3][1]][-1] == "Mount the panel"


def test_siblings_after_the_sub_step_leave_the_sub_section(tmp_path: Path) -> None:
    got = _paths(tmp_path)
    first = got[ITEMS[0][1]]
    for _, text in ITEMS[4:]:
        assert got[text] == first, (text, got[text])
        assert "Mount the panel" not in got[text]
    assert "Mount the panel" not in got["Work will take place during regular business hours."]
