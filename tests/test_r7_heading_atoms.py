"""R7 (000132, 010087): a heading that leads child lines is its own atom.

000132's SOW drafts open nine lettered groups ("A. ..." .. "I. ...") under
"Project Scope", and 010087's open "SCOPE OF WORK", "OUT OF SCOPE",
"CUSTOMER RESPONSIBILITIES". Every one of those lines went to the
suppressed sidecar at stage ``section_heading``: it lived only in its
children's section_path, so no label could govern a group. A heading that
leads child lines now comes out as a ``block_kind: heading`` atom whose
section_path is its children's, ending at itself. The document title and a
heading with nothing under it stay suppressed. Nothing else changes: every
other atom keeps its id, text and reading-order index, and no dedup folds a
heading onto a child or onto the next draft's copy of itself.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from app.parsers.docx_parser import DocxParser

fitz = pytest.importorskip("fitz")


@pytest.fixture(autouse=True)
def _no_llm(monkeypatch):
    monkeypatch.setenv("SOWSMITH_DISABLE_LLM", "1")


LEADS = ("A. Network Support", "B. Server Support", "C. End User Support")


def _child(lead: str, k: int) -> str:
    return f"Group {lead[0]} duty {k}: maintain and monitor the managed equipment"


def _draft(path: Path) -> None:
    d = Document()
    d.add_heading("Services Proposal", 0)
    d.add_heading("Project Scope", 1)
    d.add_paragraph("Provider will deliver recurring on site support services for the customer.")
    for lead in LEADS:
        p = d.add_paragraph()
        p.add_run(lead).bold = True
        for k in (1, 2):
            d.add_paragraph(_child(lead, k), style="List Bullet")
    d.add_heading("Appendix", 1)  # nothing under it
    d.add_heading("Assumptions", 1)
    d.add_paragraph("Customer will provide site access during business hours.")
    d.save(str(path))


def _loc(a) -> dict:
    return a.source_refs[0].locator or {}


def test_docx_heading_that_leads_lines_is_a_heading_atom(tmp_path):
    path = tmp_path / "sow.docx"
    _draft(path)
    atoms = DocxParser().parse_artifact("p", "a", path)
    by_text = {a.raw_text: a for a in atoms}
    for lead in LEADS:
        a = by_text[lead]
        assert _loc(a).get("block_kind") == "heading"
        assert _loc(a)["section_path"] == ["Services Proposal", "Project Scope", lead]
        assert not any(str(f).startswith("suppressed:") for f in a.review_flags)
        assert getattr(a.atom_type, "value", a.atom_type) == "deal_metadata"
        child = by_text[_child(lead, 1)]
        assert _loc(child)["section_path"] == _loc(a)["section_path"]
        assert _loc(child).get("block_kind") is None
    for led in ("Project Scope", "Assumptions"):
        assert _loc(by_text[led]).get("block_kind") == "heading", led
    # The title and a heading with nothing under it stay section metadata.
    for bare in ("Services Proposal", "Appendix"):
        assert "suppressed:section_heading" in by_text[bare].review_flags, bare
        assert _loc(by_text[bare]).get("block_kind") is None


def test_two_drafts_each_keep_their_heading_atoms(tmp_path):
    from app.core.compiler import compile_project

    deal = tmp_path / "deal"
    deal.mkdir()
    _draft(deal / "SOW v1.docx")
    _draft(deal / "SOW v2.docx")
    r = compile_project(deal, project_id="p", allow_errors=True, use_cache=False)
    for lead in LEADS:
        heads = [a for a in r.atoms if a.raw_text == lead]
        assert len(heads) == 2, (lead, len(heads))
        assert {a.source_refs[0].filename for a in heads} == {"SOW v1.docx", "SOW v2.docx"}
        for a in heads:
            assert _loc(a).get("block_kind") == "heading"
            assert not (a.value or {}).get("duplicate_of"), lead
            assert getattr(a.atom_type, "value", a.atom_type) == "deal_metadata"
    assert not [a for a in r.suppressed_atoms if a.raw_text in LEADS]


def test_pdf_heading_atom_keeps_every_other_reading_index(tmp_path):
    from app.parsers.orbitbrief_pdf import OrbitBriefPdfParser

    pdf = tmp_path / "sow.pdf"
    doc = fitz.open()
    page = doc.new_page(width=612, height=792)
    y = 60
    for head, lines in (
        ("SCOPE OF WORK", ["Provider will install two wireless access points per site.",
                           "Provider will mount the rack and dress all patch cables."]),
        ("CUSTOMER RESPONSIBILITIES", ["Customer will provide site access during business hours."]),
    ):
        page.insert_text((36, y), head, fontsize=11, fontname="hebo")
        y += 20
        for line in lines:
            page.insert_text((36, y), line, fontsize=10)
            y += 14
        y += 14
    doc.save(str(pdf))
    doc.close()
    atoms = list(OrbitBriefPdfParser().parse(pdf).atoms)
    heads = [a for a in atoms if _loc(a).get("block_kind") == "heading"]
    assert [a.raw_text for a in heads] == ["SCOPE OF WORK", "CUSTOMER RESPONSIBILITIES"]
    for a in heads:
        assert _loc(a)["section_path"][-1] == a.raw_text
    rest = [_loc(a).get("block_index") for a in atoms if _loc(a).get("block_kind") != "heading"]
    assert rest == list(range(len(rest)))
