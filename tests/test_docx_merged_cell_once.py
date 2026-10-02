"""A docx cell merged across grid columns is one cell, read once (010087).

python-docx's ``row.cells`` returns one entry per grid column, so a total-row
label spanning four columns (``w:gridSpan="4"``) came back four times and the
row atom read "ESTIMATED TOTAL FEES | ESTIMATED TOTAL FEES | ESTIMATED TOTAL
FEES | ESTIMATED TOTAL FEES | $9,504.00". The label is emitted once and the
row keeps its column alignment: the amount stays under the fees column.
"""

from __future__ import annotations

from docx import Document

from app.parsers.docx_parser import DocxParser

HEADER = ["DESCRIPTION", "STATED RATE (USD)", "BILLING INCREMENT", "EST. UNITS", "ESTIMATED FEES (USD)"]
BODY = ["Technician (est 2 hrs per site)", "$80.00", "Hourly", "40", "$3,200.00"]


def _fee_table(path):
    d = Document()
    t = d.add_table(rows=3, cols=5)
    for i, v in enumerate(HEADER):
        t.rows[0].cells[i].text = v
    for i, v in enumerate(BODY):
        t.rows[1].cells[i].text = v
    t.cell(2, 0).merge(t.cell(2, 3)).text = "ESTIMATED TOTAL FEES"
    t.cell(2, 4).text = "$3,200.00"
    assert 'w:gridSpan w:val="4"' in t._tbl.xml
    d.save(path)
    return DocxParser().parse_artifact("p", "a", path)


def _row(atoms, kind, row):
    return next(a for a in atoms if a.atom_type.value == kind
                and a.source_refs[0].locator.get("row") == row
                and "table_index" in a.source_refs[0].locator)


def test_a_spanned_label_is_read_once(tmp_path):
    atoms = _fee_table(tmp_path / "fees.docx")
    total = _row(atoms, "scope_item", 2)
    assert total.raw_text == "ESTIMATED TOTAL FEES | $3,200.00"
    assert not any(a.raw_text.count("ESTIMATED TOTAL FEES") > 1 for a in atoms)


def test_the_total_row_keeps_its_columns(tmp_path):
    atoms = _fee_table(tmp_path / "fees.docx")
    total = _row(atoms, "scope_item", 2)
    assert total.value["cells"] == {"DESCRIPTION": "ESTIMATED TOTAL FEES", "ESTIMATED FEES (USD)": "$3,200.00"}
    raw = _row(atoms, "raw_table_row", 2)
    assert raw.value["_row"] == ["ESTIMATED TOTAL FEES", "", "", "", "$3,200.00"]
    assert len(raw.value["_row"]) == len(raw.value["_columns"])


def test_an_unmerged_row_is_unchanged(tmp_path):
    atoms = _fee_table(tmp_path / "fees.docx")
    body = _row(atoms, "scope_item", 1)
    assert body.raw_text == " | ".join(BODY)
    assert body.value["cells"] == dict(zip(HEADER, BODY))
    assert _row(atoms, "raw_table_row", 1).value["_row"] == BODY
