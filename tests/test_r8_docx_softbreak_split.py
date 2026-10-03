"""A manual line break (w:br) inside a SOW paragraph ends the line before it.

Live 000132 (structure only; the text here is synthetic): a "Normal (Web)"
paragraph held a labelled line, a sentence with no full stop, a w:br, then a
second labelled line ("Label: ..."). paragraph.text reads the w:br as "\\n" and
the clause split whitespace-joined it, so the sentence and the next line came
out as one atom. In one draft that run-on then outranked the clean copy of the
same sentence (stated earlier, with its full stop) in cross-type dedup, so the
clean sentence was folded into the run-on; in the other draft the clean copy
stood. Both drafts hold the byte-identical paragraph.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from app.core.semantic_dedup import cross_type_dedup_atoms
from app.parsers.clause_split import split_clauses
from app.parsers.docx_parser import DocxParser

W = nsdecls("w")

SENTENCE = "All work under this order will be carried out in person at the client's listed sites"
LABEL_LINE = "Hours: Subject to crew availability, there is no response time target"


def _softbreak_paragraph(before: str, after: str) -> str:
    """The real paragraph's shape: bold label run, text run, a run holding only
    <w:br/>, a Strong-styled label run, then two text runs."""
    label, _, rest = after.partition(": ")
    return (
        f'<w:p {W}><w:pPr><w:pStyle w:val="NormalWeb"/><w:spacing w:line="240" w:lineRule="auto"/>'
        '<w:contextualSpacing/></w:pPr>'
        '<w:r><w:rPr><w:b/><w:bCs/></w:rPr><w:t>Onsite Support:</w:t></w:r>'
        f'<w:r><w:t xml:space="preserve"> Included. {before}</w:t></w:r>'
        '<w:r><w:br/></w:r>'
        f'<w:r><w:rPr><w:rStyle w:val="Strong"/></w:rPr><w:t>{label}:</w:t></w:r>'
        f'<w:r><w:t xml:space="preserve"> {rest.split(",")[0]}</w:t></w:r>'
        f'<w:r><w:t xml:space="preserve">,{rest.split(",", 1)[1]} </w:t></w:r></w:p>'
    )


def _atoms(tmp_path: Path, para_xml: str):
    d = Document()
    d.add_heading("Services Proposal", 1)
    d.add_heading("Project Scope", 2)
    d.add_paragraph(SENTENCE + ".")
    d.add_heading("Service Model", 3)
    d.element.body.append(parse_xml(para_xml))
    path = tmp_path / "SOW.docx"
    d.save(path)
    out = DocxParser().parse_artifact("p", "a", path)
    atoms = out if isinstance(out, list) else out.atoms
    return [a for a in atoms if a.source_refs[0].locator.get("paragraph_index") == 4]


def test_break_before_a_labelled_line_ends_the_sentence(tmp_path):
    para = _atoms(tmp_path, _softbreak_paragraph(SENTENCE, LABEL_LINE))
    assert [a.raw_text for a in para] == [
        "Onsite Support: Included.",
        SENTENCE,
        LABEL_LINE,
    ]
    assert [a.source_refs[0].locator["sentence_index"] for a in para] == [0, 1, 2]


def test_break_inside_a_sentence_stays_joined(tmp_path):
    # A w:br used as a visual wrap: the next line continues in lower case.
    xml = (
        f'<w:p {W}><w:r><w:t xml:space="preserve">Onsite Support: Included. All work under this '
        'order will be carried out in person at the client\'s listed sites and</w:t></w:r>'
        '<w:r><w:br/></w:r><w:r><w:t>during the hours the client has agreed, with no response '
        'time target unless one is set out in writing.</w:t></w:r></w:p>'
    )
    texts = [a.raw_text for a in _atoms(tmp_path, xml)]
    assert texts == [
        "Onsite Support: Included.",
        "All work under this order will be carried out in person at the client's listed sites and "
        "during the hours the client has agreed, with no response time target unless one is set out in writing.",
    ]


def test_layout_lines_are_not_hard_breaks():
    # The PDF path passes layout lines; a wrap there is not a w:br.
    text = f"Onsite Support: Included. {SENTENCE}\n{LABEL_LINE}"
    assert split_clauses(text) == ["Onsite Support: Included.", f"{SENTENCE} {LABEL_LINE}"]
    assert split_clauses(text, hard_breaks=True) == ["Onsite Support: Included.", SENTENCE, LABEL_LINE]


class _Atom:
    def __init__(self, atom_type, text):
        self.atom_type = atom_type
        self.raw_text = text
        self.text = text
        self.confidence = 0.8
        self.source_refs = [f"src::{atom_type}::{text[:10]}"]
        self.receipts = []
        self.entity_keys = []
        self.review_flags = []


def test_a_complete_sentence_never_folds_into_a_run_on():
    clean = _Atom("scope_item", SENTENCE + ".")
    run_on = _Atom("contract_term", f"{SENTENCE} {LABEL_LINE}")
    out = cross_type_dedup_atoms([clean, run_on])
    assert clean in out and run_on in out


def test_the_complete_copy_stands_whichever_type_outranks():
    for clean_type, cut_type in (("scope_item", "contract_term"), ("site_attribute", "contract_term")):
        clean = _Atom(clean_type, SENTENCE + ".")
        cut = _Atom(cut_type, SENTENCE)
        out = cross_type_dedup_atoms([clean, cut])
        assert out == [clean], (clean_type, [a.atom_type for a in out])


def test_a_sentence_and_its_paragraph_still_fold_into_the_paragraph():
    # The full stop is kept and more sentences follow: the paragraph is the
    # fuller copy of the same words (unchanged behaviour).
    first = _Atom("task", SENTENCE + ".")
    para = _Atom("exclusion", SENTENCE + ". Floor plans were shared for review and the crew will bring "
                 "its own ladders, carts and hand tools to every site.")
    out = cross_type_dedup_atoms([first, para])
    assert out == [para]
