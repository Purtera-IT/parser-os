"""A rate sheet is more than one table.

Under (or beside) a per-country matrix -- ``Country | Request | L1 Hourly ...``
-- a Deal Kit / SELL RATES sheet keeps a separate service-rate block
("L2 EUC 2 hour minimum | $96 | $63.50", "PC", "PM") with its own header row.
Its rows were bound to the matrix's header, so $96 read as "Request" and the
service "PC" read as a country: "Country: PC | Request: 50".
"""
from __future__ import annotations

from pathlib import Path

import openpyxl

from app.parsers.xlsx_parser import XlsxParser

COUNTRIES = [("United States", 50), ("Canada", 55), ("Mexico", 45), ("Brazil", 60), ("Germany", 70)]
MATRIX_HDR = ["Country", "Request", "L1 Hourly", "L2 Hourly", "L3 Hourly", "Dispatch", "After Hours"]
SERVICES = [("L2 EUC 2 hour minimum", 96, 63.5), ("L2 Networking", 110, 72), ("PC", 50, 30), ("PM", 120, 80)]


def _matrix(ws, pad=()):
    ws.append(MATRIX_HDR + list(pad))
    for c, v in COUNTRIES:
        ws.append([c, v, v + 10, v + 20, v + 30, v + 40, v + 50])


def _parse(tmp_path: Path, wb) -> list:
    p = tmp_path / "rates.xlsx"
    wb.save(p)
    return XlsxParser().parse(p)


def _texts(atoms):
    return [a.raw_text for a in atoms]


def _service_atoms(atoms):
    return [a for a in atoms if any(s[0] in a.raw_text for s in SERVICES[:2]) or "PC" in a.raw_text.split(" | ")[-3:][0:3]]


def _assert_services_not_countries(atoms):
    texts = _texts(atoms)
    for name, *_ in SERVICES:
        rows = [t for t in texts if f": {name} |" in t or f"| {name} |" in t]
        assert rows, (name, texts)
        for t in rows:
            assert f"Country: {name}" not in t, t
            assert "Request:" not in t, t


def test_service_block_under_its_own_header_after_blank_row(tmp_path):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "SELL RATES"
    _matrix(ws)
    ws.append([])
    ws.append(["Service", "Sell", "Cost"])
    for r in SERVICES:
        ws.append(list(r))
    atoms = _parse(tmp_path, wb)
    texts = _texts(atoms)
    _assert_services_not_countries(atoms)
    assert any("Service: L2 EUC 2 hour minimum | Sell: 96 | Cost: 63.5" in t for t in texts), texts
    assert any("Country: Canada | Request: 55" in t for t in texts), texts


def test_service_block_header_without_blank_row(tmp_path):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "SELL RATES"
    _matrix(ws)
    ws.append(["Service", "Sell", "Cost"])
    for r in SERVICES:
        ws.append(list(r))
    atoms = _parse(tmp_path, wb)
    _assert_services_not_countries(atoms)
    assert any("Service: PC | Sell: 50 | Cost: 30" in t for t in _texts(atoms))


def test_headerless_service_block_reads_plain(tmp_path):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "SELL RATES"
    _matrix(ws)
    ws.append([])
    for r in SERVICES:
        ws.append(list(r))
    atoms = _parse(tmp_path, wb)
    texts = _texts(atoms)
    _assert_services_not_countries(atoms)
    assert any(t.endswith("L2 EUC 2 hour minimum | 96 | 63.5") for t in texts), texts


def test_service_block_beside_the_matrix(tmp_path):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "SELL RATES"
    ws.append(MATRIX_HDR + ["", "Service", "Sell", "Cost"])
    for i, (c, v) in enumerate(COUNTRIES):
        side = list(SERVICES[i]) if i < len(SERVICES) else []
        ws.append([c, v, v + 10, v + 20, v + 30, v + 40, v + 50] + ([""] + side if side else []))
    atoms = _parse(tmp_path, wb)
    texts = _texts(atoms)
    _assert_services_not_countries(atoms)
    us = [t for t in texts if "Country: United States" in t]
    assert us and all("Service:" not in t for t in us), us
    assert any("Service: L2 EUC 2 hour minimum | Sell: 96 | Cost: 63.5" in t for t in texts), texts


def test_blank_row_inside_matrix_keeps_matrix_header(tmp_path):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "SELL RATES"
    _matrix(ws)
    ws.append([])
    ws.append(["France", 80, 90, 100, 110, 120, 130])
    atoms = _parse(tmp_path, wb)
    assert any("Country: France | Request: 80" in t for t in _texts(atoms)), _texts(atoms)


# ── "N pricing lines" next to the rows it counts ────────────────────────────


def test_pricing_lines_banner_is_suppressed_when_every_row_is_an_atom(tmp_path):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "SELL RATES"
    _matrix(ws)
    atoms = _parse(tmp_path, wb)
    (summary,) = [a for a in atoms if (a.value or {}).get("is_summary")]
    assert "pricing line" in summary.raw_text
    assert any(str(f).startswith("suppressed:") for f in summary.review_flags), summary.review_flags
    assert summary.value["_suppression"]["reason"]
    assert len([a for a in atoms if not (a.value or {}).get("is_summary")]) == len(COUNTRIES)


def test_compile_keeps_rows_and_diverts_the_banner(tmp_path):
    from app.core.compiler import compile_project

    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "SELL RATES"
    _matrix(ws)
    wb.save(tmp_path / "rates.xlsx")
    r = compile_project(tmp_path, project_id="p", allow_errors=True, use_cache=False)
    live = [a.raw_text for a in r.atoms]
    assert not [t for t in live if "pricing line" in t], live
    assert any("pricing line" in (a.raw_text or "") for a in r.suppressed_atoms)
    for country, _ in COUNTRIES:
        assert any(f"Country: {country} |" in t for t in live), (country, live)


def test_spacer_column_inside_one_table_does_not_split_its_rows():
    from app.core.schemas import ArtifactType
    from app.parsers.sheet_classifier import SheetClassification, SheetRole

    rows = [["Item", "Description", None, "Qty", "Unit Price", "Extended Price"]]
    rows += [[f"P-{i}", f"Part number {i}", None, 2 + i, 10.0 + i, (2 + i) * (10.0 + i)] for i in range(6)]
    atoms = XlsxParser()._emit_commercial_sheet_rows(
        project_id="p", artifact_id="a", artifact_type=ArtifactType.xlsx, filename="x.xlsx",
        sheet_name="Price List", rows=rows,
        classification=SheetClassification(role=SheetRole.CATALOG, suppress=True, reason="t", confidence=1.0),
    )
    rows_out = [a for a in atoms if not (a.value or {}).get("is_summary")]
    for i in range(6):
        hits = [a.raw_text for a in rows_out if f"P-{i}" in a.raw_text]
        assert len(hits) == 1, hits
        assert f"Unit Price: {10.0 + i:g}" in hits[0] or f"Unit Price: {10.0 + i}" in hits[0], hits
