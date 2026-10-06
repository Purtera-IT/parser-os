"""Every docx table whose first row is its column names gets a header pointer.

    Fees
    | Item          | Rate    | Hours | Total     |     <- column names, no atom
    | Widget setup  | $50.00  | 4     | $200.00   |
    | Gadget repair | $75.00  | 2     | $150.00   |

The header row was consumed as the columns and became no atom, so a labeler
reading "Widget setup | $50.00 | 4 | $200.00" had nothing on the page saying
what the three numbers are. #387 gave the checkbox grid's header row a pointer
atom; every table with a consumed header row now gets one the same way:
block_kind table_header, value.structure, held out of the model stages
(compiler._is_heading_atom), chatter, locator row 0, and every row atom of the
table names it (locator header_atom_id, handed on to the schema-typed row).

Nothing else moves: row atom ids are unchanged, and the pointer reads half a
step before the first row instead of taking a reading position, so every other
atom keeps its line_start (the walk keys repeated lines by it).
"""
from __future__ import annotations

from pathlib import Path

from docx import Document

from app.core.compiler import _is_heading_atom
from app.core.deal_chatter import CHATTER_FLAG
from app.core.entity_extraction import _enrich_table_atoms
from app.core.ids import stable_id
from app.core.service_router import _scope_summary
from app.parsers.docx_parser import DocxParser

HEADER = ["Item", "Rate", "Hours", "Total"]
ROWS = [
    ["Widget setup", "$50.00", "4", "$200.00"],
    ["Gadget repair", "$75.00", "2", "$150.00"],
]


def _table(doc, rows):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            t.cell(r, c).text = text
    return t


def _docx(tmp_path: Path, *, header=True, title_row=False, valued_first=False) -> Path:
    doc = Document()
    doc.add_heading("Fees", level=1)
    doc.add_paragraph("The work is billed as below.")
    rows = ([HEADER] if header else []) + ROWS
    if valued_first:
        rows = [["Version 1", "$10.00", "1", "$10.00"], *ROWS]
    t = _table(doc, rows)
    if title_row:
        t2 = _table(doc, [["Sign here"] * 3, ["Name:", "Title:", "Date:"], ["Pat", "Lead", "Jan 2"]])
        row0 = t2.rows[0].cells
        row0[0].merge(row0[2])
        row0[0].text = "Sign here"
    doc.add_paragraph("Totals are estimates.")
    p = tmp_path / "fees.docx"
    doc.save(p)
    return p


def _parse(path: Path):
    return DocxParser().parse_artifact_full(project_id="p", artifact_id="a", path=path).atoms


def _loc(a):
    return a.source_refs[0].locator


def _headers(atoms):
    return [a for a in atoms if _loc(a).get("block_kind") == "table_header"]


def test_header_row_is_one_pointer_atom(tmp_path):
    atoms = _parse(_docx(tmp_path))
    heads = _headers(atoms)
    assert len(heads) == 1
    h = heads[0]
    assert h.raw_text == "Item | Rate | Hours | Total"
    assert h.value["structure"] is True and h.value["chatter"] is True
    assert h.value["columns"] == HEADER
    assert CHATTER_FLAG in h.review_flags and "table_header" in h.review_flags
    assert _loc(h)["table_index"] == 0 and _loc(h)["row"] == 0
    assert h.id == stable_id("atm", "a", "docx_table_header", 0, 0, h.raw_text)
    # Held out of every model stage, like a heading.
    assert _is_heading_atom(h)


def test_every_row_atom_names_the_pointer(tmp_path):
    atoms = _parse(_docx(tmp_path))
    hid = _headers(atoms)[0].id
    rows = [a for a in atoms if _loc(a).get("table_index") == 0 and _loc(a).get("block_kind") != "table_header"]
    assert rows
    for a in rows:
        assert _loc(a)["header_atom_id"] == hid
        # The link rides on the locator: the row's value is as it was.
        assert "header_atom_id" not in (a.value or {})
    row_atoms = [a for a in rows if a.source_refs[0].extraction_method == "docx_table_row_v1"]
    assert {a.id for a in row_atoms} == {
        stable_id("atm", "a", "docx_row", 0, i + 1, " | ".join(r)) for i, r in enumerate(ROWS)
    }


def test_pointer_takes_no_reading_position(tmp_path):
    with_header = _parse(_docx(tmp_path))
    h = _headers(with_header)[0]
    others = [a for a in with_header if a is not h]
    # Every other atom keeps a whole-number position, 0..n-1, unbroken.
    assert sorted(_loc(a)["line_start"] for a in others) == list(range(len(others)))
    first_row = min(
        (a for a in others if _loc(a).get("table_index") == 0), key=lambda a: _loc(a)["line_start"]
    )
    assert _loc(h)["line_start"] == _loc(first_row)["line_start"] - 0.5


def test_no_pointer_without_a_header_row(tmp_path):
    # First row carries values: it is a data row (its own atom), not names.
    assert _headers(_parse(_docx(tmp_path, valued_first=True))) == []


def test_no_pointer_for_a_merged_title_row(tmp_path):
    atoms = _parse(_docx(tmp_path, title_row=True))
    # The fee table gets one; the titled signature block does not.
    assert [_loc(h)["table_index"] for h in _headers(atoms)] == [0]


def test_typed_row_inherits_the_link_not_in_its_value(tmp_path):
    doc = Document()
    _table(doc, [["SKU", "Qty", "Unit Cost"], ["AB-100", "3", "$12.00"]])
    p = tmp_path / "bom.docx"
    doc.save(p)
    atoms = _parse(p)
    hid = _headers(atoms)[0].id
    typed = _enrich_table_atoms(atoms, project_id="p")
    assert typed
    for a in typed:
        assert _loc(a)["header_atom_id"] == hid
        assert "header_atom_id" not in a.value


def test_pointer_stays_out_of_the_scope_summary(tmp_path):
    atoms = _parse(_docx(tmp_path))
    summary = _scope_summary(atoms, [{"filename": "fees.docx"}])
    assert "Item | Rate | Hours | Total" not in summary
